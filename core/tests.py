import io
import json
import sys
import tempfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
import pandas as pd
from django.apps import apps as django_apps
from django.contrib.auth import authenticate, get_user_model
from django.contrib.auth.tokens import default_token_generator
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode
from PIL import Image

from .models import (
    FIRMA_MAX_BYTE,
    AggiornamentoBC,
    CommessaPin,
    DestinazioneDocumento,
    Documento,
    EsecuzioneSchedulata,
    FirmatarioStabilimento,
    IndirizzoStabilimento,
    IndirSped,
    Notifica,
    Permesso,
    PersonaCommessa,
    Reparto,
    Revisione,
    RevisioneFileLink,
    RigaTransmittalInterno,
    RuoloFirmatarioStabilimento,
    RuoloPersonaCommessa,
    Segnalazione,
    SegnalazioneCommento,
    SegnalazioneVoto,
    Stabilimento,
    StatoEsterno,
    StatoSegnalazione,
    Testata,
    TipoIndirizzoStabilimento,
    TipoSegnalazione,
    Transmittal,
    TransmittalInterno,
)
from .services import scheduler
from .services.bc_sync import (
    BusinessCentralNonDisponibile,
    confronta_commessa,
    list_aggiornamenti,
    sincronizza_commesse,
    sincronizza_sito_costruttivo,
)
from .services.commesse import (
    MAX_PINNED_COMMESSE,
    create_commessa,
    fetch_from_bc,
    list_documenti,
    list_home_commesse,
    list_situazione,
    list_stati_esterni,
    persone_per_ruolo,
    pin_commessa,
    revisioni_by_doc_for_job,
    risolvi_file_revisione,
    risolvi_persona_commessa,
    salva_file_link,
    serialize_revisione,
)
from .services.export_grezzo import _cell as _cella_grezza
from .services.export_grezzo import build_workbook as build_dati_grezzi_workbook
from .services.export_grezzo import list_tabelle as list_tabelle_grezze
from .services.export_grezzo import resolve_tabelle as resolve_tabelle_grezze
from .services.import_old import importa_commessa_da_access
from .services.notifiche import count_notifiche, list_notifiche, segna_lette
from .services.organizzazione_commesse import (
    _dividi_nomi,
    _normalizza_job,
    backfill_persone_commessa,
    leggi_organizzazione_commesse,
    persone_per_job,
    risolvi_persone_libere,
    trova_utente_per_cognome,
)
from .services.revisione_anomalie import (
    audit_commessa,
    audit_commessa_summary,
    audit_revisione,
    classifica_revisione,
    ignora_anomalie_revisione,
    serialize_anomalie_gruppi,
)
from .services.revisione_label import (
    format_revisione_label,
    lettera_a_numero,
    numero_a_lettera,
)
from .services.revisione_sblocco import list_revisioni_sbloccabili, sblocca_revisione
from .services.revisioni_cleanup import drop_orphan_revisioni, find_orphan_indices
from .services.stato_esterno_codes import DEFAULT_STATUS_COLORS, letter_for_status_name
from .services.stato_esterno_colori import (
    FALLBACK_BG,
    MIN_CONTRAST,
    TEXT_DARK,
    TEXT_LIGHT,
    cell_colors,
    contrast_ratio,
    hex_to_rgb,
    readable_text_color,
    rgb_to_hex,
)
from .services.stato_esterno_legenda import legenda_default, legenda_stati_esterni
from .services.trasmittal_interno import (
    componi_nome,
    crea_trasmittal_interno,
    data_impegno,
    imposta_destinazioni,
    indirizzi_per_siti,
    prepara_per_dcc,
    prossimo_progressivo,
    salva_pdf,
    siti_coinvolti,
    siti_del_documento,
)

User = get_user_model()


# ── Shared fixture factory (unit tests only — uses temp dir) ──────────────────


def _make_fixtures_mock(reparto_acronimo: str = "QC"):
    """Create DB objects for unit tests that mock the filesystem.

    Args:
        reparto_acronimo: Acronimo to assign to the test reparto.

    Returns:
        Dict with created model instances keyed by name.
    """
    reparto = Reparto.objects.create(nome="Qualità e Controllo", acronimo=reparto_acronimo)
    testata = Testata.objects.create(job="25056")

    doc_qmdbi = Documento.objects.create(
        testata=testata,
        vendor_doc="25056-QMDBI",
        reparto=reparto.nome,
    )
    doc_qcpa = Documento.objects.create(
        testata=testata,
        vendor_doc="25056-01-QCPA",
        reparto=reparto.nome,
    )

    rev_qmdbi_0 = Revisione.objects.create(
        documento=doc_qmdbi,
        rev_no=0,
        dis_act_date="2025-09-05",
        rec_act_date="2025-11-17",
    )
    rev_qcpa_0 = Revisione.objects.create(
        documento=doc_qcpa,
        rev_no=0,
        dis_act_date="2025-09-05",
        rec_act_date="2025-10-14",
    )
    rev_qcpa_1 = Revisione.objects.create(
        documento=doc_qcpa,
        rev_no=1,
        dis_act_date="2025-11-07",
        rec_act_date="2025-11-10",
    )
    rev_qcpa_2 = Revisione.objects.create(
        documento=doc_qcpa,
        rev_no=2,
        dis_act_date="2025-11-13",
        rec_act_date="2025-11-16",
    )

    return {
        "reparto": reparto,
        "testata": testata,
        "doc_qmdbi": doc_qmdbi,
        "doc_qcpa": doc_qcpa,
        "rev_qmdbi_0": rev_qmdbi_0,
        "rev_qcpa_0": rev_qcpa_0,
        "rev_qcpa_1": rev_qcpa_1,
        "rev_qcpa_2": rev_qcpa_2,
    }


# ── Unit tests (mock filesystem) ─────────────────────────────────────────────


class RisolviFileRevisioneTest(TestCase):
    """Unit tests for ``risolvi_file_revisione`` using a mocked filesystem."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.objs = _make_fixtures_mock()

    def tearDown(self):
        self.patcher.stop()

    def _ricevuti(self, date_str: str) -> Path:
        d = Path(self.tmp) / "25056" / "PROGETTO" / "QC" / "RICEVUTI" / date_str
        d.mkdir(parents=True, exist_ok=True)
        return d

    def _base(self) -> Path:
        d = Path(self.tmp) / "25056" / "PROGETTO" / "QC"
        d.mkdir(parents=True, exist_ok=True)
        return d

    def test_found_single_file_in_ricevuti(self):
        d = self._ricevuti("2025-11-17")
        (d / "25056-QMDBI Rev0.pdf").touch()

        result = risolvi_file_revisione(self.objs["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "found")
        self.assertEqual(len(result["files"]), 1)
        self.assertEqual(result["files"][0]["nome"], "25056-QMDBI Rev0.pdf")

    def test_multiple_files_in_ricevuti(self):
        d = self._ricevuti("2025-11-17")
        (d / "25056-QMDBI Rev0.pdf").touch()
        (d / "25056-QMDBI Rev0 Annotated.pdf").touch()

        result = risolvi_file_revisione(self.objs["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "multiple")
        self.assertEqual(len(result["files"]), 2)

    def test_not_found_no_matching_file(self):
        d = self._ricevuti("2025-11-17")
        (d / "OTHER-DOC.pdf").touch()

        result = risolvi_file_revisione(self.objs["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "not_found")

    def test_not_found_missing_directory(self):
        result = risolvi_file_revisione(self.objs["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "not_found")

    def test_found_in_base_path_when_only_sent(self):
        """Revision with dis_act_date only → look in base path."""
        rev_sent = Revisione.objects.create(
            documento=self.objs["doc_qcpa"],
            rev_no=3,
            dis_act_date="2025-12-01",
            rec_act_date=None,
        )
        d = self._base()
        (d / "25056-01-QCPA Rev3.pdf").touch()

        result = risolvi_file_revisione(rev_sent.pk)
        self.assertEqual(result["status"], "found")
        self.assertEqual(result["files"][0]["estensione"], "pdf")

    def test_not_found_no_dates(self):
        """Revision with neither date → not_found."""
        rev_none = Revisione.objects.create(documento=self.objs["doc_qcpa"], rev_no=4)
        result = risolvi_file_revisione(rev_none.pk)
        self.assertEqual(result["status"], "not_found")

    def test_manual_link_takes_priority(self):
        d = self._ricevuti("2025-11-17")
        (d / "25056-QMDBI Rev0.pdf").touch()
        manual_path = str(Path(self.tmp) / "25056" / "PROGETTO" / "QC" / "25056-QMDBI-MANUAL.pdf")
        RevisioneFileLink.objects.create(
            revisione=self.objs["rev_qmdbi_0"],
            percorso=manual_path,
        )

        result = risolvi_file_revisione(self.objs["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "manual_linked")
        self.assertEqual(result["files"][0]["percorso"], manual_path)

    def test_qcpa_each_revision_in_own_date_folder(self):
        """Each QCPA revision must resolve in its own RICEVUTI date folder."""
        cases = [
            ("rev_qcpa_0", "2025-10-14"),
            ("rev_qcpa_1", "2025-11-10"),
            ("rev_qcpa_2", "2025-11-16"),
        ]
        for rev_key, date_str in cases:
            d = self._ricevuti(date_str)
            (d / f"25056-01-QCPA_{date_str}.pdf").touch()
            result = risolvi_file_revisione(self.objs[rev_key].pk)
            self.assertEqual(result["status"], "found", f"Expected found for {rev_key}")


class SalvaFileLinkTest(TestCase):
    """Unit tests for ``salva_file_link``."""

    def setUp(self):
        objs = _make_fixtures_mock()
        self.rev = objs["rev_qmdbi_0"]

    def test_create_link(self):
        path = r"Z:\JOBS\25056\PROGETTO\QC\25056-QMDBI Rev0.pdf"
        link = salva_file_link(self.rev.pk, path)
        self.assertEqual(link.percorso, path)
        self.assertEqual(RevisioneFileLink.objects.count(), 1)

    def test_update_existing_link(self):
        salva_file_link(self.rev.pk, r"Z:\JOBS\25056\old.pdf")
        salva_file_link(self.rev.pk, r"Z:\JOBS\25056\new.pdf")
        self.assertEqual(RevisioneFileLink.objects.count(), 1)
        self.assertEqual(
            RevisioneFileLink.objects.get(revisione=self.rev).percorso,
            r"Z:\JOBS\25056\new.pdf",
        )


class FileBrowserApiTest(TestCase):
    """Unit tests for the fileserver browse API using a mocked filesystem."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.user = User.objects.create_user(
            "testuser_browse", password="pw", permesso=Permesso.WRITING
        )
        self.client = Client()
        self.client.force_login(self.user)

    def tearDown(self):
        self.patcher.stop()

    def test_browse_root(self):
        (Path(self.tmp) / "25056").mkdir()
        response = self.client.get("/api/fileserver/browse/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("dirs", data)
        self.assertIn("files", data)
        self.assertIsNone(data["parent"])
        self.assertIn("25056", [d["nome"] for d in data["dirs"]])

    def test_browse_subdir(self):
        subdir = Path(self.tmp) / "25056"
        subdir.mkdir()
        (subdir / "test.pdf").touch()
        response = self.client.get("/api/fileserver/browse/", {"path": str(subdir)})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["parent"], self.tmp)
        self.assertIn("test.pdf", [f["nome"] for f in data["files"]])

    def test_browse_outside_root_forbidden(self):
        response = self.client.get("/api/fileserver/browse/", {"path": "C:\\Windows"})
        self.assertEqual(response.status_code, 403)

    def test_browse_nonexistent_path(self):
        response = self.client.get(
            "/api/fileserver/browse/",
            {"path": str(Path(self.tmp) / "doesnotexist")},
        )
        self.assertEqual(response.status_code, 404)

    def test_browse_requires_auth(self):
        response = Client().get("/api/fileserver/browse/")
        self.assertIn(response.status_code, [302, 401, 403])


class RevisioneFileMockApiTest(TestCase):
    """Unit tests for the revision file API endpoints using a mocked filesystem."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.user = User.objects.create_user(
            "testuser_api", password="pw", permesso=Permesso.WRITING
        )
        self.client = Client()
        self.client.force_login(self.user)
        objs = _make_fixtures_mock()
        self.rev = objs["rev_qmdbi_0"]

    def tearDown(self):
        self.patcher.stop()

    def test_resolve_found(self):
        d = Path(self.tmp) / "25056" / "PROGETTO" / "QC" / "RICEVUTI" / "2025-11-17"
        d.mkdir(parents=True)
        (d / "25056-QMDBI Rev0.pdf").touch()

        response = self.client.get(f"/api/revisioni/{self.rev.pk}/file/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "found")
        self.assertEqual(len(data["files"]), 1)

    def test_resolve_not_found(self):
        response = self.client.get(f"/api/revisioni/{self.rev.pk}/file/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "not_found")

    def test_resolve_nonexistent_revision(self):
        response = self.client.get("/api/revisioni/999999/file/")
        self.assertEqual(response.status_code, 404)

    def test_save_link(self):
        path = str(Path(self.tmp) / "25056" / "test.pdf")
        response = self.client.post(
            f"/api/revisioni/{self.rev.pk}/file/link/",
            data=json.dumps({"percorso": path}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["percorso"], path)
        self.assertTrue(RevisioneFileLink.objects.filter(revisione=self.rev).exists())

    def test_save_link_missing_percorso(self):
        response = self.client.post(
            f"/api/revisioni/{self.rev.pk}/file/link/",
            data=json.dumps({}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)

    def test_save_link_requires_auth(self):
        response = Client().post(
            f"/api/revisioni/{self.rev.pk}/file/link/",
            data=json.dumps({"percorso": "Z:\\x.pdf"}),
            content_type="application/json",
        )
        self.assertIn(response.status_code, [302, 401, 403])


def _trasmittal_dir(jobs_root, job: str) -> Path:
    folder = Path(jobs_root) / job / "PROGETTO" / "DCC" / "TRANSMITTAL"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _trasmittal_path(jobs_root, job: str, numero: int) -> Path:
    return (
        Path(jobs_root)
        / job
        / "PROGETTO"
        / "DCC"
        / "TRANSMITTAL"
        / f"Transmittal {job}-{numero}.pdf"
    )


def _write_trasmittal(jobs_root, job, numero, content=b"%PDF", filename=None):
    folder = _trasmittal_dir(jobs_root, job)
    name = filename or f"Transmittal {job}-{numero}.pdf"
    path = folder / name
    path.write_bytes(content)
    return path


class ListaTrasmittalTest(TestCase):
    """Unit tests for transmittal archive listing against a temp folder."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()

    def test_lists_matching_job_sorted_desc(self):
        folder = _trasmittal_dir(self.tmp, "25089")
        (folder / "Transmittal 25089-3.pdf").write_bytes(b"%PDF")
        (folder / "Transmittal 25089-21.pdf").write_bytes(b"%PDF")
        (folder / "Transmittal 25056-1.pdf").write_bytes(b"%PDF")
        (folder / "readme.txt").write_text("x")
        _write_trasmittal(self.tmp, "25056", 9)
        from core.services.trasmittal_archivio import lista_trasmittal

        items = lista_trasmittal("25089")
        self.assertEqual([i["id"] for i in items], [21, 3])
        self.assertEqual(items[0]["nome"], "Transmittal 25089-21.pdf")

    def test_case_insensitive_filename(self):
        _write_trasmittal(self.tmp, "25089", 2, filename="TRANSMITTAL 25089-2.PDF")
        from core.services.trasmittal_archivio import lista_trasmittal

        items = lista_trasmittal("25089")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["id"], 2)

    def test_italian_spelling_trasmittal(self):
        _write_trasmittal(self.tmp, "25089", 2, filename="Trasmittal 25089-2.pdf")
        from core.services.trasmittal_archivio import lista_trasmittal

        items = lista_trasmittal("25089")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["id"], 2)
        self.assertEqual(items[0]["nome"], "Trasmittal 25089-2.pdf")

    def test_leading_zeros_in_number(self):
        _write_trasmittal(self.tmp, "25089", 1, filename="Transmittal 25089-01.pdf")
        from core.services.trasmittal_archivio import lista_trasmittal

        items = lista_trasmittal("25089")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["id"], 1)
        self.assertEqual(items[0]["nome"], "Transmittal 25089-01.pdf")

    def test_missing_folder_returns_empty(self):
        from core.services.trasmittal_archivio import lista_trasmittal

        self.assertEqual(lista_trasmittal("25089"), [])

    def test_percorso_found(self):
        _write_trasmittal(self.tmp, "25089", 7, content=b"%PDF-1.4")
        from core.services.trasmittal_archivio import percorso_trasmittal

        path = percorso_trasmittal("25089", 7)
        self.assertTrue(path.is_file())
        self.assertEqual(path.name, "Transmittal 25089-7.pdf")
        self.assertEqual(path, _trasmittal_path(self.tmp, "25089", 7))

    def test_percorso_resolves_padded_filename(self):
        written = _write_trasmittal(self.tmp, "25089", 1, filename="Transmittal 25089-01.pdf")
        from core.services.trasmittal_archivio import percorso_trasmittal

        path = percorso_trasmittal("25089", 1)
        self.assertEqual(path, written)
        self.assertEqual(path.name, "Transmittal 25089-01.pdf")

    def test_sync_imports_italian_padded_filename(self):
        Testata.objects.create(job="25089")
        _write_trasmittal(self.tmp, "25089", 1, filename="Trasmittal 25089-01.pdf")
        from core.services.trasmittal_archivio import sync_trasmittal_da_cartella

        created = sync_trasmittal_da_cartella("25089")
        self.assertEqual(created, 1)
        row = Transmittal.objects.get(testata_id="25089", numero=1)
        self.assertIsNotNone(row.data_emissione)

    def test_percorso_wrong_job_not_found(self):
        _write_trasmittal(self.tmp, "25089", 7)
        from core.services.trasmittal_archivio import percorso_trasmittal

        with self.assertRaises(FileNotFoundError):
            percorso_trasmittal("25056", 7)

    def test_lists_from_da_spedire_when_dcc_missing(self):
        folder = Path(self.tmp) / "25089" / "PROGETTO" / "DCC" / "DA SPEDIRE" / "TRANSMITTAL"
        folder.mkdir(parents=True)
        (folder / "Transmittal 25089-5.pdf").write_bytes(b"%PDF")
        from core.services.trasmittal_archivio import (
            cartella_trasmittal,
            lista_trasmittal,
            trasmittal_in_da_spedire,
        )

        items = lista_trasmittal("25089")
        self.assertEqual([i["id"] for i in items], [5])
        self.assertTrue(trasmittal_in_da_spedire("25089"))
        self.assertEqual(cartella_trasmittal("25089"), folder)

    def test_prefers_dcc_over_da_spedire(self):
        _write_trasmittal(self.tmp, "25089", 1)
        fallback = Path(self.tmp) / "25089" / "PROGETTO" / "DCC" / "DA SPEDIRE" / "TRANSMITTAL"
        fallback.mkdir(parents=True)
        (fallback / "Transmittal 25089-9.pdf").write_bytes(b"%PDF")
        from core.services.trasmittal_archivio import lista_trasmittal, trasmittal_in_da_spedire

        items = lista_trasmittal("25089")
        self.assertEqual([i["id"] for i in items], [1])
        self.assertFalse(trasmittal_in_da_spedire("25089"))


class TrasmittalStoricoApiTest(TestCase):
    """API tests for transmittal history and file serve."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.user = User.objects.create_user(
            "testuser_trasmittal", password="pw", permesso=Permesso.READING
        )
        self.client = Client()
        self.client.force_login(self.user)
        Testata.objects.create(job="25089")

    def tearDown(self):
        self.patcher.stop()

    def test_storico_lists_files(self):
        _write_trasmittal(self.tmp, "25089", 1)
        _write_trasmittal(self.tmp, "25089", 4)
        response = self.client.get("/api/commesse/25089/trasmittal/storico/")
        self.assertEqual(response.status_code, 200)
        items = response.json()["items"]
        self.assertEqual([i["id"] for i in items], [4, 1])
        self.assertTrue(all(i.get("data") for i in items))
        self.assertEqual(Transmittal.objects.filter(testata_id="25089").count(), 2)

    def test_storico_lists_db_rows_without_file(self):
        Transmittal.objects.create(
            testata_id="25089",
            numero=8,
            data_emissione="2026-09-02",
        )
        response = self.client.get("/api/commesse/25089/trasmittal/storico/")
        self.assertEqual(response.status_code, 200)
        items = response.json()["items"]
        self.assertEqual(items[0]["id"], 8)
        self.assertEqual(items[0]["codice"], "25089-8")
        self.assertEqual(items[0]["data"], "2026-09-02")

    def test_storico_does_not_overwrite_existing_row(self):
        Transmittal.objects.create(
            testata_id="25089",
            numero=1,
            data_emissione="2020-01-01",
        )
        _write_trasmittal(self.tmp, "25089", 1)
        response = self.client.get("/api/commesse/25089/trasmittal/storico/")
        self.assertEqual(response.status_code, 200)
        row = Transmittal.objects.get(testata_id="25089", numero=1)
        self.assertEqual(str(row.data_emissione), "2020-01-01")

    def test_storico_skips_sync_when_job_already_has_rows(self):
        Transmittal.objects.create(
            testata_id="25089",
            numero=1,
            data_emissione="2020-01-01",
        )
        _write_trasmittal(self.tmp, "25089", 9)
        response = self.client.get("/api/commesse/25089/trasmittal/storico/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual([i["id"] for i in response.json()["items"]], [1])
        self.assertFalse(Transmittal.objects.filter(testata_id="25089", numero=9).exists())

    def test_storico_empty_includes_cartella_and_formato(self):
        from core.services.trasmittal_archivio import (
            cartella_trasmittal_canonica,
            formato_nome_file,
        )

        response = self.client.get("/api/commesse/25089/trasmittal/storico/")
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["items"], [])
        self.assertEqual(body["cartella"], str(cartella_trasmittal_canonica("25089")))
        self.assertEqual(body["formato"], formato_nome_file("25089"))
        self.assertEqual(body["formato"], "Transmittal 25089-{n}.pdf")
        self.assertFalse(body["in_da_spedire"])

    def test_storico_reads_from_da_spedire_and_flags(self):
        folder = Path(self.tmp) / "25089" / "PROGETTO" / "DCC" / "DA SPEDIRE" / "TRANSMITTAL"
        folder.mkdir(parents=True)
        (folder / "Transmittal 25089-4.pdf").write_bytes(b"%PDF")
        response = self.client.get("/api/commesse/25089/trasmittal/storico/?sync=1")
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["in_da_spedire"])
        self.assertEqual([i["id"] for i in body["items"]], [4])
        self.assertTrue(Transmittal.objects.filter(testata_id="25089", numero=4).exists())

    def test_storico_force_sync_imports_new_files(self):
        Transmittal.objects.create(
            testata_id="25089",
            numero=1,
            data_emissione="2020-01-01",
        )
        _write_trasmittal(self.tmp, "25089", 9)
        response = self.client.get("/api/commesse/25089/trasmittal/storico/?sync=1")
        self.assertEqual(response.status_code, 200)
        self.assertEqual([i["id"] for i in response.json()["items"]], [9, 1])
        self.assertTrue(Transmittal.objects.filter(testata_id="25089", numero=9).exists())

    def test_storico_unknown_job(self):
        response = self.client.get("/api/commesse/NOPE/trasmittal/storico/")
        self.assertEqual(response.status_code, 404)

    def test_storico_requires_auth(self):
        response = Client().get("/api/commesse/25089/trasmittal/storico/")
        self.assertIn(response.status_code, [302, 401, 403])

    def test_serve_pdf_inline(self):
        _write_trasmittal(self.tmp, "25089", 12, content=b"%PDF-fake")
        response = self.client.get("/api/commesse/25089/trasmittal/12/file/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn(b"%PDF-fake", b"".join(response.streaming_content))

    def test_serve_missing_id(self):
        response = self.client.get("/api/commesse/25089/trasmittal/99/file/")
        self.assertEqual(response.status_code, 404)

    def test_serve_other_job_file_hidden(self):
        _write_trasmittal(self.tmp, "11111", 1)
        response = self.client.get("/api/commesse/25089/trasmittal/1/file/")
        self.assertEqual(response.status_code, 404)

    def test_prossimo_api(self):
        from core.services.trasmittal_archivio import percorso_previsto

        _write_trasmittal(self.tmp, "25089", 3)
        Transmittal.objects.create(testata_id="25089", numero=5, data_emissione="2026-01-01")
        response = self.client.get("/api/commesse/25089/trasmittal/prossimo/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.json(),
            {
                "numero": 6,
                "codice": "25089-6",
                "percorso": str(percorso_previsto("25089", 6)),
            },
        )


class TransmittalArchivioServiceTest(TestCase):
    """DB + filesystem tests for transmittal numbering and emission archive."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.testata = Testata.objects.create(job="25089")
        self.doc = Documento.objects.create(testata=self.testata, vendor_doc="25089-A")
        self.rev = Revisione.objects.create(documento=self.doc, rev_no=0)

    def tearDown(self):
        self.patcher.stop()

    def test_prossimo_numero_from_files_and_db(self):
        from core.services.trasmittal_archivio import prossimo_numero

        _write_trasmittal(self.tmp, "25089", 3)
        Transmittal.objects.create(testata=self.testata, numero=5, data_emissione="2026-01-01")
        self.assertEqual(prossimo_numero("25089"), 6)

    def test_emetti_e_archivia_writes_file_and_links_revs(self):
        from core.services.trasmittal_archivio import emetti_e_archivia

        result = emetti_e_archivia("25089", [self.doc.pk], "2026-09-02", b"%PDF-emit")
        self.assertEqual(result["trasmittal"]["id"], 1)
        self.assertEqual(result["trasmittal"]["codice"], "25089-1")
        dest = _trasmittal_path(self.tmp, "25089", 1)
        self.assertTrue(dest.is_file())
        self.assertEqual(dest.read_bytes(), b"%PDF-emit")
        row = Transmittal.objects.get(testata=self.testata, numero=1)
        self.assertEqual(str(row.data_emissione), "2026-09-02")
        self.assertEqual(list(row.revisioni.values_list("pk", flat=True)), [self.rev.pk])
        self.rev.refresh_from_db()
        self.assertEqual(self.rev.int_status, "inviato_al_cliente")

    def test_annulla_restores_and_reuses_numero(self):
        from core.services.trasmittal_archivio import (
            annulla_trasmittal,
            emetti_e_archivia,
            lista_storico,
            prossimo_numero,
        )

        emetti_e_archivia("25089", [self.doc.pk], "2026-09-02", b"%PDF-emit")
        dest = _trasmittal_path(self.tmp, "25089", 1)
        self.assertTrue(dest.is_file())
        items = lista_storico("25089")
        self.assertTrue(items[0]["annullabile"])

        annulla_trasmittal("25089", 1)
        self.rev.refresh_from_db()
        self.assertEqual(self.rev.int_status, "")
        self.assertIsNone(self.rev.dis_act_date)
        self.assertIsNone(self.rev.rec_plan_date)
        self.assertFalse(Transmittal.objects.filter(testata=self.testata, numero=1).exists())
        self.assertFalse(dest.is_file())
        self.assertEqual(prossimo_numero("25089"), 1)

    def test_annulla_rejects_not_latest(self):
        from core.services.trasmittal_archivio import TrasmittalAnnullaError, emetti_e_archivia

        emetti_e_archivia("25089", [self.doc.pk], "2026-09-02", b"%PDF-1")
        doc2 = Documento.objects.create(testata=self.testata, vendor_doc="25089-B")
        Revisione.objects.create(documento=doc2, rev_no=0)
        emetti_e_archivia("25089", [doc2.pk], "2026-09-03", b"%PDF-2")
        from core.services.trasmittal_archivio import annulla_trasmittal

        with self.assertRaises(TrasmittalAnnullaError):
            annulla_trasmittal("25089", 1)

    def test_annulla_rejects_legacy_without_revs(self):
        from core.services.trasmittal_archivio import TrasmittalAnnullaError, annulla_trasmittal

        Transmittal.objects.create(testata=self.testata, numero=1, data_emissione="2026-01-01")
        with self.assertRaises(TrasmittalAnnullaError):
            annulla_trasmittal("25089", 1)

    def test_annulla_rejects_already_received(self):
        from core.services.trasmittal_archivio import (
            TrasmittalAnnullaError,
            annulla_trasmittal,
            emetti_e_archivia,
        )

        emetti_e_archivia("25089", [self.doc.pk], "2026-09-02", b"%PDF-emit")
        self.rev.refresh_from_db()
        self.rev.int_status = "ricevuto"
        self.rev.rec_act_date = "2026-09-10"
        self.rev.save(update_fields=["int_status", "rec_act_date"])
        with self.assertRaises(TrasmittalAnnullaError):
            annulla_trasmittal("25089", 1)


class EmissioneArchiviaApiTest(TestCase):
    """POST /api/emissione/ archives the PDF and returns transmittal id."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.patcher = patch("django.conf.settings.FILESERVER_JOBS_PATH", self.tmp)
        self.patcher.start()
        self.pdf_patcher = patch("core.views.genera_trasmittal_pdf", return_value=b"%PDF-api")
        self.pdf_patcher.start()
        self.user = User.objects.create_user(
            "testuser_emit_tr", password="pw", permesso=Permesso.WRITING
        )
        self.client = Client()
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="25089")
        self.doc = Documento.objects.create(testata=self.testata, vendor_doc="25089-A")
        Revisione.objects.create(documento=self.doc, rev_no=0)

    def tearDown(self):
        self.pdf_patcher.stop()
        self.patcher.stop()

    def test_emissione_saves_pdf_and_row(self):
        response = self.client.post(
            "/api/emissione/",
            data=json.dumps(
                {
                    "doc_ids": [self.doc.pk],
                    "dis_act_date": "2026-09-02",
                    "job": "25089",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["ok"])
        self.assertEqual(body["trasmittal"]["id"], 1)
        self.assertTrue(_trasmittal_path(self.tmp, "25089", 1).is_file())
        self.assertTrue(Transmittal.objects.filter(testata_id="25089", numero=1).exists())

    def test_annulla_api(self):
        emit = self.client.post(
            "/api/emissione/",
            data=json.dumps(
                {"doc_ids": [self.doc.pk], "dis_act_date": "2026-09-02", "job": "25089"}
            ),
            content_type="application/json",
        )
        self.assertEqual(emit.status_code, 200)
        response = self.client.post("/api/commesse/25089/trasmittal/1/annulla/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        self.assertFalse(Transmittal.objects.filter(testata_id="25089", numero=1).exists())
        self.assertFalse(_trasmittal_path(self.tmp, "25089", 1).is_file())

    def test_annulla_api_requires_write(self):
        User.objects.create_user(
            "reader_tr",
            email="reader_tr@example.com",
            password="pw",
            permesso=Permesso.READING,
        )
        reader = Client()
        reader.force_login(User.objects.get(username="reader_tr"))
        Transmittal.objects.create(testata=self.testata, numero=1, data_emissione="2026-09-02")
        response = reader.post("/api/commesse/25089/trasmittal/1/annulla/")
        self.assertEqual(response.status_code, 403)


class UserStabilimentoTest(TestCase):
    def test_stabilimento_optional(self):
        user = User.objects.create_user("stab_none", password="pw")
        self.assertIsNone(user.stabilimento_id)

    def test_stabilimento_assign(self):
        stab = Stabilimento.objects.create(nome="Bergamo")
        user = User.objects.create_user("stab_yes", password="pw", stabilimento=stab)
        self.assertEqual(user.stabilimento.nome, "Bergamo")


class StabilimentoSiglaCodiceBCTests(TestCase):
    """Sigla e codice_bc: unici ma non tutti gli stabilimenti li hanno."""

    def test_sigla_unica(self):
        Stabilimento.objects.create(nome="Uno", sigla="BG")

        with self.assertRaises(IntegrityError), transaction.atomic():
            Stabilimento.objects.create(nome="Due", sigla="BG")

    def test_codice_bc_unico(self):
        Stabilimento.objects.create(nome="Uno", codice_bc=1)

        with self.assertRaises(IntegrityError), transaction.atomic():
            Stabilimento.objects.create(nome="Due", codice_bc=1)

    def test_piu_stabilimenti_con_sigla_e_codice_null_convivono(self):
        # Non tutti gli stabilimenti sono siti costruttivi: NULL è ammesso più volte.
        Stabilimento.objects.create(nome="Milano")
        Stabilimento.objects.create(nome="Roma")

        self.assertEqual(
            set(Stabilimento.objects.values_list("nome", flat=True)), {"Milano", "Roma"}
        )


class FirmatarioStabilimentoTests(TestCase):
    """Un solo firmatario per stabilimento e ruolo; lo stesso utente può firmare più siti."""

    def setUp(self):
        self.stab = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.altro_stab = Stabilimento.objects.create(nome="Schio", sigla="VI", codice_bc=5)
        self.utente = User.objects.create_user("firmatario", password="pw")

    def test_un_solo_firmatario_per_ruolo_e_stabilimento(self):
        FirmatarioStabilimento.objects.create(
            stabilimento=self.stab,
            ruolo=RuoloFirmatarioStabilimento.PRODUZIONE,
            utente=self.utente,
        )

        with self.assertRaises(IntegrityError), transaction.atomic():
            FirmatarioStabilimento.objects.create(
                stabilimento=self.stab,
                ruolo=RuoloFirmatarioStabilimento.PRODUZIONE,
                utente=self.utente,
            )

    def test_lo_stesso_utente_firma_per_piu_stabilimenti(self):
        FirmatarioStabilimento.objects.create(
            stabilimento=self.stab,
            ruolo=RuoloFirmatarioStabilimento.QUALITA,
            utente=self.utente,
        )
        FirmatarioStabilimento.objects.create(
            stabilimento=self.altro_stab,
            ruolo=RuoloFirmatarioStabilimento.QUALITA,
            utente=self.utente,
        )

        self.assertEqual(self.utente.firmatario_di.count(), 2)


class MigrazioneSigleCodiciBCTests(TestCase):
    """La migrazione dati 0045 assegna sigla e codice_bc per nome."""

    def _modulo(self):
        import importlib

        return importlib.import_module("core.migrations.0045_popola_sigla_codice_bc_stabilimenti")

    def test_assegna_i_codici_giusti(self):
        for nome in ("Valbrembo", "Albignasego", "Marghera", "Ricengo", "Schio", "Milano"):
            Stabilimento.objects.create(nome=nome)

        self._modulo().popola(django_apps, None)

        attesi = {
            "Valbrembo": ("BG", 1),
            "Albignasego": ("PD", 2),
            "Marghera": ("VE", 3),
            "Ricengo": ("CR", 4),
            "Schio": ("VI", 5),
            "Milano": (None, None),
        }
        for nome, (sigla, codice_bc) in attesi.items():
            stab = Stabilimento.objects.get(nome=nome)
            self.assertEqual((stab.sigla, stab.codice_bc), (sigla, codice_bc))

    def test_reverse_svuota_solo_i_nomi_noti(self):
        Stabilimento.objects.create(nome="Ricengo", sigla="CR", codice_bc=4)
        Stabilimento.objects.create(nome="Altro", sigla="XX", codice_bc=99)

        self._modulo().svuota(django_apps, None)

        self.assertEqual(
            (
                Stabilimento.objects.get(nome="Ricengo").sigla,
                Stabilimento.objects.get(nome="Ricengo").codice_bc,
            ),
            (None, None),
        )
        # Uno stabilimento fuori dall'elenco della migrazione non è toccato.
        altro = Stabilimento.objects.get(nome="Altro")
        self.assertEqual((altro.sigla, altro.codice_bc), ("XX", 99))


class MigrazioneTpiStrutturatoTests(SimpleTestCase):
    """La migrazione dati 0051 converte il vecchio TPI testo libero in (bool, destinatario).

    La forma storica del campo (``tpi_testo``, prima della migrazione) non
    esiste più nel modello corrente: non si può costruire con un
    ``RigaTransmittalInterno`` reale. Le funzioni della migrazione operano
    solo su ``apps.get_model(...).objects.all()`` e ``riga.save(update_fields=...)``,
    quindi bastano dei doppi minimi con quella stessa forma.
    """

    def _modulo(self):
        import importlib

        return importlib.import_module(
            "core.migrations.0051_riga_trasmittal_interno_tpi_strutturato"
        )

    class _RigaFinta:
        def __init__(self, tpi_testo):
            self.tpi_testo = tpi_testo
            self.tpi_destinatario = ""
            self.tpi_bool = False
            self.salvata_con = None

        def save(self, update_fields):
            self.salvata_con = update_fields

    class _AppsFinto:
        def __init__(self, righe):
            self._righe = righe

        def get_model(self, app_label, model_name):
            righe = self._righe

            class _ManagerFinto:
                def all(self):
                    return righe

            return type("RigaTransmittalInternoFinto", (), {"objects": _ManagerFinto()})

    def test_tpi_testo_valorizzato_diventa_tpi_true_col_testo_come_destinatario(self):
        riga = self._RigaFinta(tpi_testo="AI")

        self._modulo().popola_tpi_strutturato(self._AppsFinto([riga]), None)

        self.assertTrue(riga.tpi_bool)
        self.assertEqual(riga.tpi_destinatario, "AI")
        self.assertEqual(riga.salvata_con, ["tpi_destinatario", "tpi_bool"])

    def test_tpi_testo_vuoto_diventa_tpi_false(self):
        riga = self._RigaFinta(tpi_testo="")

        self._modulo().popola_tpi_strutturato(self._AppsFinto([riga]), None)

        self.assertFalse(riga.tpi_bool)
        self.assertEqual(riga.tpi_destinatario, "")

    def test_tpi_testo_solo_spazi_diventa_tpi_false(self):
        riga = self._RigaFinta(tpi_testo="   ")

        self._modulo().popola_tpi_strutturato(self._AppsFinto([riga]), None)

        self.assertFalse(riga.tpi_bool)
        self.assertEqual(riga.tpi_destinatario, "")


class IndirizziPerSitiTests(TestCase):
    """indirizzi_per_siti: TO/CC per un elenco di siti, per il trasmittal interno."""

    def setUp(self):
        self.valbrembo = Stabilimento.objects.create(nome="Valbrembo", codice_bc=1)
        self.albignasego = Stabilimento.objects.create(nome="Albignasego", codice_bc=2)

    def _indirizzo(self, stabilimento, email, tipo, attivo=True):
        return IndirizzoStabilimento.objects.create(
            stabilimento=stabilimento, email=email, tipo=tipo, attivo=attivo
        )

    def test_input_vuoto(self):
        self._indirizzo(self.valbrembo, "a@b.it", TipoIndirizzoStabilimento.TO)

        self.assertEqual(indirizzi_per_siti([]), {"to": [], "cc": []})
        self.assertEqual(indirizzi_per_siti(None), {"to": [], "cc": []})

    def test_codici_inesistenti(self):
        self.assertEqual(indirizzi_per_siti([999]), {"to": [], "cc": []})

    def test_piu_siti_con_indirizzi_sovrapposti(self):
        self._indirizzo(self.valbrembo, "to1@b.it", TipoIndirizzoStabilimento.TO)
        self._indirizzo(self.albignasego, "to2@b.it", TipoIndirizzoStabilimento.TO)
        # Stesso indirizzo su entrambi i siti, maiuscole diverse: un solo TO.
        self._indirizzo(self.valbrembo, "comune@b.it", TipoIndirizzoStabilimento.TO)
        self._indirizzo(self.albignasego, "Comune@B.it", TipoIndirizzoStabilimento.TO)

        risultato = indirizzi_per_siti([1, 2])

        # Deduplica case-insensitive: dei due "comune" ne resta uno solo,
        # qualunque sia la maiuscola/minuscola sopravvissuta.
        self.assertEqual(len(risultato["to"]), 3)
        normalizzati = {e.lower() for e in risultato["to"]}
        self.assertEqual(normalizzati, {"to1@b.it", "to2@b.it", "comune@b.it"})
        self.assertEqual(risultato["cc"], [])

    def test_stesso_indirizzo_in_to_e_cc_compare_solo_in_to(self):
        self._indirizzo(self.valbrembo, "doppio@b.it", TipoIndirizzoStabilimento.TO)
        self._indirizzo(self.albignasego, "DOPPIO@b.it", TipoIndirizzoStabilimento.CC)
        self._indirizzo(self.albignasego, "solo_cc@b.it", TipoIndirizzoStabilimento.CC)

        risultato = indirizzi_per_siti([1, 2])

        self.assertEqual(risultato["to"], ["doppio@b.it"])
        self.assertEqual(risultato["cc"], ["solo_cc@b.it"])

    def test_indirizzi_inattivi_esclusi(self):
        self._indirizzo(self.valbrembo, "attivo@b.it", TipoIndirizzoStabilimento.TO)
        self._indirizzo(self.valbrembo, "inattivo@b.it", TipoIndirizzoStabilimento.TO, attivo=False)

        risultato = indirizzi_per_siti([1])

        self.assertEqual(risultato["to"], ["attivo@b.it"])


class DestinazioneDocumentoTests(TestCase):
    """Stabilimenti destinatari della copia cartacea di un documento."""

    def setUp(self):
        self.valbrembo = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.albignasego = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        self.marghera = Stabilimento.objects.create(nome="Marghera", sigla="VE", codice_bc=3)
        self.milano = Stabilimento.objects.create(nome="Milano")  # senza codice_bc

        testata = Testata.objects.create(job="99001")
        self.doc1 = Documento.objects.create(testata=testata, vendor_doc="99001-DOC1")
        self.doc2 = Documento.objects.create(testata=testata, vendor_doc="99001-DOC2")

    def test_unicita_documento_stabilimento(self):
        DestinazioneDocumento.objects.create(documento=self.doc1, stabilimento=self.valbrembo)

        with self.assertRaises(IntegrityError), transaction.atomic():
            DestinazioneDocumento.objects.create(documento=self.doc1, stabilimento=self.valbrembo)

    def test_clean_rifiuta_uno_stabilimento_senza_codice_bc(self):
        destinazione = DestinazioneDocumento(documento=self.doc1, stabilimento=self.milano)

        with self.assertRaises(ValidationError):
            destinazione.full_clean()

    def test_imposta_destinazioni_sostituisce_e_non_accumula(self):
        imposta_destinazioni(self.doc1, [1, 2])
        self.assertEqual([s.codice_bc for s in siti_del_documento(self.doc1)], [1, 2])

        imposta_destinazioni(self.doc1, [3])

        self.assertEqual([s.codice_bc for s in siti_del_documento(self.doc1)], [3])

    def test_imposta_destinazioni_lista_vuota_azzera(self):
        imposta_destinazioni(self.doc1, [1, 2])

        imposta_destinazioni(self.doc1, [])

        self.assertEqual(siti_del_documento(self.doc1), [])

    def test_siti_del_documento_ordinati_per_codice_bc(self):
        imposta_destinazioni(self.doc1, [3, 1, 2])

        self.assertEqual([s.codice_bc for s in siti_del_documento(self.doc1)], [1, 2, 3])

    def test_siti_coinvolti_deduplica_su_documenti_con_siti_sovrapposti(self):
        imposta_destinazioni(self.doc1, [1, 2])
        imposta_destinazioni(self.doc2, [2, 3])

        risultato = siti_coinvolti([self.doc1, self.doc2])

        self.assertEqual([s.codice_bc for s in risultato], [1, 2, 3])


class TrasmittalInternoArchivioTests(TestCase):
    """Archivio del trasmittal interno: progressivo, nome, snapshot dei siti, destinatari."""

    def setUp(self):
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.pd = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        self.ve = Stabilimento.objects.create(nome="Marghera", sigla="VE", codice_bc=3)

        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.TO
        )
        IndirizzoStabilimento.objects.create(
            stabilimento=self.pd, email="pd@b.it", tipo=TipoIndirizzoStabilimento.TO
        )
        # Stesso indirizzo di BG, ma CC su un altro sito: nei destinatari deve
        # restare solo in TO, senza comparire due volte.
        IndirizzoStabilimento.objects.create(
            stabilimento=self.ve, email="bg@b.it", tipo=TipoIndirizzoStabilimento.CC
        )

        self.testata = Testata.objects.create(job="99010")
        self.altra_testata = Testata.objects.create(job="99011")
        self.doc1 = Documento.objects.create(testata=self.testata, vendor_doc="99010-DOC1")
        self.doc2 = Documento.objects.create(testata=self.testata, vendor_doc="99010-DOC2")

        imposta_destinazioni(self.doc1, [1, 2])
        imposta_destinazioni(self.doc2, [3])

        self.utente = User.objects.create_user("trasmittal_utente", password="pw")

    def _righe(self, *documenti):
        return [{"documento": doc, "revisione": "1"} for doc in documenti]

    # -- prossimo_progressivo --

    def test_prossimo_progressivo_riparte_da_1_il_giorno_dopo(self):
        crea_trasmittal_interno(
            self.testata, self._righe(self.doc1), self.utente, data=date(2026, 9, 14)
        )

        self.assertEqual(prossimo_progressivo(self.testata, date(2026, 9, 14)), 2)
        self.assertEqual(prossimo_progressivo(self.testata, date(2026, 9, 15)), 1)

    def test_prossimo_progressivo_non_collide_fra_commesse(self):
        crea_trasmittal_interno(
            self.testata, self._righe(self.doc1), self.utente, data=date(2026, 9, 14)
        )

        self.assertEqual(prossimo_progressivo(self.altra_testata, date(2026, 9, 14)), 1)

    # -- componi_nome --

    def test_componi_nome_formato_atteso(self):
        self.assertEqual(componi_nome(self.testata, date(2026, 9, 14), 3), "99010_2026-09-14_E3")

    # -- snapshot dei siti --

    def test_i_siti_della_riga_sono_uno_snapshot(self):
        trasmittal = crea_trasmittal_interno(
            self.testata, self._righe(self.doc1), self.utente, data=date(2026, 9, 14)
        )
        riga = trasmittal.righe.get(documento=self.doc1)
        self.assertEqual(sorted(s.codice_bc for s in riga.siti.all()), [1, 2])

        imposta_destinazioni(self.doc1, [3])

        self.assertEqual(sorted(s.codice_bc for s in riga.siti.all()), [1, 2])

    # -- destinatari di stabilimento --

    def test_destinatari_di_stabilimento_dedotti_dall_unione_dei_siti_senza_duplicati(self):
        trasmittal = crea_trasmittal_interno(
            self.testata, self._righe(self.doc1, self.doc2), self.utente, data=date(2026, 9, 14)
        )

        destinatari = list(trasmittal.destinatari.values_list("email", "tipo", "origine"))

        self.assertEqual(
            sorted(destinatari),
            sorted(
                [
                    ("bg@b.it", "to", "stabilimento"),
                    ("pd@b.it", "to", "stabilimento"),
                ]
            ),
        )

    # -- unicità documento nella lettera --

    def test_un_documento_non_puo_comparire_due_volte_nella_stessa_lettera(self):
        trasmittal = TransmittalInterno.objects.create(
            testata=self.testata,
            data=date(2026, 9, 14),
            progressivo=1,
            nome=componi_nome(self.testata, date(2026, 9, 14), 1),
            creato_da=self.utente,
        )
        RigaTransmittalInterno.objects.create(
            trasmittal=trasmittal, documento=self.doc1, revisione="1", posizione=1
        )

        with self.assertRaises(IntegrityError), transaction.atomic():
            RigaTransmittalInterno.objects.create(
                trasmittal=trasmittal, documento=self.doc1, revisione="2", posizione=2
            )


class RigaTransmittalInternoTpiTests(TestCase):
    """TPI (b)(c) del modulo: SÌ/NO all'ispettore terzo, e se sì chi.

    ``tpi`` e ``tpi_destinatario`` non possono essere in disaccordo: un
    destinatario ha senso solo se ``tpi`` è attivo, e se è attivo va
    specificato chi.
    """

    def setUp(self):
        self.testata = Testata.objects.create(job="99012")
        self.doc = Documento.objects.create(testata=self.testata, vendor_doc="99012-DOC1")
        self.utente = User.objects.create_user("riga_tpi_utente", password="pw")
        self.trasmittal = TransmittalInterno.objects.create(
            testata=self.testata,
            data=date(2026, 9, 14),
            progressivo=1,
            nome=componi_nome(self.testata, date(2026, 9, 14), 1),
            creato_da=self.utente,
        )

    def _riga(self, **overrides):
        base = {
            "trasmittal": self.trasmittal,
            "documento": self.doc,
            "revisione": "1",
            "posizione": 1,
        }
        base.update(overrides)
        return RigaTransmittalInterno(**base)

    def test_tpi_spento_con_destinatario_e_un_errore_di_validazione(self):
        riga = self._riga(tpi=False, tpi_destinatario="AI")

        with self.assertRaises(ValidationError):
            riga.full_clean()

    def test_tpi_acceso_senza_destinatario_e_un_errore_di_validazione(self):
        riga = self._riga(tpi=True, tpi_destinatario="")

        with self.assertRaises(ValidationError):
            riga.full_clean()

    def test_tpi_spento_senza_destinatario_e_valido(self):
        riga = self._riga(tpi=False, tpi_destinatario="")

        riga.full_clean()  # non deve sollevare

    def test_tpi_acceso_con_destinatario_e_valido(self):
        riga = self._riga(tpi=True, tpi_destinatario="No.Bo.")

        riga.full_clean()  # non deve sollevare


class DestinatariTrasmittalInternoRuoliTests(TestCase):
    """Risoluzione dei destinatari da PM/PE/QCI e dalla regola export@ per documenti SHn."""

    def setUp(self):
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.TO
        )

        self.testata = Testata.objects.create(job="99012")
        self.doc = Documento.objects.create(testata=self.testata, vendor_doc="99012-01-QCPA")
        imposta_destinazioni(self.doc, [1])

        self.utente = User.objects.create_user("destinatari_utente", password="pw")
        self.pm = User.objects.create_user("destinatari_pm", "pm@b.it", "pw")
        self.pe = User.objects.create_user("destinatari_pe", "pe@b.it", "pw")
        self.qci = User.objects.create_user("destinatari_qci", "qci@b.it", "pw")

    def _righe(self, *documenti):
        return [{"documento": doc, "revisione": "1"} for doc in documenti]

    def test_pm_in_to_pe_e_qci_in_cc(self):
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, utente=self.pm
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PE, utente=self.pe
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.QCI, utente=self.qci
        )

        trasmittal = crea_trasmittal_interno(self.testata, self._righe(self.doc), self.utente)

        destinatari = set(trasmittal.destinatari.values_list("email", "tipo", "origine"))
        self.assertEqual(
            destinatari,
            {
                ("bg@b.it", "to", "stabilimento"),
                ("pm@b.it", "to", "pm"),
                ("pe@b.it", "cc", "pe"),
                ("qci@b.it", "cc", "qci"),
            },
        )

    def test_ruolo_non_valorizzato_non_e_un_errore(self):
        trasmittal = crea_trasmittal_interno(self.testata, self._righe(self.doc), self.utente)

        self.assertEqual(
            set(trasmittal.destinatari.values_list("email", "tipo", "origine")),
            {("bg@b.it", "to", "stabilimento")},
        )

    def test_documento_shn_aggiunge_export_in_cc(self):
        doc_shn = Documento.objects.create(testata=self.testata, vendor_doc="99012-01-ESH1")
        imposta_destinazioni(doc_shn, [1])

        trasmittal = crea_trasmittal_interno(
            self.testata, self._righe(self.doc, doc_shn), self.utente
        )

        self.assertIn(
            ("export@brembanarolle.com", "cc", "export"),
            set(trasmittal.destinatari.values_list("email", "tipo", "origine")),
        )

    def test_documento_non_shn_non_aggiunge_export(self):
        trasmittal = crea_trasmittal_interno(self.testata, self._righe(self.doc), self.utente)

        self.assertNotIn(
            "export@brembanarolle.com", trasmittal.destinatari.values_list("email", flat=True)
        )

    def test_pm_gia_indirizzo_di_stabilimento_una_sola_occorrenza_in_to(self):
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, utente=self.pm
        )
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="pm@b.it", tipo=TipoIndirizzoStabilimento.CC
        )

        trasmittal = crea_trasmittal_interno(self.testata, self._righe(self.doc), self.utente)

        destinatari = list(trasmittal.destinatari.values_list("email", "tipo"))
        self.assertEqual(destinatari.count(("pm@b.it", "to")), 1)
        self.assertNotIn(("pm@b.it", "cc"), destinatari)

    def test_stesso_indirizzo_in_to_e_cc_resta_in_to(self):
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.CC
        )

        trasmittal = crea_trasmittal_interno(self.testata, self._righe(self.doc), self.utente)

        destinatari = list(trasmittal.destinatari.values_list("email", "tipo"))
        self.assertEqual(destinatari.count(("bg@b.it", "to")), 1)
        self.assertNotIn(("bg@b.it", "cc"), destinatari)


class DataImpegnoTests(SimpleTestCase):
    """data_impegno: termine per completare la distribuzione, non la data di firma."""

    def test_lunedi_martedi(self):
        self.assertEqual(data_impegno(date(2026, 9, 14)), date(2026, 9, 15))

    def test_giovedi_venerdi(self):
        self.assertEqual(data_impegno(date(2026, 9, 17)), date(2026, 9, 18))

    def test_venerdi_lunedi_tre_giorni_dopo(self):
        self.assertEqual(data_impegno(date(2026, 9, 18)), date(2026, 9, 21))

    def test_sabato_lunedi(self):
        self.assertEqual(data_impegno(date(2026, 9, 19)), date(2026, 9, 21))

    def test_domenica_lunedi(self):
        self.assertEqual(data_impegno(date(2026, 9, 20)), date(2026, 9, 21))


class TrasmittalInternoPdfTests(TestCase):
    """Generazione del PDF del trasmittal interno (form MQ 7.5-04)."""

    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        impostazioni = override_settings(MEDIA_ROOT=media.name)
        impostazioni.enable()
        self.addCleanup(impostazioni.disable)

        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.pd = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)

        self.produzione = User.objects.create_user(
            "pdf_produzione", "pdf-produzione@b.it", "pw", first_name="Mario", last_name="Rossi"
        )
        self.produzione.firma = SimpleUploadedFile("firma.png", _png(), "image/png")
        self.produzione.save()
        # Lo stesso utente firma per la Qualità di entrambi i siti: nel PDF
        # deve comparire una sola volta.
        self.qualita = User.objects.create_user(
            "pdf_qualita", "pdf-qualita@b.it", "pw", first_name="Anna", last_name="Bianchi"
        )
        FirmatarioStabilimento.objects.create(
            stabilimento=self.bg,
            ruolo=RuoloFirmatarioStabilimento.PRODUZIONE,
            utente=self.produzione,
        )
        FirmatarioStabilimento.objects.create(
            stabilimento=self.bg, ruolo=RuoloFirmatarioStabilimento.QUALITA, utente=self.qualita
        )
        FirmatarioStabilimento.objects.create(
            stabilimento=self.pd, ruolo=RuoloFirmatarioStabilimento.QUALITA, utente=self.qualita
        )

        self.testata = Testata.objects.create(job="99020", job_detail="Prova PDF interno")
        self.doc1 = Documento.objects.create(testata=self.testata, vendor_doc="99020-DOC1")
        self.doc2 = Documento.objects.create(testata=self.testata, vendor_doc="99020-DOC2")
        imposta_destinazioni(self.doc1, [1, 2])
        imposta_destinazioni(self.doc2, [1])

        self.trasmittal = crea_trasmittal_interno(
            self.testata,
            [
                {
                    "documento": self.doc1,
                    "revisione": "3",
                    "copie": 2,
                    "tpi": True,
                    "tpi_destinatario": "ABC",
                    "note": "x",
                },
                {"documento": self.doc2, "revisione": "B", "cliente": False},
            ],
            self.produzione,
            data=date(2026, 9, 14),
            note="Note libere della lettera.",
        )

    def test_il_pdf_si_genera_e_non_e_vuoto(self):
        from src.pdf import genera_trasmittal_interno_pdf

        contenuto = genera_trasmittal_interno_pdf(self.trasmittal)

        self.assertTrue(contenuto.startswith(b"%PDF"))
        self.assertGreater(len(contenuto), 1000)

    def test_firme_deduplicate_per_utente_condiviso_tra_stabilimenti(self):
        from src.pdf import _firmatari_per_ruolo

        firmatari_qualita = _firmatari_per_ruolo(
            [self.bg, self.pd], RuoloFirmatarioStabilimento.QUALITA
        )

        self.assertEqual([f.utente_id for f in firmatari_qualita], [self.qualita.pk])

    # -- colonna TPI: "NO" se spento, il destinatario se acceso --

    def test_colonna_tpi_mostra_no_se_spento(self):
        from src.pdf import _int_row_value

        riga = self.trasmittal.righe.get(documento=self.doc2)  # tpi=False di default

        self.assertEqual(_int_row_value(riga, "tpi"), "NO")

    def test_colonna_tpi_mostra_il_destinatario_se_acceso(self):
        from src.pdf import _int_row_value

        riga = self.trasmittal.righe.get(documento=self.doc1)  # tpi=True, dest="ABC"

        self.assertEqual(_int_row_value(riga, "tpi"), "ABC")

    def test_tabella_email_mostra_no_e_destinatario_come_il_pdf(self):
        from core.services.trasmittal_interno import _tabella_testo_righe

        tabella = _tabella_testo_righe(self.trasmittal)

        self.assertIn("NO", tabella)
        self.assertIn("ABC", tabella)

    def test_firmatario_senza_immagine_non_solleva_eccezioni(self):
        from src.pdf import genera_trasmittal_interno_pdf

        # self.qualita non ha un'immagine di firma caricata.
        contenuto = genera_trasmittal_interno_pdf(self.trasmittal)

        self.assertTrue(contenuto.startswith(b"%PDF"))


class SalvaPdfTrasmittalInternoTests(TestCase):
    """salva_pdf: scrittura sul fileserver, sempre dentro la cartella JOBS."""

    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        impostazioni = override_settings(MEDIA_ROOT=media.name)
        impostazioni.enable()
        self.addCleanup(impostazioni.disable)

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        jobs_patcher = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        jobs_patcher.enable()
        self.addCleanup(jobs_patcher.disable)

        self.utente = User.objects.create_user("salva_pdf_utente", "salva-pdf@b.it", "pw")

    def _trasmittal(self, job):
        testata = Testata.objects.create(job=job)
        return TransmittalInterno.objects.create(
            testata=testata,
            data=date(2026, 9, 14),
            progressivo=1,
            nome=componi_nome(testata, date(2026, 9, 14), 1),
            creato_da=self.utente,
        )

    def test_salva_pdf_rifiuta_un_percorso_fuori_dalla_jobs_root(self):
        trasmittal = self._trasmittal("../fuori")

        with self.assertRaises(PermissionError):
            salva_pdf(trasmittal)


class PreparaPerDccTests(TestCase):
    """prepara_per_dcc: copia i PDF dei documenti cliente in una cartella dedicata."""

    def setUp(self):
        from .services.fileserver import get_base_path

        self.get_base_path = get_base_path

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        patcher = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        patcher.enable()
        self.addCleanup(patcher.disable)

        self.reparto = Reparto.objects.create(nome="Qualità e Controllo", acronimo="QMD")
        self.testata = Testata.objects.create(job="99030")
        self.utente = User.objects.create_user("dcc_utente", "dcc@b.it", "pw")

        self.doc_ok = Documento.objects.create(
            testata=self.testata, vendor_doc="99030-QMDBI", reparto=self.reparto.nome
        )
        self.doc_non_cliente = Documento.objects.create(
            testata=self.testata, vendor_doc="99030-QCPA", reparto=self.reparto.nome
        )
        self.doc_missing = Documento.objects.create(
            testata=self.testata, vendor_doc="99030-DWG01", reparto=self.reparto.nome
        )

        base = get_base_path("99030", "QMD")
        base.mkdir(parents=True)
        (base / f"{self.doc_ok.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-doc-ok")
        (base / f"{self.doc_non_cliente.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-doc-non-cliente")
        # doc_missing: nessun file sul finto fileserver.

    def _trasmittal(self, righe, testata=None):
        return crea_trasmittal_interno(
            testata or self.testata, righe, self.utente, data=date(2026, 9, 14)
        )

    def test_copia_solo_i_documenti_con_cliente_true(self):
        trasmittal = self._trasmittal(
            [
                {"documento": self.doc_ok, "revisione": "1", "cliente": True},
                {"documento": self.doc_non_cliente, "revisione": "1", "cliente": False},
            ]
        )

        esito = prepara_per_dcc(trasmittal)

        self.assertEqual([c["vendor_doc"] for c in esito["copiati"]], [self.doc_ok.vendor_doc])
        self.assertEqual(esito["mancanti"], [])
        cartella = Path(esito["cartella"])
        self.assertEqual(
            [p.name for p in cartella.iterdir()], [f"{self.doc_ok.vendor_doc} Rev A.pdf"]
        )

    def test_nessuna_riga_cliente_nessuna_cartella_creata(self):
        trasmittal = self._trasmittal(
            [{"documento": self.doc_non_cliente, "revisione": "1", "cliente": False}]
        )

        esito = prepara_per_dcc(trasmittal)

        self.assertEqual(esito, {"cartella": None, "copiati": [], "mancanti": []})
        cartella_attesa = self.get_base_path("99030", "DCC") / "DA SPEDIRE" / trasmittal.nome
        self.assertFalse(cartella_attesa.exists())

    def test_documento_mancante_gli_altri_vengono_copiati(self):
        trasmittal = self._trasmittal(
            [
                {"documento": self.doc_ok, "revisione": "1", "cliente": True},
                {"documento": self.doc_missing, "revisione": "1", "cliente": True},
            ]
        )

        esito = prepara_per_dcc(trasmittal)

        self.assertEqual([c["vendor_doc"] for c in esito["copiati"]], [self.doc_ok.vendor_doc])
        self.assertEqual(
            [m["vendor_doc"] for m in esito["mancanti"]], [self.doc_missing.vendor_doc]
        )

    def test_seconda_esecuzione_idempotente(self):
        trasmittal = self._trasmittal(
            [{"documento": self.doc_ok, "revisione": "1", "cliente": True}]
        )

        esito1 = prepara_per_dcc(trasmittal)
        esito2 = prepara_per_dcc(trasmittal)

        self.assertEqual(esito1, esito2)
        cartella = Path(esito1["cartella"])
        self.assertEqual(len(list(cartella.iterdir())), 1)

    def test_percorso_fuori_dalla_jobs_root_rifiutato(self):
        testata_fuori = Testata.objects.create(job="../fuori")
        doc = Documento.objects.create(
            testata=testata_fuori, vendor_doc="X", reparto=self.reparto.nome
        )
        trasmittal = self._trasmittal(
            [{"documento": doc, "revisione": "1", "cliente": True}], testata=testata_fuori
        )

        with self.assertRaises(PermissionError):
            prepara_per_dcc(trasmittal)

    def test_non_interferisce_con_trasmittal_archivio(self):
        # La cartella "DA SPEDIRE/<nome lettera>" è sorella di "DA SPEDIRE/TRANSMITTAL",
        # non ci finisce dentro: trasmittal_archivio non deve accorgersene.
        from .services.trasmittal_archivio import cartella_trasmittal, lista_trasmittal

        trasmittal = self._trasmittal(
            [{"documento": self.doc_ok, "revisione": "1", "cliente": True}]
        )

        prepara_per_dcc(trasmittal)

        self.assertEqual(lista_trasmittal("99030"), [])
        self.assertEqual(
            cartella_trasmittal("99030"), self.get_base_path("99030", "DCC") / "TRANSMITTAL"
        )


class TrasmittalInternoDestinazioniViewTests(TestCase):
    """Pagina e API della griglia destinazioni cartacee del trasmittal interno."""

    def setUp(self):
        self.client = Client()
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.pd = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        Stabilimento.objects.create(nome="Milano")  # non è un sito costruttivo

        self.reparto_ut = Reparto.objects.create(nome="Ufficio Tecnico", acronimo="UT")
        self.reparto_qc = Reparto.objects.create(nome="Qualità e Controllo", acronimo="QMD")

        self.testata = Testata.objects.create(job="99040", sito_costruttivo=self.bg)
        self.doc_alfa = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99040-ALFA",
            doc_title="Documento Alfa",
            reparto=self.reparto_ut.nome,
        )
        self.doc_beta = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99040-BETA",
            doc_title="Documento Beta",
            reparto=self.reparto_ut.nome,
        )
        self.doc_non_ut = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99040-QC",
            reparto=self.reparto_qc.nome,
        )

        self.writer = User.objects.create_user(
            "tidv_writer", "tidv-writer@b.it", "pw", permesso=Permesso.WRITING
        )
        self.reader = User.objects.create_user(
            "tidv_reader", "tidv-reader@b.it", "pw", permesso=Permesso.READING
        )

    def _pagina_url(self, job=None):
        return f"/commesse/{job or self.testata.job}/trasmittal-interno/"

    def _api_destinazioni_url(self, job=None):
        return f"/api/commesse/{job or self.testata.job}/trasmittal-interno/destinazioni/"

    def _api_documento_url(self, documento_id, job=None):
        return (
            f"/api/commesse/{job or self.testata.job}/trasmittal-interno/"
            f"documenti/{documento_id}/destinazioni/"
        )

    def _api_bulk_url(self, codice_bc, job=None):
        return (
            f"/api/commesse/{job or self.testata.job}/trasmittal-interno/"
            f"stabilimenti/{codice_bc}/destinazioni/"
        )

    # -- pagina --

    def test_pagina_richiede_login(self):
        risposta = self.client.get(self._pagina_url())
        self.assertNotEqual(risposta.status_code, 200)

    def test_pagina_risponde_e_mostra_i_documenti_della_commessa(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(self._pagina_url())

        self.assertEqual(risposta.status_code, 200)
        self.assertContains(risposta, "Trasmittal interno")

    # -- prerequisito: sito costruttivo --

    def test_commessa_senza_sito_costruttivo_accesso_diretto_rifiutato(self):
        Testata.objects.create(job="99041")  # niente sito_costruttivo
        self.client.force_login(self.reader)

        risposta = self.client.get(self._pagina_url(job="99041"))

        self.assertEqual(risposta.status_code, 403)

    def test_card_bloccata_in_pagina_commessa_senza_sito_costruttivo(self):
        Testata.objects.create(job="99042")
        self.client.force_login(self.reader)

        risposta = self.client.get("/commesse/99042/")

        # Il nome della classe compare anche nel <style> (regole CSS), quindi
        # va cercato sull'attributo class della card, non come sottostringa
        # libera nell'intera risposta.
        self.assertContains(risposta, 'class="section-card section-card-locked"')
        self.assertContains(risposta, "Completa prima il sito costruttivo della commessa")

    def test_card_non_bloccata_in_pagina_commessa_con_sito_costruttivo(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(f"/commesse/{self.testata.job}/")

        self.assertNotContains(risposta, 'class="section-card section-card-locked"')

    # -- API: lettura griglia --

    def test_api_destinazioni_richiede_login(self):
        risposta = self.client.get(self._api_destinazioni_url())
        self.assertNotEqual(risposta.status_code, 200)

    def test_api_destinazioni_elenca_solo_i_documenti_ut(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(self._api_destinazioni_url())

        self.assertEqual(risposta.status_code, 200)
        dati = risposta.json()
        self.assertEqual(
            {d["vendor_doc"] for d in dati["documenti"]},
            {self.doc_alfa.vendor_doc, self.doc_beta.vendor_doc},
        )
        self.assertEqual(
            [s["sigla"] for s in dati["stabilimenti"]], ["BG", "PD"]
        )  # niente Milano (senza codice_bc)

    # -- API: salvataggio per documento --

    def test_salvataggio_persiste_e_la_rilettura_mostra_le_destinazioni_spuntate(self):
        self.client.force_login(self.writer)

        risposta = self.client.post(
            self._api_documento_url(self.doc_alfa.pk),
            data=json.dumps({"codici_bc": [1, 2]}),
            content_type="application/json",
        )
        self.assertEqual(risposta.status_code, 200)

        rilettura = self.client.get(self._api_destinazioni_url()).json()
        alfa = next(d for d in rilettura["documenti"] if d["id"] == self.doc_alfa.pk)
        self.assertEqual(sorted(alfa["codici_bc"]), [1, 2])
        beta = next(d for d in rilettura["documenti"] if d["id"] == self.doc_beta.pk)
        self.assertEqual(beta["codici_bc"], [])

    def test_salvataggio_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._api_documento_url(self.doc_alfa.pk),
            data=json.dumps({"codici_bc": [1]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 403)
        self.assertEqual(list(self.doc_alfa.destinazioni.all()), [])

    def test_salvataggio_su_documento_non_ut_rifiutato(self):
        self.client.force_login(self.writer)

        risposta = self.client.post(
            self._api_documento_url(self.doc_non_ut.pk),
            data=json.dumps({"codici_bc": [1]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 404)

    # -- API: selezione in blocco --

    def test_selezione_in_blocco_agisce_solo_sui_documenti_filtrati(self):
        self.client.force_login(self.writer)
        imposta_destinazioni(self.doc_beta, [2])  # beta è già su PD, fuori dal filtro

        risposta = self.client.post(
            self._api_bulk_url(codice_bc=1),
            data=json.dumps({"documento_ids": [self.doc_alfa.pk], "attiva": True}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(
            sorted(self.doc_alfa.destinazioni.values_list("stabilimento__codice_bc", flat=True)),
            [1],
        )
        # doc_beta non era nel filtro (documento_ids): resta invariato.
        self.assertEqual(
            list(self.doc_beta.destinazioni.values_list("stabilimento__codice_bc", flat=True)),
            [2],
        )

    def test_selezione_in_blocco_puo_disattivare(self):
        self.client.force_login(self.writer)
        imposta_destinazioni(self.doc_alfa, [1, 2])
        imposta_destinazioni(self.doc_beta, [1])

        risposta = self.client.post(
            self._api_bulk_url(codice_bc=1),
            data=json.dumps(
                {"documento_ids": [self.doc_alfa.pk, self.doc_beta.pk], "attiva": False}
            ),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(
            sorted(self.doc_alfa.destinazioni.values_list("stabilimento__codice_bc", flat=True)),
            [2],
        )
        self.assertEqual(
            list(self.doc_beta.destinazioni.values_list("stabilimento__codice_bc", flat=True)), []
        )

    def test_selezione_in_blocco_ignora_id_estranei_a_questa_commessa(self):
        altra_testata = Testata.objects.create(job="99043", sito_costruttivo=self.bg)
        doc_altra = Documento.objects.create(
            testata=altra_testata,
            vendor_doc="99043-ALFA",
            reparto=self.reparto_ut.nome,
        )
        self.client.force_login(self.writer)

        risposta = self.client.post(
            self._api_bulk_url(codice_bc=1),
            data=json.dumps({"documento_ids": [self.doc_alfa.pk, doc_altra.pk], "attiva": True}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(list(doc_altra.destinazioni.all()), [])
        self.assertEqual(
            list(self.doc_alfa.destinazioni.values_list("stabilimento__codice_bc", flat=True)),
            [1],
        )

    def test_selezione_in_blocco_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._api_bulk_url(codice_bc=1),
            data=json.dumps({"documento_ids": [self.doc_alfa.pk], "attiva": True}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 403)
        self.assertEqual(list(self.doc_alfa.destinazioni.all()), [])


# ── Integration tests (real fileserver Z:\JOBS) ───────────────────────────────


def _make_integration_fixtures():
    """Create DB fixtures pointing to the real fileserver structure.

    Both documents are under Z:\\JOBS\\25056\\PROGETTO\\QC\\ as observed on the
    real fileserver. This creates fresh records in the test DB.

    Returns:
        Dict with Revisione instances keyed by name.
    """
    import datetime

    reparto, _ = Reparto.objects.get_or_create(
        acronimo="QC",
        defaults={"nome": "Qualità e Controllo"},
    )
    testata, _ = Testata.objects.get_or_create(
        job="25056",
        defaults={"client": "Test", "job_detail": "Integration test"},
    )
    doc_qmdbi, _ = Documento.objects.get_or_create(
        testata=testata,
        vendor_doc="25056-QMDBI",
        defaults={"reparto": reparto.nome, "doc_title": "QMDBI"},
    )
    doc_qcpa, _ = Documento.objects.get_or_create(
        testata=testata,
        vendor_doc="25056-01-QCPA",
        defaults={"reparto": reparto.nome, "doc_title": "QCPA"},
    )

    rev_qmdbi_0, _ = Revisione.objects.get_or_create(
        documento=doc_qmdbi,
        rev_no=0,
        defaults={
            "dis_act_date": datetime.date(2025, 9, 5),
            "rec_act_date": datetime.date(2025, 11, 17),
        },
    )
    rev_qcpa_0, _ = Revisione.objects.get_or_create(
        documento=doc_qcpa,
        rev_no=0,
        defaults={
            "dis_act_date": datetime.date(2025, 9, 5),
            "rec_act_date": datetime.date(2025, 10, 14),
        },
    )
    rev_qcpa_1, _ = Revisione.objects.get_or_create(
        documento=doc_qcpa,
        rev_no=1,
        defaults={
            "dis_act_date": datetime.date(2025, 11, 7),
            "rec_act_date": datetime.date(2025, 11, 10),
        },
    )
    rev_qcpa_2, _ = Revisione.objects.get_or_create(
        documento=doc_qcpa,
        rev_no=2,
        defaults={
            "dis_act_date": datetime.date(2025, 11, 13),
            "rec_act_date": datetime.date(2025, 11, 16),
        },
    )

    return {
        "rev_qmdbi_0": rev_qmdbi_0,
        "rev_qcpa_0": rev_qcpa_0,
        "rev_qcpa_1": rev_qcpa_1,
        "rev_qcpa_2": rev_qcpa_2,
    }


class TrasmittalInternoLetteraTests(TestCase):
    """Creazione della lettera: selezione, anteprima, conferma, elenco emesse."""

    def setUp(self):
        self.client = Client()

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        patcher = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        patcher.enable()
        self.addCleanup(patcher.disable)

        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.TO
        )
        self.reparto_ut = Reparto.objects.create(nome="Ufficio Tecnico", acronimo="UT")
        self.testata = Testata.objects.create(job="99070", sito_costruttivo=self.bg)

        self.doc_ok = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99070-ALFA",
            doc_title="Documento Alfa",
            reparto=self.reparto_ut.nome,
        )
        Revisione.objects.create(documento=self.doc_ok, rev_no=0)
        imposta_destinazioni(self.doc_ok, [self.bg.codice_bc])
        from core.services.fileserver import get_base_path

        base = get_base_path(self.testata.job, "UT")
        base.mkdir(parents=True)
        (base / f"{self.doc_ok.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-fake")

        self.doc_senza_revisione = Documento.objects.create(
            testata=self.testata, vendor_doc="99070-BETA", reparto=self.reparto_ut.nome
        )
        self.doc_senza_file = Documento.objects.create(
            testata=self.testata, vendor_doc="99070-GAMMA", reparto=self.reparto_ut.nome
        )
        Revisione.objects.create(documento=self.doc_senza_file, rev_no=0)

        self.pm_utente = User.objects.create_user(
            "til_pm", "til-pm@b.it", "pw", first_name="Mario", last_name="PM"
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, utente=self.pm_utente
        )

        self.writer = User.objects.create_user(
            "til_writer", "til-writer@b.it", "pw", permesso=Permesso.WRITING
        )
        self.reader = User.objects.create_user(
            "til_reader", "til-reader@b.it", "pw", permesso=Permesso.READING
        )

    def _url(self, path):
        return f"/api/commesse/{self.testata.job}/trasmittal-interno/{path}"

    def _riga(self, documento, **overrides):
        base = {
            "documento_id": documento.pk,
            "revisione": "A",
            "copie": 2,
            "tpi": False,
            "note": "",
            "cliente": True,
        }
        base.update(overrides)
        return base

    # -- selezione --

    def test_documento_senza_revisione_non_selezionabile(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(self._url("selezione/"))

        voce = next(
            d for d in risposta.json()["documenti"] if d["id"] == self.doc_senza_revisione.pk
        )
        self.assertFalse(voce["selezionabile"])
        self.assertEqual(voce["motivo"], "Nessuna revisione registrata.")

    def test_documento_senza_file_non_selezionabile(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(self._url("selezione/"))

        voce = next(d for d in risposta.json()["documenti"] if d["id"] == self.doc_senza_file.pk)
        self.assertFalse(voce["selezionabile"])
        self.assertEqual(voce["motivo"], "Nessun file trovato sul fileserver.")

    def test_documento_valido_selezionabile(self):
        self.client.force_login(self.reader)

        risposta = self.client.get(self._url("selezione/"))

        voce = next(d for d in risposta.json()["documenti"] if d["id"] == self.doc_ok.pk)
        self.assertTrue(voce["selezionabile"])
        self.assertEqual(voce["motivo"], "")

    # -- anteprima --

    def test_anteprima_mostra_destinatari_risolti_con_origine(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        destinatari = risposta.json()["destinatari"]
        self.assertIn({"email": "bg@b.it", "tipo": "to", "origine": "stabilimento"}, destinatari)
        self.assertIn({"email": "til-pm@b.it", "tipo": "to", "origine": "pm"}, destinatari)
        self.assertEqual(sorted(risposta.json()["ruoli_mancanti"]), ["PE", "QCI"])

    def test_anteprima_include_nome_e_percorso_pdf_previsti(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok, cliente=False)]}),
            content_type="application/json",
        )

        corpo = risposta.json()
        self.assertTrue(corpo["nome_preview"].startswith(f"{self.testata.job}_"))
        self.assertTrue(corpo["percorso_pdf_preview"].endswith(f"{corpo['nome_preview']}.pdf"))
        self.assertIsNone(corpo["percorso_dcc_preview"])

    def test_anteprima_percorso_dcc_previsto_solo_con_un_documento_cliente(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok, cliente=True)]}),
            content_type="application/json",
        )

        corpo = risposta.json()
        self.assertIsNotNone(corpo["percorso_dcc_preview"])
        self.assertIn(corpo["nome_preview"], corpo["percorso_dcc_preview"])

    def test_anteprima_rifiuta_documento_non_selezionabile(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_senza_revisione)]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 400)
        self.assertIn("revisione", risposta.json()["error"].lower())

    # -- conferma --

    def test_conferma_crea_trasmittal_salva_pdf_popola_dcc_e_invia_email(self):
        self.client.force_login(self.writer)

        risposta = self.client.post(
            self._url("emetti/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok)], "note": "riga 1\nriga 2"}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        corpo = risposta.json()
        self.assertTrue(corpo["ok"])
        self.assertTrue(corpo["pdf"]["ok"])
        self.assertTrue(Path(corpo["pdf"]["percorso"]).is_file())
        self.assertTrue(corpo["dcc"]["ok"])
        self.assertEqual(len(corpo["dcc"]["copiati"]), 1)
        self.assertTrue(corpo["email"]["ok"])

        self.assertEqual(TransmittalInterno.objects.filter(testata=self.testata).count(), 1)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["bg@b.it", "til-pm@b.it"])
        self.assertEqual(len(mail.outbox[0].attachments), 1)

    def test_destinatari_modificati_in_anteprima_sono_quelli_usati_e_registrati(self):
        self.client.force_login(self.writer)
        destinatari_modificati = [
            {"email": "extra@b.it", "tipo": "to", "origine": "manuale"},
        ]

        risposta = self.client.post(
            self._url("emetti/"),
            data=json.dumps(
                {"righe": [self._riga(self.doc_ok)], "destinatari": destinatari_modificati}
            ),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        trasmittal = TransmittalInterno.objects.get(pk=risposta.json()["trasmittal_id"])
        registrati = list(trasmittal.destinatari.values_list("email", "tipo", "origine"))
        self.assertEqual(registrati, [("extra@b.it", "to", "manuale")])
        self.assertEqual(mail.outbox[0].to, ["extra@b.it"])
        # I destinatari auto-risolti (stabilimento, PM) non compaiono più:
        # quelli confermati in anteprima sono gli unici usati.
        self.assertNotIn("bg@b.it", mail.outbox[0].to)

    def test_conferma_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._url("emetti/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 403)
        self.assertEqual(TransmittalInterno.objects.count(), 0)

    def test_invio_email_fallito_la_lettera_resta_in_archivio_e_l_errore_e_visibile(self):
        self.client.force_login(self.writer)

        with patch(
            "django.core.mail.EmailMessage.send", side_effect=Exception("SMTP non raggiungibile")
        ):
            risposta = self.client.post(
                self._url("emetti/"),
                data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
                content_type="application/json",
            )

        self.assertEqual(risposta.status_code, 200)
        corpo = risposta.json()
        self.assertTrue(corpo["ok"])
        self.assertFalse(corpo["email"]["ok"])
        self.assertIn("SMTP non raggiungibile", corpo["email"]["errore"])
        # La lettera resta in archivio nonostante l'invio email fallito, e gli
        # altri passi indipendenti (pdf, dcc) restano chiaramente riusciti.
        self.assertTrue(TransmittalInterno.objects.filter(pk=corpo["trasmittal_id"]).exists())
        self.assertTrue(corpo["pdf"]["ok"])
        self.assertTrue(corpo["dcc"]["ok"])
        self.assertEqual(len(mail.outbox), 0)

    # -- elenco lettere emesse --

    def test_elenco_lettere_emesse_in_ordine(self):
        self.client.force_login(self.writer)
        for _ in range(2):
            self.client.post(
                self._url("emetti/"),
                data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
                content_type="application/json",
            )

        risposta = self.client.get(self._url("lettere/"))

        lettere = risposta.json()["lettere"]
        self.assertEqual(len(lettere), 2)
        self.assertEqual(lettere[0]["nome"], "99070_" + timezone.localdate().isoformat() + "_E2")
        self.assertEqual(lettere[1]["nome"], "99070_" + timezone.localdate().isoformat() + "_E1")
        self.assertEqual(lettere[0]["n_documenti"], 1)
        self.assertEqual(lettere[0]["creato_da"], self.writer.nome_completo)


class TrasmittalInternoEmissioneApiTests(TestCase):
    """Anteprima, emissione e retry per-passo (API): il flusso "Nuova lettera"

    gira interamente in modal nella pagina ``trasmittal-interno/``, quindi
    ``self.client`` (che non esegue JavaScript) non può esercitarlo dal vivo.
    L'evidenziazione dello stepper, "Emetti non invia finché il pannello non
    è confermato" (oltre alla garanzia di non-scrittura testata qui) e
    "l'assegnazione siti inline non perde lo stato delle altre righe" restano
    verifiche manuali/di browser — le garanzie server-side sottostanti sono
    invece testate qui, direttamente sulle API, e in
    ``TrasmittalInternoDestinazioniViewTests``.
    """

    def setUp(self):
        self.client = Client()

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        patcher = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        patcher.enable()
        self.addCleanup(patcher.disable)

        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.TO
        )
        self.reparto_ut = Reparto.objects.create(nome="Ufficio Tecnico", acronimo="UT")
        self.testata = Testata.objects.create(job="99075", sito_costruttivo=self.bg)

        self.doc_ok = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99075-ALFA",
            doc_title="Documento Alfa",
            reparto=self.reparto_ut.nome,
        )
        Revisione.objects.create(documento=self.doc_ok, rev_no=0)
        imposta_destinazioni(self.doc_ok, [self.bg.codice_bc])
        from core.services.fileserver import get_base_path

        base = get_base_path(self.testata.job, "UT")
        base.mkdir(parents=True)
        (base / f"{self.doc_ok.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-fake")

        self.writer = User.objects.create_user(
            "tinp_writer", "tinp-writer@b.it", "pw", permesso=Permesso.WRITING
        )
        self.reader = User.objects.create_user(
            "tinp_reader", "tinp-reader@b.it", "pw", permesso=Permesso.READING
        )

    def _url_detail(self):
        return f"/commesse/{self.testata.job}/trasmittal-interno/"

    def _api_url(self, path):
        return f"/api/commesse/{self.testata.job}/trasmittal-interno/{path}"

    def _riga(self, documento, **overrides):
        base = {
            "documento_id": documento.pk,
            "revisione": "A",
            "copie": 1,
            "tpi": False,
            "note": "",
            "cliente": True,
        }
        base.update(overrides)
        return base

    def test_pagina_trasmittal_interno_richiede_sito_costruttivo(self):
        self.testata.sito_costruttivo = None
        self.testata.save(update_fields=["sito_costruttivo"])
        self.client.force_login(self.reader)

        risposta = self.client.get(self._url_detail())

        self.assertEqual(risposta.status_code, 403)

    # -- revisione modificata --

    def test_revisione_modificata_viene_usata_nella_lettera(self):
        self.client.force_login(self.writer)

        risposta = self.client.post(
            self._api_url("emetti/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok, revisione="C")]}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        trasmittal_id = risposta.json()["trasmittal_id"]
        riga = RigaTransmittalInterno.objects.get(
            trasmittal_id=trasmittal_id, documento=self.doc_ok
        )
        self.assertEqual(riga.revisione, "C")

    # -- "Emetti" non invia finché il pannello non è confermato --

    def test_anteprima_non_scrive_nulla(self):
        self.client.force_login(self.reader)

        self.client.post(
            self._api_url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
            content_type="application/json",
        )

        self.assertEqual(TransmittalInterno.objects.count(), 0)

    def test_anteprima_include_sigle_stabilimento(self):
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._api_url("anteprima/"),
            data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
            content_type="application/json",
        )

        sigle = risposta.json()["sigle_stabilimento"]
        self.assertEqual(sigle.get("bg@b.it"), ["BG"])

    # -- esito con un passo fallito: retry per-passo --

    def _emetti_con_email_fallita(self):
        with patch(
            "django.core.mail.EmailMessage.send", side_effect=Exception("SMTP non raggiungibile")
        ):
            risposta = self.client.post(
                self._api_url("emetti/"),
                data=json.dumps({"righe": [self._riga(self.doc_ok)]}),
                content_type="application/json",
            )
        return risposta.json()

    def test_retry_email_riesce_dopo_un_fallimento(self):
        self.client.force_login(self.writer)
        corpo = self._emetti_con_email_fallita()
        self.assertFalse(corpo["email"]["ok"])
        self.assertEqual(len(mail.outbox), 0)

        risposta = self.client.post(
            self._api_url(f"lettere/{corpo['trasmittal_id']}/retry/"),
            data=json.dumps({"step": "email"}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        corpo_retry = risposta.json()
        self.assertTrue(corpo_retry["ok"])
        self.assertTrue(corpo_retry["esito"]["ok"])
        self.assertEqual(len(mail.outbox), 1)

    def test_retry_richiede_login(self):
        risposta = self.client.post(
            self._api_url("lettere/1/retry/"),
            data=json.dumps({"step": "email"}),
            content_type="application/json",
        )
        self.assertEqual(risposta.status_code, 401)

    def test_retry_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.writer)
        corpo = self._emetti_con_email_fallita()
        self.client.force_login(self.reader)

        risposta = self.client.post(
            self._api_url(f"lettere/{corpo['trasmittal_id']}/retry/"),
            data=json.dumps({"step": "email"}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 403)

    def test_retry_lettera_di_un_altra_commessa_404(self):
        self.client.force_login(self.writer)
        corpo = self._emetti_con_email_fallita()

        altra = Testata.objects.create(job="99076", sito_costruttivo=self.bg)
        risposta = self.client.post(
            f"/api/commesse/{altra.job}/trasmittal-interno/lettere/{corpo['trasmittal_id']}/retry/",
            data=json.dumps({"step": "email"}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 404)

    def test_retry_step_non_valido_400(self):
        self.client.force_login(self.writer)
        corpo = self._emetti_con_email_fallita()

        risposta = self.client.post(
            self._api_url(f"lettere/{corpo['trasmittal_id']}/retry/"),
            data=json.dumps({"step": "qualcosa"}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 400)


class AnnullaTrasmittalInternoTests(TestCase):
    """Annullamento di una lettera emessa per errore."""

    def setUp(self):
        self.client = Client()

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        patcher = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        patcher.enable()
        self.addCleanup(patcher.disable)

        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg, email="bg@b.it", tipo=TipoIndirizzoStabilimento.TO
        )
        self.reparto_ut = Reparto.objects.create(nome="Ufficio Tecnico", acronimo="UT")
        self.testata = Testata.objects.create(job="99080", sito_costruttivo=self.bg)

        from core.services.fileserver import get_base_path

        self.get_base_path = get_base_path
        base = get_base_path(self.testata.job, "UT")
        base.mkdir(parents=True)

        self.doc1 = Documento.objects.create(
            testata=self.testata, vendor_doc="99080-ALFA", reparto=self.reparto_ut.nome
        )
        Revisione.objects.create(documento=self.doc1, rev_no=0)
        imposta_destinazioni(self.doc1, [self.bg.codice_bc])
        (base / f"{self.doc1.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-1")

        self.doc2 = Documento.objects.create(
            testata=self.testata, vendor_doc="99080-BETA", reparto=self.reparto_ut.nome
        )
        Revisione.objects.create(documento=self.doc2, rev_no=0)
        imposta_destinazioni(self.doc2, [self.bg.codice_bc])
        (base / f"{self.doc2.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-2")

        self.writer = User.objects.create_user(
            "ann_writer", "ann-writer@b.it", "pw", permesso=Permesso.WRITING
        )
        self.reader = User.objects.create_user(
            "ann_reader", "ann-reader@b.it", "pw", permesso=Permesso.READING
        )

    def _url(self, path):
        return f"/api/commesse/{self.testata.job}/trasmittal-interno/{path}"

    def _emetti(self, documento):
        risposta = self.client.post(
            self._url("emetti/"),
            data=json.dumps(
                {"righe": [{"documento_id": documento.pk, "revisione": "A", "cliente": True}]}
            ),
            content_type="application/json",
        )
        return risposta.json()

    def test_annullamento_dell_ultimo_riesce_e_rimuove_record_pdf_e_cartella_dcc(self):
        self.client.force_login(self.writer)
        emesso = self._emetti(self.doc1)
        pdf_path = Path(emesso["pdf"]["percorso"])
        cartella_dcc = Path(emesso["dcc"]["cartella"])
        self.assertTrue(pdf_path.is_file())
        self.assertTrue(cartella_dcc.is_dir())

        risposta = self.client.post(self._url(f"lettere/{emesso['trasmittal_id']}/annulla/"))

        self.assertEqual(risposta.status_code, 200)
        corpo = risposta.json()
        self.assertTrue(corpo["ok"])
        self.assertTrue(corpo["pdf_rimosso"])
        self.assertTrue(corpo["cartella_dcc_rimossa"])
        self.assertIn("email", corpo["email_avviso"].lower())
        self.assertFalse(TransmittalInterno.objects.filter(pk=emesso["trasmittal_id"]).exists())
        self.assertFalse(pdf_path.exists())
        self.assertFalse(cartella_dcc.exists())

    def test_annullamento_di_uno_non_ultimo_rifiutato(self):
        self.client.force_login(self.writer)
        primo = self._emetti(self.doc1)
        self._emetti(self.doc2)  # secondo, ora è lui l'ultimo

        risposta = self.client.post(self._url(f"lettere/{primo['trasmittal_id']}/annulla/"))

        self.assertEqual(risposta.status_code, 409)
        self.assertIn("ultimo", risposta.json()["error"].lower())
        self.assertTrue(TransmittalInterno.objects.filter(pk=primo["trasmittal_id"]).exists())
        self.assertTrue(Path(primo["pdf"]["percorso"]).is_file())

    def test_progressivo_successivo_riparte_dal_numero_liberato(self):
        self.client.force_login(self.writer)
        emesso = self._emetti(self.doc1)
        self.assertTrue(emesso["nome"].endswith("_E1"))
        self.client.post(self._url(f"lettere/{emesso['trasmittal_id']}/annulla/"))

        rifatto = self._emetti(self.doc2)

        self.assertTrue(rifatto["nome"].endswith("_E1"))

    def test_cartella_dcc_con_contenuto_estraneo_non_viene_rimossa(self):
        self.client.force_login(self.writer)
        emesso = self._emetti(self.doc1)
        cartella_dcc = Path(emesso["dcc"]["cartella"])
        estraneo = cartella_dcc / "documento_non_nostro.pdf"
        estraneo.write_bytes(b"%PDF-estraneo")

        risposta = self.client.post(self._url(f"lettere/{emesso['trasmittal_id']}/annulla/"))

        self.assertEqual(risposta.status_code, 200)
        corpo = risposta.json()
        self.assertFalse(corpo["cartella_dcc_rimossa"])
        # Il file di questo trasmittal è stato rimosso, quello estraneo no.
        self.assertTrue(cartella_dcc.is_dir())
        self.assertEqual([p.name for p in cartella_dcc.iterdir()], ["documento_non_nostro.pdf"])
        self.assertTrue(estraneo.exists())

    def test_annullamento_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.writer)
        emesso = self._emetti(self.doc1)
        self.client.logout()
        self.client.force_login(self.reader)

        risposta = self.client.post(self._url(f"lettere/{emesso['trasmittal_id']}/annulla/"))

        self.assertEqual(risposta.status_code, 403)
        self.assertTrue(TransmittalInterno.objects.filter(pk=emesso["trasmittal_id"]).exists())

    def test_annullamento_richiede_login(self):
        emesso_writer = Client()
        emesso_writer.force_login(self.writer)
        emesso = emesso_writer.post(
            self._url("emetti/"),
            data=json.dumps(
                {"righe": [{"documento_id": self.doc1.pk, "revisione": "A", "cliente": True}]}
            ),
            content_type="application/json",
        ).json()

        risposta = self.client.post(self._url(f"lettere/{emesso['trasmittal_id']}/annulla/"))

        self.assertNotEqual(risposta.status_code, 200)

    def test_annullabile_solo_sull_ultimo_nella_lista_lettere(self):
        self.client.force_login(self.writer)
        primo = self._emetti(self.doc1)
        secondo = self._emetti(self.doc2)

        lettere = {
            voce["id"]: voce for voce in self.client.get(self._url("lettere/")).json()["lettere"]
        }

        self.assertFalse(lettere[primo["trasmittal_id"]]["annullabile"])
        self.assertTrue(lettere[secondo["trasmittal_id"]]["annullabile"])


class TrasmittalInternoEndToEndTests(TestCase):
    """End-to-end: percorre il flusso utente del trasmittal interno con i dati
    reali del vecchio strumento Excel/VBA (anagrafiche, indirizzi, firmatari)
    e confronta il risultato con il comportamento atteso di quello strumento.

    Trova divergenze, non le corregge (vedi il report consegnato con la PR).
    Nessun invio di posta reale: solo il backend locmem di Django (mai
    sostituito in questa classe), fileserver su una directory temporanea,
    nessun accesso a Business Central.
    """

    def setUp(self):
        self.client = Client()

        # Garanzia esplicita: mai un backend email reale in questi test.
        from django.core.mail import get_connection

        self.assertIn("locmem", type(get_connection()).__module__)

        self.jobs_root = tempfile.TemporaryDirectory()
        self.addCleanup(self.jobs_root.cleanup)
        fs_patch = override_settings(FILESERVER_JOBS_PATH=self.jobs_root.name)
        fs_patch.enable()
        self.addCleanup(fs_patch.disable)

        # ── Stabilimenti ──
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.pd = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        self.ve = Stabilimento.objects.create(nome="Marghera", sigla="VE", codice_bc=3)
        self.cr = Stabilimento.objects.create(nome="Ricengo", sigla="CR", codice_bc=4)
        self.vi = Stabilimento.objects.create(nome="Schio", sigla="VI", codice_bc=5)
        self.milano = Stabilimento.objects.create(nome="Milano")  # nessun codice sito

        # ── Indirizzi TO/CC per sito (dati del vecchio strumento) ──
        _TO = {
            self.bg: ["fdamiani", "locatellim"],
            self.pd: ["egomiero", "mtoniolo", "dgiunchi", "dgigante"],
            self.ve: ["egomiero", "mtoniolo", "dgiunchi", "dgigante"],  # identici a PD
            self.cr: ["epaparazzo", "msolazzo", "mmaggi", "fcorradini", "poro"],
            self.vi: ["mgasparini", "anovella", "thossain"],
        }
        _CC = {
            self.bg: ["dpasserini", "quality", "mcheccolin", "fbaldin", "mcarminati", "mpersoneni"],
            self.pd: ["dpasserini", "lterrassan", "mgalli", "asandona"],
            self.ve: ["dpasserini", "psaccarola", "mgalli", "asandona"],
            self.cr: ["dpasserini", "egritti", "rlucini", "egalbiati"],
            self.vi: ["dpasserini", "lferracin", "aottoboni", "apunturieri"],
        }
        for stabilimento, nomi in _TO.items():
            for nome in nomi:
                IndirizzoStabilimento.objects.create(
                    stabilimento=stabilimento,
                    email=f"{nome}@brembanarolle.com",
                    tipo=TipoIndirizzoStabilimento.TO,
                )
        for stabilimento, nomi in _CC.items():
            for nome in nomi:
                IndirizzoStabilimento.objects.create(
                    stabilimento=stabilimento,
                    email=f"{nome}@brembanarolle.com",
                    tipo=TipoIndirizzoStabilimento.CC,
                )

        # ── Firmatari Produzione/Qualità per sito ──
        def _utente_firmatario(nome):
            return User.objects.create_user(nome.lower(), f"{nome.lower()}@brembanarolle.com", "pw")

        self.firmatari_utenti = {
            nome: _utente_firmatario(nome)
            for nome in (
                "MLocatelli",
                "FBaldin",
                "DGiunchi",
                "MGalli",
                "POro",
                "RLucini",
                "THossain",
                "AOttoboni",
            )
        }
        _FIRMATARI = {
            self.bg: ("MLocatelli", "FBaldin"),
            self.pd: ("DGiunchi", "MGalli"),
            self.ve: ("DGiunchi", "MGalli"),  # stessi di PD
            self.cr: ("POro", "RLucini"),
            self.vi: ("THossain", "AOttoboni"),
        }
        for stabilimento, (produzione, qualita) in _FIRMATARI.items():
            FirmatarioStabilimento.objects.create(
                stabilimento=stabilimento,
                ruolo=RuoloFirmatarioStabilimento.PRODUZIONE,
                utente=self.firmatari_utenti[produzione],
            )
            FirmatarioStabilimento.objects.create(
                stabilimento=stabilimento,
                ruolo=RuoloFirmatarioStabilimento.QUALITA,
                utente=self.firmatari_utenti[qualita],
            )
        # Nessun firmatario ha un'immagine di firma caricata (punto 9): il
        # campo firma resta vuoto per tutti, di proposito.

        # ── Commessa e documenti UT ──
        self.reparto_ut = Reparto.objects.create(nome="Ufficio Tecnico", acronimo="UT")
        self.testata = Testata.objects.create(job="99090", sito_costruttivo=self.bg)

        from core.services.fileserver import get_base_path

        base = get_base_path(self.testata.job, "UT")
        base.mkdir(parents=True)

        self.doc_a = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99090-01-DWGA",
            doc_title="Documento A",
            reparto="Ufficio Tecnico",
        )
        self.doc_b = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99090-01-DWGB",
            doc_title="Documento B",
            reparto="Ufficio Tecnico",
        )
        self.doc_c = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99090-01-ESH1",  # pattern SHn: ?????-??-ESH*
            doc_title="Documento C (SHn)",
            reparto="Ufficio Tecnico",
        )
        self.doc_d = Documento.objects.create(
            testata=self.testata,
            vendor_doc="99090-01-DWGD",
            doc_title="Documento D",
            reparto="Ufficio Tecnico",
        )
        for doc in (self.doc_a, self.doc_b, self.doc_c, self.doc_d):
            Revisione.objects.create(documento=doc, rev_no=0)
        for doc in (self.doc_a, self.doc_b, self.doc_c):  # DOC-D: file mancante di proposito
            (base / f"{doc.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-fake")

        # ── PM/PE/QCI ──
        self.pm_utente = User.objects.create_user(
            "mtoniolo", "mtoniolo@brembanarolle.com", "pw", first_name="Mauro", last_name="Toniolo"
        )
        self.pe_utente = User.objects.create_user(
            "e2e_pe", "pe.nuovo@brembanarolle.com", "pw", first_name="Paolo", last_name="Erre"
        )
        self.qci_utente = User.objects.create_user(
            "e2e_qci", "qci.nuovo@brembanarolle.com", "pw", first_name="Quinto", last_name="Ci"
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, utente=self.pm_utente
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PE, utente=self.pe_utente
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.QCI, utente=self.qci_utente
        )

        self.writer = User.objects.create_user(
            "e2e_writer", "e2e-writer@b.it", "pw", permesso=Permesso.WRITING
        )
        self.client.force_login(self.writer)

        # ── Percorso 1-3: dalla pagina commessa alla griglia destinazioni,
        #    tutto attraverso il test client (mai chiamate dirette ai servizi) ──
        self._imposta_destinazioni(self.doc_a, [self.pd.codice_bc, self.ve.codice_bc])
        self._imposta_destinazioni(self.doc_b, [self.bg.codice_bc])
        self._imposta_destinazioni(self.doc_c, [self.bg.codice_bc])
        # DOC-D: nessuna destinazione, di proposito.

    # ── Helper: navigazione del flusso via test client ──────────────────────

    def _url(self, path):
        return f"/api/commesse/{self.testata.job}/trasmittal-interno/{path}"

    def _imposta_destinazioni(self, documento, codici_bc):
        risposta = self.client.post(
            self._url(f"documenti/{documento.pk}/destinazioni/"),
            data=json.dumps({"codici_bc": codici_bc}),
            content_type="application/json",
        )
        self.assertEqual(risposta.status_code, 200, risposta.content)
        return risposta.json()

    def _riga(self, documento, **overrides):
        base = {
            "documento_id": documento.pk,
            "revisione": "0",
            "copie": 2,
            "tpi": False,
            "note": "",
            "cliente": True,
        }
        base.update(overrides)
        return base

    def _anteprima(self, righe, note=""):
        risposta = self.client.post(
            self._url("anteprima/"),
            data=json.dumps({"righe": righe, "note": note}),
            content_type="application/json",
        )
        self.assertEqual(risposta.status_code, 200, risposta.content)
        return risposta.json()

    def _emetti(self, righe, note="", destinatari=None):
        payload = {"righe": righe, "note": note}
        if destinatari is not None:
            payload["destinatari"] = destinatari
        risposta = self.client.post(
            self._url("emetti/"),
            data=json.dumps(payload),
            content_type="application/json",
        )
        return risposta

    def _righe_abcd(self, con_c=True):
        righe = [self._riga(self.doc_a, cliente=True), self._riga(self.doc_b, cliente=False)]
        if con_c:
            righe.append(self._riga(self.doc_c, cliente=True))
        return righe

    # ── Percorso: card della commessa e pagina dedicata ─────────────────────

    def test_01_card_sbloccata_e_pagina_raggiungibile(self):
        pagina_commessa = self.client.get(f"/commesse/{self.testata.job}/")
        self.assertEqual(pagina_commessa.status_code, 200)
        self.assertNotContains(pagina_commessa, 'class="section-card section-card-locked"')
        self.assertContains(pagina_commessa, f"/commesse/{self.testata.job}/trasmittal-interno/")

        pagina_trasmittal = self.client.get(f"/commesse/{self.testata.job}/trasmittal-interno/")
        self.assertEqual(pagina_trasmittal.status_code, 200)

    # ── Punto 19: Milano non è selezionabile come destinazione ──────────────

    def test_02_milano_non_selezionabile_come_destinazione(self):
        risposta = self.client.get(self._url("destinazioni/"))
        sigle = [s["sigla"] for s in risposta.json()["stabilimenti"]]
        self.assertNotIn(None, sigle)
        self.assertEqual(sorted(sigle), ["BG", "CR", "PD", "VE", "VI"])

    # ── Punto 14: documento senza file segnalato in selezione, non blocca ───

    def test_03_doc_d_segnalato_in_selezione_gli_altri_selezionabili(self):
        risposta = self.client.get(self._url("selezione/"))
        stati = {d["id"]: d for d in risposta.json()["documenti"]}

        for doc in (self.doc_a, self.doc_b, self.doc_c):
            self.assertTrue(stati[doc.pk]["selezionabile"], stati[doc.pk])

        self.assertFalse(stati[self.doc_d.pk]["selezionabile"])
        self.assertEqual(stati[self.doc_d.pk]["motivo"], "Nessun file trovato sul fileserver.")

    def test_03b_lettera_si_crea_senza_doc_d_nonostante_la_sua_presenza_in_elenco(self):
        risposta = self._emetti(self._righe_abcd())
        self.assertEqual(risposta.status_code, 200, risposta.content)
        self.assertTrue(risposta.json()["ok"])

    # ── Punti 1-2-3: dedup TO/CC su PD+VE+BG (indirizzi identici per PD/VE) ──

    def test_04_to_pd_ve_compaiono_una_volta_ciascuno(self):
        anteprima = self._anteprima(self._righe_abcd())
        to = [d["email"] for d in anteprima["destinatari"] if d["tipo"] == "to"]
        for nome in ("egomiero", "mtoniolo", "dgiunchi", "dgigante"):
            email = f"{nome}@brembanarolle.com"
            self.assertEqual(to.count(email), 1, f"{email}: atteso 1, trovato {to.count(email)}")

    def test_05_dpasserini_una_volta_sola_nonostante_tre_siti(self):
        anteprima = self._anteprima(self._righe_abcd())
        emails = [d["email"] for d in anteprima["destinatari"]]
        self.assertEqual(emails.count("dpasserini@brembanarolle.com"), 1)

    def test_06_mgalli_asandona_una_volta_ciascuno(self):
        anteprima = self._anteprima(self._righe_abcd())
        emails = [d["email"] for d in anteprima["destinatari"]]
        self.assertEqual(emails.count("mgalli@brembanarolle.com"), 1)
        self.assertEqual(emails.count("asandona@brembanarolle.com"), 1)

    # ── Punto 4: PM già in TO di PD → una sola occorrenza, in TO ────────────

    def test_07_pm_mtoniolo_una_sola_occorrenza_in_to(self):
        anteprima = self._anteprima(self._righe_abcd())
        occorrenze = [
            d for d in anteprima["destinatari"] if d["email"] == "mtoniolo@brembanarolle.com"
        ]
        self.assertEqual(len(occorrenze), 1)
        self.assertEqual(occorrenze[0]["tipo"], "to")

    # ── Punto 5: un indirizzo in TO e in CC resta solo in TO ─────────────────
    # I dati di partenza dati dal task, applicati alla lettera BG+PD+VE, non
    # producono di per sé una collisione TO/CC (i nominativi TO e CC dei tre
    # siti coinvolti sono tutti distinti — vedi il report). Verifico quindi
    # la regola con un caso sintetico minimo, sullo stesso sito BG già in
    # gioco, aggiungendo un indirizzo apposta presente sia in TO che in CC.

    def test_08_indirizzo_in_to_e_in_cc_resta_solo_in_to(self):
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg,
            email="doppio@brembanarolle.com",
            tipo=TipoIndirizzoStabilimento.CC,
        )
        IndirizzoStabilimento.objects.create(
            stabilimento=self.bg,
            email="doppio@brembanarolle.com",
            tipo=TipoIndirizzoStabilimento.TO,
        )
        anteprima = self._anteprima(self._righe_abcd())
        occorrenze = [
            d for d in anteprima["destinatari"] if d["email"] == "doppio@brembanarolle.com"
        ]
        self.assertEqual(len(occorrenze), 1)
        self.assertEqual(occorrenze[0]["tipo"], "to")

    # ── Punto 6: DOC-C (SHn) → export@ in CC; rimosso, sparisce ──────────────

    def test_09_export_presente_con_doc_c_assente_senza(self):
        anteprima_con_c = self._anteprima(self._righe_abcd(con_c=True))
        emails_con_c = [d["email"] for d in anteprima_con_c["destinatari"]]
        self.assertIn("export@brembanarolle.com", emails_con_c)

        anteprima_senza_c = self._anteprima(self._righe_abcd(con_c=False))
        emails_senza_c = [d["email"] for d in anteprima_senza_c["destinatari"]]
        self.assertNotIn("export@brembanarolle.com", emails_senza_c)

    # ── Punto 7: PM/PE/QCI non valorizzati → lettera creata comunque ────────

    def test_10_lettera_creata_senza_pm_pe_qci_valorizzati(self):
        altra = Testata.objects.create(job="99091", sito_costruttivo=self.bg)
        doc = Documento.objects.create(
            testata=altra, vendor_doc="99091-01-DWGX", reparto="Ufficio Tecnico"
        )
        Revisione.objects.create(documento=doc, rev_no=0)
        from core.services.fileserver import get_base_path

        base = get_base_path(altra.job, "UT")
        base.mkdir(parents=True)
        (base / f"{doc.vendor_doc} Rev A.pdf").write_bytes(b"%PDF-fake")
        self.client.post(
            f"/api/commesse/{altra.job}/trasmittal-interno/documenti/{doc.pk}/destinazioni/",
            data=json.dumps({"codici_bc": [self.bg.codice_bc]}),
            content_type="application/json",
        )

        risposta = self.client.post(
            f"/api/commesse/{altra.job}/trasmittal-interno/emetti/",
            data=json.dumps(
                {"righe": [{"documento_id": doc.pk, "revisione": "0", "cliente": True}]}
            ),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200, risposta.content)
        self.assertTrue(risposta.json()["ok"])

    # ── Punto 8: firmatari PD+VE deduplicati (DGiunchi/MGalli comuni) ───────
    # Nessuna libreria di estrazione testo da PDF è disponibile in questo
    # ambiente (fitz non installato — vedi il fallimento preesistente in
    # SituazioneApiTestCase): verifico quindi la stessa funzione di dedup
    # usata da genera_trasmittal_interno_pdf direttamente sui siti reali
    # della lettera, non analizzando i byte del PDF.

    def test_11_firmatari_pd_ve_deduplicati_per_ruolo(self):
        from src.pdf import _firmatari_per_ruolo

        siti = [self.pd, self.ve]
        produzione = _firmatari_per_ruolo(siti, RuoloFirmatarioStabilimento.PRODUZIONE)
        qualita = _firmatari_per_ruolo(siti, RuoloFirmatarioStabilimento.QUALITA)

        self.assertEqual([f.utente_id for f in produzione], [self.firmatari_utenti["DGiunchi"].pk])
        self.assertEqual([f.utente_id for f in qualita], [self.firmatari_utenti["MGalli"].pk])

    # ── Punto 9: firmatario senza immagine → PDF generato comunque ──────────

    def test_12_pdf_generato_senza_immagini_di_firma(self):
        for firmatario in self.firmatari_utenti.values():
            self.assertFalse(firmatario.firma)

        risposta = self.client.post(
            self._url("anteprima/pdf/"),
            data=json.dumps({"righe": self._righe_abcd()}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(risposta["Content-Type"], "application/pdf")
        self.assertTrue(risposta.content.startswith(b"%PDF"))
        self.assertGreater(len(risposta.content), 1000)

    # ── Punto 10: nome file e percorso ───────────────────────────────────────

    def test_13_nome_file_pdf_e_percorso(self):
        risposta = self._emetti(self._righe_abcd())
        corpo = risposta.json()
        self.assertTrue(corpo["pdf"]["ok"])
        percorso = Path(corpo["pdf"]["percorso"])
        oggi = timezone.localdate().isoformat()
        self.assertEqual(percorso.name, f"99090_{oggi}_E1.pdf")
        self.assertIn(str(Path("99090") / "Progetto" / "UT" / "Transmittal"), str(percorso))
        self.assertTrue(percorso.is_file())

    # ── Punto 11: progressivo giornaliero, reset il giorno dopo ─────────────

    def test_14_progressivo_giornaliero_e_reset_il_giorno_dopo(self):
        giorno1 = date(2026, 9, 15)  # martedì
        giorno2 = date(2026, 9, 16)

        with patch("core.services.trasmittal_interno.timezone.localdate", return_value=giorno1):
            prima = self._emetti(self._righe_abcd()).json()
            seconda = self._emetti([self._riga(self.doc_b, cliente=False)]).json()
        with patch("core.services.trasmittal_interno.timezone.localdate", return_value=giorno2):
            terza = self._emetti([self._riga(self.doc_c, cliente=True)]).json()

        self.assertTrue(prima["nome"].endswith("_E1"))
        self.assertTrue(seconda["nome"].endswith("_E2"))
        self.assertTrue(terza["nome"].endswith("_E1"))

    # ── Punto 12: data di distribuzione (venerdì → lunedì, altrimenti dopo) ──
    # La data di distribuzione mostrata all'utente riflette data_impegno()
    # (giorno lavorativo successivo, lunedì se venerdì), non la data di
    # emissione: verifico che generare il PDF invochi davvero data_impegno.

    def test_15_data_impegno_e_invocata_generando_il_pdf(self):
        from core.services.trasmittal_interno import data_impegno
        from src.pdf import genera_trasmittal_interno_pdf

        venerdi = date(2026, 9, 18)
        self.assertEqual(data_impegno(venerdi), date(2026, 9, 21))  # lunedì, per DataImpegnoTests

        with patch("core.services.trasmittal_interno.timezone.localdate", return_value=venerdi):
            esito = self._emetti(self._righe_abcd()).json()
        trasmittal = TransmittalInterno.objects.get(pk=esito["trasmittal_id"])
        self.assertEqual(trasmittal.data, venerdi)

        with patch(
            "core.services.trasmittal_interno.data_impegno", side_effect=data_impegno
        ) as mock_impegno:
            genera_trasmittal_interno_pdf(trasmittal)

        self.assertTrue(
            mock_impegno.called,
            "data_impegno() non è invocata generando il PDF: la colonna DATE di "
            "'PAPER COPIES DISTRIBUTION' mostrerebbe la data di emissione invece "
            "del giorno lavorativo successivo (lunedì se venerdì).",
        )

    # ── Punti 15-16: oggetto e corpo dell'email ──────────────────────────────

    def test_16_oggetto_email_formato_atteso(self):
        esito = self._emetti(self._righe_abcd()).json()
        self.assertEqual(len(mail.outbox), 1)
        atteso = f"DOCUMENT TRANSMITTAL [Form MQ 7.5-04 Rev.0]: {esito['nome']}"
        self.assertEqual(mail.outbox[0].subject, atteso)

    def test_17_corpo_email_contiene_le_righe_il_percorso_e_il_promemoria_firma(self):
        esito = self._emetti(self._righe_abcd()).json()
        self.assertEqual(len(mail.outbox), 1)
        corpo = mail.outbox[0].body
        for doc in (self.doc_a, self.doc_b, self.doc_c):
            self.assertIn(doc.vendor_doc, corpo)
        # Non solo i documenti: anche gli altri valori di riga, come nel PDF.
        self.assertIn("YES", corpo)  # CLIENT di doc_a/doc_c
        self.assertIn("NO", corpo)  # CLIENT di doc_b
        self.assertIn(esito["pdf"]["percorso"], corpo)
        self.assertIn("firmat", corpo.lower())

    # ── Punto 17: destinatari modificati in anteprima → usati e registrati ──

    def test_18_destinatari_modificati_in_anteprima_usati_e_registrati(self):
        confermati = [
            {"email": "destinatario.scelto@brembanarolle.com", "tipo": "to", "origine": "manuale"},
        ]
        risposta = self._emetti(self._righe_abcd(), destinatari=confermati)
        esito = risposta.json()

        self.assertEqual(mail.outbox[0].to, ["destinatario.scelto@brembanarolle.com"])
        trasmittal = TransmittalInterno.objects.get(pk=esito["trasmittal_id"])
        registrati = list(trasmittal.destinatari.values_list("email", "tipo"))
        self.assertEqual(registrati, [("destinatario.scelto@brembanarolle.com", "to")])

    # ── Punto 18: il PDF è allegato ──────────────────────────────────────────

    def test_19_pdf_allegato_alla_email(self):
        esito = self._emetti(self._righe_abcd()).json()
        self.assertEqual(len(mail.outbox[0].attachments), 1)
        nome_allegato, contenuto, mimetype = mail.outbox[0].attachments[0]
        self.assertEqual(nome_allegato, f"{esito['nome']}.pdf")
        self.assertEqual(mimetype, "application/pdf")
        self.assertTrue(contenuto.startswith(b"%PDF"))

    # ── Punto 13: DOC-B (CLIENT NO) non in DCC; DOC-A/DOC-C sì ──────────────

    def test_20_doc_b_escluso_dal_dcc_doc_a_e_c_inclusi(self):
        esito = self._emetti(self._righe_abcd())
        corpo = esito.json()
        copiati = {c["vendor_doc"] for c in corpo["dcc"]["copiati"]}

        self.assertIn(self.doc_a.vendor_doc, copiati)
        self.assertIn(self.doc_c.vendor_doc, copiati)
        self.assertNotIn(self.doc_b.vendor_doc, copiati)

        cartella = Path(corpo["dcc"]["cartella"])
        nomi_file = [p.name for p in cartella.iterdir()]
        self.assertTrue(any(self.doc_a.vendor_doc in n for n in nomi_file))
        self.assertTrue(any(self.doc_c.vendor_doc in n for n in nomi_file))
        self.assertFalse(any(self.doc_b.vendor_doc in n for n in nomi_file))

    # ── Punto 20: modificare le destinazioni dopo l'emissione non cambia lo
    #    snapshot della lettera già emessa ────────────────────────────────

    def test_21_destinazioni_modificate_dopo_emissione_non_toccano_lo_snapshot(self):
        esito = self._emetti(self._righe_abcd())
        trasmittal = TransmittalInterno.objects.get(pk=esito.json()["trasmittal_id"])
        riga_a = trasmittal.righe.get(documento=self.doc_a)
        siti_originali = sorted(s.sigla for s in riga_a.siti.all())
        self.assertEqual(siti_originali, ["PD", "VE"])

        self._imposta_destinazioni(self.doc_a, [self.bg.codice_bc])

        riga_a.refresh_from_db()
        siti_dopo_modifica = sorted(s.sigla for s in riga_a.siti.all())
        self.assertEqual(siti_dopo_modifica, ["PD", "VE"])  # invariato: snapshot

        # Le destinazioni "vive" del documento, invece, sono cambiate davvero.
        risposta = self.client.get(self._url("destinazioni/"))
        doc_a_live = next(d for d in risposta.json()["documenti"] if d["id"] == self.doc_a.pk)
        self.assertEqual(doc_a_live["codici_bc"], [self.bg.codice_bc])

    # ── Punto 21: annullamento libera PDF, cartella DCC e progressivo ──────

    def test_22_annullamento_rimuove_pdf_cartella_dcc_e_libera_il_progressivo(self):
        esito = self._emetti(self._righe_abcd())
        corpo = esito.json()
        pdf_path = Path(corpo["pdf"]["percorso"])
        cartella_dcc = Path(corpo["dcc"]["cartella"])
        self.assertTrue(pdf_path.is_file())
        self.assertTrue(cartella_dcc.is_dir())

        risposta = self.client.post(self._url(f"lettere/{corpo['trasmittal_id']}/annulla/"))

        self.assertEqual(risposta.status_code, 200, risposta.content)
        annulla_corpo = risposta.json()
        self.assertTrue(annulla_corpo["pdf_rimosso"])
        self.assertTrue(annulla_corpo["cartella_dcc_rimossa"])
        self.assertFalse(pdf_path.exists())
        self.assertFalse(cartella_dcc.exists())
        self.assertFalse(TransmittalInterno.objects.filter(pk=corpo["trasmittal_id"]).exists())

        rifatto = self._emetti([self._riga(self.doc_b, cliente=False)]).json()
        self.assertTrue(rifatto["nome"].endswith("_E1"))  # progressivo liberato, non E2


class FileserverIntegrationTest(TestCase):
    """Integration tests against the real fileserver at Z:\\JOBS.

    These tests create their own DB fixtures in the test database and access
    the real fileserver at Z:\\JOBS (no mocking). The test is skipped
    automatically if the fileserver is not reachable.

    Expected results based on real fileserver content at Z:\\JOBS\\25056\\PROGETTO\\QC\\:
        - QMDBI Rev.0  (rec=2025-11-17): 1 file  → found
        - QCPA  Rev.0  (rec=2025-10-14): 1 file  → found
        - QCPA  Rev.1  (rec=2025-11-10): 3 files → multiple
        - QCPA  Rev.2  (rec=2025-11-16): 1 file  → found
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from django.conf import settings

        cls.jobs_root = Path(settings.FILESERVER_JOBS_PATH)
        cls.fileserver_available = cls.jobs_root.exists()

    def setUp(self):
        if not self.fileserver_available:
            self.skipTest(f"Fileserver non raggiungibile: {self.jobs_root}")
        self.fixtures = _make_integration_fixtures()
        self.user = User.objects.create_user(
            "integration_user", password="pw", permesso=Permesso.WRITING
        )
        self.client = Client()
        self.client.force_login(self.user)

    def tearDown(self):
        RevisioneFileLink.objects.filter(revisione__in=self.fixtures.values()).delete()

    # ── risolvi_file_revisione ────────────────────────────────────────────────

    def test_qmdbi_rev0_found(self):
        """QMDBI Rev.0 ricevuta 17/11/25 → deve trovare 1 file."""
        result = risolvi_file_revisione(self.fixtures["rev_qmdbi_0"].pk)
        self.assertIn(
            result["status"],
            ("found", "multiple"),
            f"Atteso found/multiple, ottenuto {result['status']}.\n"
            f"Cerca in: Z:\\JOBS\\25056\\PROGETTO\\QC\\RICEVUTI\\2025-11-17\\",
        )
        if result["status"] == "found":
            self.assertEqual(len(result["files"]), 1)
            nome = result["files"][0]["nome"]
            self.assertTrue(
                nome.lower().startswith("25056-qmdbi"),
                f"Nome file inatteso: {nome!r}",
            )

    def test_qcpa_rev0_found(self):
        """QCPA Rev.0 ricevuta 14/10/25 → deve trovare 1 file."""
        result = risolvi_file_revisione(self.fixtures["rev_qcpa_0"].pk)
        self.assertIn(
            result["status"],
            ("found", "multiple"),
            f"Atteso found/multiple, ottenuto {result['status']}.\n"
            f"Cerca in: Z:\\JOBS\\25056\\PROGETTO\\QC\\RICEVUTI\\2025-10-14\\",
        )

    def test_qcpa_rev1_multiple(self):
        """QCPA Rev.1 ricevuta 10/11/25 → più file presenti → multiple."""
        result = risolvi_file_revisione(self.fixtures["rev_qcpa_1"].pk)
        self.assertEqual(
            result["status"],
            "multiple",
            f"Atteso multiple (3 file noti), ottenuto {result['status']}.\n"
            f"File trovati: {[f['nome'] for f in result['files']]}",
        )
        self.assertEqual(
            len(result["files"]),
            3,
            f"Attesi 3 file, trovati {len(result['files'])}: "
            f"{[f['nome'] for f in result['files']]}",
        )

    def test_qcpa_rev2_found(self):
        """QCPA Rev.2 ricevuta 16/11/25 → deve trovare 1 file."""
        result = risolvi_file_revisione(self.fixtures["rev_qcpa_2"].pk)
        self.assertIn(
            result["status"],
            ("found", "multiple"),
            f"Atteso found/multiple, ottenuto {result['status']}.\n"
            f"Cerca in: Z:\\JOBS\\25056\\PROGETTO\\QC\\RICEVUTI\\2025-11-16\\",
        )

    def test_manual_link_overrides_filesystem(self):
        """Se salvo un link manuale, risolvi deve restituire quello, non il file FS."""
        real_file = str(
            self.jobs_root
            / "25056"
            / "PROGETTO"
            / "QC"
            / "RICEVUTI"
            / "2025-11-17"
            / "25056-QMDBI rev.0 - commented.pdf"
        )
        salva_file_link(self.fixtures["rev_qmdbi_0"].pk, real_file)
        result = risolvi_file_revisione(self.fixtures["rev_qmdbi_0"].pk)
        self.assertEqual(result["status"], "manual_linked")
        self.assertEqual(result["files"][0]["percorso"], real_file)

    # ── API endpoint ─────────────────────────────────────────────────────────

    def test_api_resolve_qmdbi_rev0(self):
        """GET /api/revisioni/{pk}/file/ deve restituire found o multiple."""
        pk = self.fixtures["rev_qmdbi_0"].pk
        response = self.client.get(f"/api/revisioni/{pk}/file/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn(
            data["status"],
            ("found", "multiple"),
            f"Ottenuto {data['status']}; files={data.get('files')}",
        )

    def test_api_resolve_qcpa_rev1_multiple(self):
        """GET /api/revisioni/{pk}/file/ per Rev.1 QCPA deve restituire multiple."""
        pk = self.fixtures["rev_qcpa_1"].pk
        response = self.client.get(f"/api/revisioni/{pk}/file/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(
            data["status"],
            "multiple",
            f"Atteso multiple, ottenuto {data['status']}; "
            f"files={[f['nome'] for f in data.get('files', [])]}",
        )

    def test_api_save_and_resolve_manual_link(self):
        """POST link → GET risolve manual_linked."""
        pk = self.fixtures["rev_qmdbi_0"].pk
        real_file = str(
            self.jobs_root
            / "25056"
            / "PROGETTO"
            / "QC"
            / "RICEVUTI"
            / "2025-11-17"
            / "25056-QMDBI rev.0 - commented.pdf"
        )
        resp = self.client.post(
            f"/api/revisioni/{pk}/file/link/",
            data=json.dumps({"percorso": real_file}),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)

        resp2 = self.client.get(f"/api/revisioni/{pk}/file/")
        self.assertEqual(resp2.status_code, 200)
        self.assertEqual(resp2.json()["status"], "manual_linked")


class PermessiTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.admin_user = User.objects.create_user(
            "admin_user", "admin@example.com", "pw", permesso=Permesso.ADMIN
        )
        self.writing_user = User.objects.create_user(
            "writing_user", "writing@example.com", "pw", permesso=Permesso.WRITING
        )
        self.reading_user = User.objects.create_user(
            "reading_user", "reading@example.com", "pw", permesso=Permesso.READING
        )

    def test_reading_can_list_commesse(self):
        self.client.force_login(self.reading_user)
        response = self.client.get("/api/commesse/")
        self.assertEqual(response.status_code, 200)

    def test_reading_cannot_create_commessa(self):
        self.client.force_login(self.reading_user)
        response = self.client.post(
            "/api/commesse/",
            data=json.dumps({"job": "READ-001"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["error"], "Permesso negato.")

    def test_writing_can_create_commessa(self):
        self.client.force_login(self.writing_user)
        response = self.client.post(
            "/api/commesse/",
            data=json.dumps({"job": "WRITE-001"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)

    def test_delete_commessa_sends_request_email_without_deleting(self):
        Testata.objects.create(job="DEL-001", job_detail="Prova eliminazione")
        self.client.force_login(self.writing_user)
        response = self.client.delete("/api/commesse/DEL-001/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        self.assertTrue(Testata.objects.filter(job="DEL-001").exists())
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["plitta@brembanarolle.com"])
        self.assertIn("writing_user", mail.outbox[0].body)
        self.assertIn("DEL-001", mail.outbox[0].body)
        self.assertIn("ha richiesto l'eliminazione della commessa", mail.outbox[0].body)

    def test_admin_delete_commessa_removes_it(self):
        Testata.objects.create(job="DEL-ADM", job_detail="Eliminazione admin")
        self.client.force_login(self.admin_user)
        response = self.client.delete("/api/commesse/DEL-ADM/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["ok"])
        self.assertFalse(Testata.objects.filter(job="DEL-ADM").exists())
        self.assertEqual(len(mail.outbox), 0)

    def test_reading_cannot_access_admin(self):
        self.client.force_login(self.reading_user)
        response = self.client.get("/admin/")
        self.assertIn(response.status_code, (302, 403))

    def test_admin_can_access_admin(self):
        self.client.force_login(self.admin_user)
        response = self.client.get("/admin/")
        self.assertEqual(response.status_code, 200)

    def test_registration_defaults_to_reading(self):
        response = self.client.post(
            "/register/",
            data={
                "username": "new_reader",
                "email": "reader@brembanarolle.com",
                "first_name": "Nuovo",
                "last_name": "Lettore",
                "password1": "securepass123",
                "password2": "securepass123",
            },
        )
        self.assertEqual(response.status_code, 302)
        user = User.objects.get(username="new_reader")
        self.assertEqual(user.permesso, Permesso.READING)
        self.assertFalse(user.is_staff)

    def test_registration_rejects_non_company_email(self):
        response = self.client.post(
            "/register/",
            data={
                "username": "external_user",
                "email": "user@example.com",
                "first_name": "Esterno",
                "last_name": "Utente",
                "password1": "securepass123",
                "password2": "securepass123",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "@brembanarolle.com")
        self.assertFalse(User.objects.filter(username="external_user").exists())

    def test_admin_syncs_is_staff(self):
        self.assertTrue(self.admin_user.is_staff)
        self.assertTrue(self.writing_user.is_staff)
        self.assertFalse(self.reading_user.is_staff)


class RevisioniCleanupTests(TestCase):
    def test_single_orphan_after_approved_with_only_dis_plan(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "C",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": "2025-02-01",
                    "RecPlanDate": "2025-02-10",
                    "RecActDate": "2025-02-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 2,
                    "DisPlanDate": "2025-03-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        orphans = find_orphan_indices(df)
        self.assertEqual(orphans, {2})
        cleaned = drop_orphan_revisioni(df)
        self.assertEqual(len(cleaned), 2)

    def test_not_orphan_when_previous_status_needs_new_rev(self):
        # Status C has crea_nuova_rev=True by default → keep last placeholder.
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "C",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())

    def test_not_orphan_when_dis_act_date_present(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": "2025-02-01",
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())

    def test_not_orphan_without_dis_plan_date(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": None,
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())


class RevisioneAnomalieTests(TestCase):
    def setUp(self):
        self.approved = StatoEsterno.objects.create(nome="Approved", colore="#00B050")
        self.rejected = StatoEsterno.objects.create(
            nome="Rejected - Work can not proceed", colore="#D61D09"
        )
        self.testata = Testata.objects.create(job="25012", time_cli_doc_rev=21)
        self.doc = Documento.objects.create(
            testata=self.testata,
            vendor_doc="25012-01-4005",
            doc_title="Test doc",
        )

    def test_empty_revision_not_anomalous(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=1, dis_plan_date="2026-03-01")
        self.assertEqual(audit_revisione(rev), [])

    def test_risposta_senza_ricezione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        codes = [a.codice for a in audit_revisione(rev, time_cli=21)]
        self.assertIn("RISPOSTA_SENZA_RICEZIONE", codes)
        self.assertIn("DISPATCH_SENZA_RIENTRO_PREV", codes)

    def test_inviato_senza_dispatch(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=4,
            int_status="inviato_al_cliente",
        )
        codes = [a.codice for a in audit_revisione(rev)]
        self.assertEqual(codes, ["INVIATO_SENZA_DISPATCH"])

    def test_empty_after_approved_not_anomalous(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date="2025-01-01",
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        rev1 = Revisione.objects.create(documento=self.doc, rev_no=1)
        self.assertEqual(audit_revisione(rev1), [])
        self.assertEqual(audit_commessa("25012"), [])

    def test_post_ricezione_pending_emission_not_anomalous(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date="2025-01-01",
            rec_act_date="2025-01-10",
            ext_status=self.rejected,
            int_status="ricevuto",
        )
        rev1 = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            dis_plan_date="2025-01-24",
        )
        self.assertEqual(audit_revisione(rev1), [])
        self.assertEqual(audit_commessa("25012"), [])

    def test_int_status_incongruente_ricevuto_senza_data(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="ricevuto",
        )
        codes = [a.codice for a in audit_revisione(rev)]
        self.assertIn("INT_STATUS_INCONGRUENTE", codes)

    def test_classifica_conclusi_fixes_risposta_senza_ricezione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        result = classifica_revisione(
            "25012",
            rev.pk,
            "conclusi",
            {"rec_act_date": "2026-03-10"},
        )
        self.assertEqual(result["revisione"]["int_status"], "ricevuto")
        self.assertEqual(result["revisione"]["rec_act_date"], "2026-03-10")
        rev.refresh_from_db()
        self.assertEqual(audit_revisione(rev, time_cli=21), [])

    def test_classifica_da_emettere_clears_inviato_senza_dispatch(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=4,
            int_status="inviato_al_cliente",
        )
        classifica_revisione("25012", rev.pk, "da_emettere")
        rev.refresh_from_db()
        self.assertEqual(rev.int_status, "")
        self.assertEqual(audit_revisione(rev), [])

    def test_gruppi_marks_latest_revision(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="ricevuto",
            rec_act_date="2025-01-10",
            dis_act_date="2025-01-01",
            ext_status=self.approved,
        )
        rev1 = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            int_status="inviato_al_cliente",
        )
        gruppi = serialize_anomalie_gruppi("25012", audit_commessa("25012"))
        self.assertEqual(len(gruppi), 1)
        self.assertEqual(gruppi[0]["revisione_id"], rev1.pk)
        self.assertTrue(gruppi[0]["is_latest"])
        self.assertEqual(gruppi[0]["suggest_bucket"], "da_ricevere")
        self.assertEqual(gruppi[0]["situazione_attuale"]["label"], "Da ricevere")

    def test_ignora_anomalie_esclude_revisione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        self.assertEqual(len(audit_commessa("25012")), 2)
        ignora_anomalie_revisione("25012", rev.pk)
        rev.refresh_from_db()
        self.assertTrue(rev.ignora_anomalie)
        self.assertEqual(audit_commessa("25012"), [])
        self.assertEqual(audit_commessa_summary("25012")["count"], 0)


def _costruisci_xlsm_organizzazione(percorso, righe):
    """Un file Excel "Organizzazione Commesse" minimo, per i test.

    righe: lista di dict con chiavi Job/PM/PE/WE/QCI (e opzionalmente altre
    colonne del foglio reale, ignorate se assenti).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commesse"
    ws.append(["", "", "", "PROJECT TEAM", "", "", "", "", "", "", ""])
    ws.append(
        [
            "Job no.",
            "DWG / ITEM",
            "Client ",
            "PM",
            "PE",
            "WE",
            "QCI",
            "Plant",
            "Delivery",
            "Note",
            "Closed",
        ]
    )
    for riga in righe:
        ws.append(
            [
                riga.get("Job"),
                riga.get("DWG"),
                riga.get("Client"),
                riga.get("PM"),
                riga.get("PE"),
                riga.get("WE"),
                riga.get("QCI"),
                riga.get("Plant"),
                riga.get("Delivery"),
                riga.get("Note"),
                riga.get("Closed"),
            ]
        )
    wb.save(percorso)


class ImportOldTests(TestCase):
    def _frames(self):
        return {
            "testata": pd.DataFrame(
                [
                    {
                        "Job": "99999",
                        "Client": "Cliente Test",
                        "POno": "PO-1",
                        "JobDetail": "Dettaglio",
                        "DeliveryDate": None,
                        "DeliveryTerm": "",
                        "Requisition": "",
                        "TimeCliDocRev": None,
                        "TimeVenDocRev": None,
                        "RevLetFlag": False,
                    }
                ]
            ),
            "indirsped": pd.DataFrame(),
            "dettaglio": pd.DataFrame(
                [
                    {
                        "ItemNo": "1",
                        "VendorDoc": "99999-01",
                        "ClientDocNo": "",
                        "ContractorDocNo": "CTR-99",
                        "ClientDocClass": "",
                        "DocTitle": "Titolo",
                        "DocPenalty": False,
                        "DocPayment": False,
                        "RevGen": False,
                        "Reparto": 4,
                        "Remarks": "",
                    }
                ]
            ),
            "revisioni": pd.DataFrame(
                [
                    {
                        "VendorDoc": "99999-01",
                        "RevNo": 0,
                        "RevLet": "",
                        "DisPlanDate": None,
                        "DisActDate": "2025-01-01",
                        "RecPlanDate": None,
                        "RecActDate": "2025-01-10",
                        "Status": "A",
                    },
                    {
                        "VendorDoc": "99999-01",
                        "RevNo": 1,
                        "RevLet": "",
                        "DisPlanDate": None,
                        "DisActDate": None,
                        "RecPlanDate": None,
                        "RecActDate": None,
                        "Status": None,
                    },
                ]
            ),
        }

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_includes_all_revisioni_and_reports_orphans(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        result = importa_commessa_da_access("99999")
        self.assertEqual(result["documenti"], 1)
        self.assertEqual(result["revisioni"], 2)
        self.assertEqual(result["revisioni_orfane_escluse"], 0)
        self.assertEqual(Revisione.objects.count(), 2)
        self.assertIn("anomalie", result)
        # Rev 0 has DisActDate + RecActDate → ricevuto; rev 1 has no dates → Da inviare
        self.assertEqual(Revisione.objects.filter(rev_no=0).get().int_status, "ricevuto")
        self.assertEqual(Revisione.objects.filter(rev_no=1).get().int_status, "")
        doc = Documento.objects.get(vendor_doc="99999-01")
        self.assertEqual(doc.contractor_doc_no, "CTR-99")

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_sets_ricevuto_when_client_answered_without_dates(self, mock_fetch):
        StatoEsterno.objects.create(nome="Approved", colore="#00B050", crea_nuova_rev=False)
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "A",
                    "DisPlanDate": None,
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": "A",
                }
            ]
        )
        mock_fetch.return_value = frames
        importa_commessa_da_access("99999")
        rev = Revisione.objects.get()
        self.assertEqual(rev.int_status, "ricevuto")
        self.assertIsNotNone(rev.ext_status)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_excludes_orphan_after_approved(self, mock_fetch):
        StatoEsterno.objects.create(nome="Approved", colore="#00B050", crea_nuova_rev=False)
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "",
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 1,
                    "RevLet": "",
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        mock_fetch.return_value = frames
        result = importa_commessa_da_access("99999")
        self.assertEqual(result["revisioni_orfane_escluse"], 1)
        self.assertEqual(result["revisioni"], 1)
        self.assertEqual(Revisione.objects.count(), 1)
        self.assertEqual(Revisione.objects.get().rev_no, 0)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_sets_inviato_when_only_dis_act_date(self, mock_fetch):
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "A",
                    "DisPlanDate": None,
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                }
            ]
        )
        mock_fetch.return_value = frames
        importa_commessa_da_access("99999")
        rev = Revisione.objects.get()
        self.assertEqual(rev.int_status, "inviato_al_cliente")
        self.assertEqual(rev.dis_act_date.isoformat(), "2025-01-01")
        self.assertIsNone(rev.rec_act_date)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_rejects_existing_job(self, mock_fetch):
        Testata.objects.create(job="99999")
        with self.assertRaises(ValueError):
            importa_commessa_da_access("99999")
        mock_fetch.assert_not_called()

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_legge_persone_da_excel(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        xlsm_path = Path(media.name) / "organizzazione.xlsx"
        _costruisci_xlsm_organizzazione(
            xlsm_path,
            [{"Job": "99999", "PM": "ROSSI", "PE": "BIANCHI/VERDI", "WE": None, "QCI": "N/A"}],
        )
        # ROSSI ha un utente registrato con quel cognome: va abbinato. BIANCHI
        # e VERDI no: restano testo libero, da risolvere a mano.
        rossi = User.objects.create_user(
            "rossi_test", "rossi@b.it", "pw", first_name="Mario", last_name="Rossi"
        )

        with override_settings(ORGANIZZAZIONE_COMMESSE_XLSM_PATH=str(xlsm_path)):
            with self.captureOnCommitCallbacks(execute=True):
                importa_commessa_da_access("99999")

        testata = Testata.objects.get(job="99999")
        pm = testata.persone.get(ruolo="pm")
        self.assertEqual(pm.utente_id, rossi.pk)
        self.assertEqual(pm.nome_libero, "")

        liberi = testata.persone.filter(ruolo="pe")
        self.assertEqual(sorted(liberi.values_list("nome_libero", flat=True)), ["BIANCHI", "VERDI"])
        self.assertTrue(all(p.utente_id is None for p in liberi))

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_file_excel_assente_non_blocca_import(self, mock_fetch):
        mock_fetch.return_value = self._frames()

        with override_settings(
            ORGANIZZAZIONE_COMMESSE_XLSM_PATH=r"Z:\percorso\che\non\esiste.xlsm"
        ):
            with self.captureOnCommitCallbacks(execute=True):
                result = importa_commessa_da_access("99999")

        self.assertEqual(result["documenti"], 1)
        testata = Testata.objects.get(job="99999")
        self.assertEqual(testata.persone.count(), 0)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_risolve_sito_costruttivo_da_bc(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        fake = _FakeBusinessCentralSito({"99999": "1"})

        with patch("core.services.bc_sync._apri_connessione", return_value=fake):
            with self.captureOnCommitCallbacks(execute=True):
                importa_commessa_da_access("99999")

        testata = Testata.objects.get(job="99999")
        self.assertEqual(testata.sito_costruttivo_id, bg.pk)


class PersonaCommessaTests(TestCase):
    """Vincoli del modello: utente XOR nome libero, più persone per ruolo."""

    def setUp(self):
        self.testata = Testata.objects.create(job="88001")
        self.utente = User.objects.create_user("pc_utente", "pc-utente@b.it", "pw")

    def test_richiede_utente_o_nome_libero(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            PersonaCommessa.objects.create(testata=self.testata, ruolo=RuoloPersonaCommessa.PM)

    def test_rifiuta_utente_e_nome_libero_insieme(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            PersonaCommessa.objects.create(
                testata=self.testata,
                ruolo=RuoloPersonaCommessa.PM,
                utente=self.utente,
                nome_libero="Mario Rossi",
            )

    def test_piu_persone_stesso_ruolo(self):
        altro = User.objects.create_user("pc_altro", "pc-altro@b.it", "pw")
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PE, utente=self.utente
        )
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PE, utente=altro
        )

        self.assertEqual(
            PersonaCommessa.objects.filter(
                testata=self.testata, ruolo=RuoloPersonaCommessa.PE
            ).count(),
            2,
        )

    def test_cascata_alla_cancellazione_testata(self):
        PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.WE, utente=self.utente
        )

        self.testata.delete()

        self.assertEqual(PersonaCommessa.objects.count(), 0)

    def test_nome_visualizzato(self):
        con_utente = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.QCI, utente=self.utente
        )
        libero = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.QCI, nome_libero="Anna Verdi"
        )

        self.assertEqual(con_utente.nome_visualizzato, self.utente.nome_completo)
        self.assertEqual(libero.nome_visualizzato, "Anna Verdi")

    def test_nome_visualizzato_mostra_nome_e_cognome_non_lo_username(self):
        utente = User.objects.create_user(
            "pc_nome_cognome", "pc-nc@b.it", "pw", first_name="Mario", last_name="Rossi"
        )
        persona = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, utente=utente
        )

        self.assertEqual(persona.nome_visualizzato, "Mario Rossi")


class DividiNomiTests(SimpleTestCase):
    """_dividi_nomi: split di una cella grezza del foglio in singoli nomi."""

    def test_nome_singolo_con_spazio_finale(self):
        self.assertEqual(_dividi_nomi("CAPPELLOTTO "), ["CAPPELLOTTO"])

    def test_split_su_slash(self):
        self.assertEqual(_dividi_nomi("LUCINI / PILONI"), ["LUCINI", "PILONI"])

    def test_split_su_a_capo(self):
        self.assertEqual(_dividi_nomi("LUCINI\nPILONI"), ["LUCINI", "PILONI"])

    def test_scarta_parti_vuote(self):
        self.assertEqual(_dividi_nomi("LUCINI//PILONI"), ["LUCINI", "PILONI"])

    def test_scarta_n_a_case_insensitive(self):
        self.assertEqual(_dividi_nomi("N/A"), [])
        self.assertEqual(_dividi_nomi("n/a"), [])

    def test_valore_vuoto(self):
        self.assertEqual(_dividi_nomi(None), [])
        self.assertEqual(_dividi_nomi(""), [])

    def test_caso_sporco_non_solleva_eccezioni(self):
        # Limite noto: non c'è un modo affidabile di interpretare questo
        # caso, importa che non sollevi eccezioni.
        risultato = _dividi_nomi("BG: IMPALLOMENI \nVI: ")
        self.assertIsInstance(risultato, list)


class NormalizzaJobTests(SimpleTestCase):
    """_normalizza_job: normalizza una cella 'Job no.' al formato Testata.job."""

    def test_stringa(self):
        self.assertEqual(_normalizza_job("22105"), "22105")

    def test_intero(self):
        self.assertEqual(_normalizza_job(22116), "22116")

    def test_float_intero(self):
        self.assertEqual(_normalizza_job(22116.0), "22116")

    def test_prefisso_lettera(self):
        self.assertEqual(_normalizza_job("I23004"), "I23004")

    def test_none(self):
        self.assertEqual(_normalizza_job(None), "")


class LeggiOrganizzazioneCommesseTests(TestCase):
    """leggi_organizzazione_commesse / persone_per_job: lettura del foglio Excel."""

    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        self.xlsm_path = Path(media.name) / "organizzazione.xlsx"

    def _leggi(self, righe):
        _costruisci_xlsm_organizzazione(self.xlsm_path, righe)
        return leggi_organizzazione_commesse(str(self.xlsm_path))

    def test_job_stringa_e_intero(self):
        dati = self._leggi(
            [
                {"Job": "22105", "PM": "CAPPELLOTTO"},
                {"Job": 22116, "PM": "ROSSI"},
            ]
        )
        self.assertEqual(dati["22105"]["pm"], ["CAPPELLOTTO"])
        self.assertEqual(dati["22116"]["pm"], ["ROSSI"])

    def test_job_con_prefisso_lettera(self):
        dati = self._leggi([{"Job": "I23004", "PM": "N/A"}])
        self.assertIn("I23004", dati)
        self.assertEqual(dati["I23004"]["pm"], [])

    def test_ruolo_multiplo_split_su_slash(self):
        dati = self._leggi([{"Job": "23068", "PM": "BORELLI/QUIPPERETTI"}])
        self.assertEqual(dati["23068"]["pm"], ["BORELLI", "QUIPPERETTI"])

    def test_n_a_trattato_come_vuoto(self):
        dati = self._leggi([{"Job": "99001", "QCI": "N/A"}])
        self.assertEqual(dati["99001"]["qci"], [])

    def test_ruolo_vuoto(self):
        dati = self._leggi([{"Job": "99002"}])
        self.assertEqual(dati["99002"], {"pm": [], "pe": [], "we": [], "qci": []})

    def test_righe_con_lo_stesso_job_si_uniscono(self):
        dati = self._leggi(
            [
                {"Job": "23068", "PM": "BORELLI/QUIPPERETTI"},
                {"Job": "23068", "PM": "BORELLI"},
            ]
        )
        # Deduplicato: BORELLI compare una sola volta anche se in entrambe le righe.
        self.assertEqual(dati["23068"]["pm"], ["BORELLI", "QUIPPERETTI"])

    def test_persone_per_job_non_trovato_restituisce_none(self):
        self._leggi([{"Job": "99003", "PM": "ROSSI"}])
        self.assertIsNone(persone_per_job("00000", str(self.xlsm_path)))

    def test_persone_per_job_trovato(self):
        self._leggi([{"Job": "99003", "PM": "ROSSI"}])
        self.assertEqual(persone_per_job("99003", str(self.xlsm_path))["pm"], ["ROSSI"])

    def test_file_non_trovato(self):
        with self.assertRaises(FileNotFoundError):
            leggi_organizzazione_commesse(r"Z:\percorso\che\non\esiste.xlsm")


class CreateCommessaConPersoneTests(TestCase):
    """create_commessa: creazione atomica di Testata + PersonaCommessa."""

    def setUp(self):
        self.utente = User.objects.create_user("cc_utente", "cc-utente@b.it", "pw")

    def _dati(self, job, **extra):
        return {"job": job, "client": "Cliente Test", **extra}

    def test_utente_registrato(self):
        t = create_commessa(self._dati("77001", persone={"pm": [{"utente_id": self.utente.pk}]}))

        riga = PersonaCommessa.objects.get(testata=t, ruolo=RuoloPersonaCommessa.PM)
        self.assertEqual(riga.utente_id, self.utente.pk)
        self.assertEqual(riga.nome_libero, "")

    def test_nome_libero(self):
        t = create_commessa(self._dati("77002", persone={"pe": [{"nome": "Mario Rossi"}]}))

        riga = PersonaCommessa.objects.get(testata=t, ruolo=RuoloPersonaCommessa.PE)
        self.assertIsNone(riga.utente_id)
        self.assertEqual(riga.nome_libero, "Mario Rossi")

    def test_piu_persone_sullo_stesso_ruolo(self):
        altro = User.objects.create_user("cc_altro", "cc-altro@b.it", "pw")
        t = create_commessa(
            self._dati(
                "77003",
                persone={
                    "qci": [{"utente_id": self.utente.pk}, {"utente_id": altro.pk}],
                },
            )
        )

        self.assertEqual(
            PersonaCommessa.objects.filter(testata=t, ruolo=RuoloPersonaCommessa.QCI).count(), 2
        )

    def test_senza_persone_non_crea_righe(self):
        t = create_commessa(self._dati("77004"))

        self.assertEqual(PersonaCommessa.objects.filter(testata=t).count(), 0)

    def test_nome_libero_vuoto_ignorato(self):
        t = create_commessa(self._dati("77005", persone={"we": [{"nome": "   "}]}))

        self.assertEqual(PersonaCommessa.objects.filter(testata=t).count(), 0)

    def test_utente_id_inesistente_solleva_integrity_error(self):
        from django.db import connection

        # Postgres non verifica sempre il vincolo di chiave esterna in modo
        # sincrono dentro un savepoint annidato (quello di TestCase): un
        # check_constraints() esplicito forza a farlo qui, invece di
        # scoprirlo solo al rollback di fine test.
        with self.assertRaises(IntegrityError), transaction.atomic():
            create_commessa(self._dati("77006", persone={"pm": [{"utente_id": 999999}]}))
            connection.check_constraints()

    def test_risolve_sito_costruttivo_da_bc_alla_creazione(self):
        bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        fake = _FakeBusinessCentralSito({"77007": "1"})

        with patch("core.services.bc_sync._apri_connessione", return_value=fake):
            with self.captureOnCommitCallbacks(execute=True):
                t = create_commessa(self._dati("77007"))

        t.refresh_from_db()
        self.assertEqual(t.sito_costruttivo_id, bg.pk)

    def test_bc_non_raggiungibile_non_impedisce_la_creazione(self):
        fake = _FakeBusinessCentralSito({}, conn=False)

        with patch("core.services.bc_sync._apri_connessione", return_value=fake):
            with self.captureOnCommitCallbacks(execute=True):
                t = create_commessa(self._dati("77008"))

        self.assertTrue(Testata.objects.filter(job="77008").exists())
        t.refresh_from_db()
        self.assertIsNone(t.sito_costruttivo_id)


class PersonePerRuoloTests(TestCase):
    """persone_per_ruolo: persone per ruolo, con lo stato di abbinamento, per l'header."""

    def test_elenca_le_persone_per_ruolo_con_lo_stato_di_abbinamento(self):
        t = Testata.objects.create(job="66001")
        u1 = User.objects.create_user(
            "ppr1", "ppr1@b.it", "pw", first_name="Mario", last_name="Rossi"
        )
        abbinata = PersonaCommessa.objects.create(
            testata=t, ruolo=RuoloPersonaCommessa.PM, utente=u1
        )
        libera = PersonaCommessa.objects.create(
            testata=t, ruolo=RuoloPersonaCommessa.PM, nome_libero="Libero Bianchi"
        )

        risultato = persone_per_ruolo(t)

        self.assertEqual(
            risultato["pm"],
            [
                {"id": abbinata.pk, "nome": "Mario Rossi", "abbinato": True},
                {"id": libera.pk, "nome": "Libero Bianchi", "abbinato": False},
            ],
        )
        self.assertEqual(risultato["pe"], [])
        self.assertEqual(set(risultato), {"pm", "pe", "qci", "we"})


class RisolviPersonaCommessaTests(TestCase):
    """risolvi_persona_commessa: collega manualmente una voce a testo libero a un utente."""

    def setUp(self):
        self.testata = Testata.objects.create(job="66010")
        self.persona = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.WE, nome_libero="Baldelli"
        )
        self.utente = User.objects.create_user("rpc1", "rpc1@b.it", "pw", last_name="Baldelli")

    def test_collega_l_utente_e_svuota_il_nome_libero(self):
        persona = risolvi_persona_commessa(self.persona.pk, self.utente.pk)

        self.assertEqual(persona.utente_id, self.utente.pk)
        self.assertEqual(persona.nome_libero, "")

    def test_persona_inesistente(self):
        with self.assertRaises(PersonaCommessa.DoesNotExist):
            risolvi_persona_commessa(999999, self.utente.pk)

    def test_utente_inesistente(self):
        with self.assertRaises(User.DoesNotExist):
            risolvi_persona_commessa(self.persona.pk, 999999)


class BackfillPersoneCommessaTests(TestCase):
    """backfill_persone_commessa e il comando che lo espone."""

    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        self.xlsm_path = Path(media.name) / "organizzazione.xlsx"
        _costruisci_xlsm_organizzazione(
            self.xlsm_path,
            [
                {"Job": "55001", "PM": "ROSSI", "PE": "BIANCHI"},
                {"Job": "55002", "PM": "VERDI"},
            ],
        )
        patcher = override_settings(ORGANIZZAZIONE_COMMESSE_XLSM_PATH=str(self.xlsm_path))
        patcher.enable()
        self.addCleanup(patcher.disable)

        self.t1 = Testata.objects.create(job="55001")
        self.t2 = Testata.objects.create(job="55002")
        self.t3 = Testata.objects.create(job="55003")  # non nel foglio

    def test_crea_persone_per_commesse_esistenti(self):
        report = backfill_persone_commessa()

        self.assertEqual(sorted(report["aggiornate"]), ["55001", "55002"])
        self.assertEqual(report["non_trovate"], ["55003"])
        self.assertEqual(
            set(self.t1.persone.values_list("ruolo", "nome_libero")),
            {("pm", "ROSSI"), ("pe", "BIANCHI")},
        )

    def test_salta_ruolo_gia_popolato(self):
        PersonaCommessa.objects.create(
            testata=self.t1, ruolo=RuoloPersonaCommessa.PM, nome_libero="Già corretto"
        )

        report = backfill_persone_commessa()

        # PM non toccato (era già popolato); PE viene comunque riempito.
        pm = list(self.t1.persone.filter(ruolo="pm").values_list("nome_libero", flat=True))
        self.assertEqual(pm, ["Già corretto"])
        self.assertTrue(self.t1.persone.filter(ruolo="pe", nome_libero="BIANCHI").exists())
        self.assertEqual(report["ruoli_saltati"], 1)

    def test_rieseguibile_senza_duplicare(self):
        backfill_persone_commessa()
        prima = PersonaCommessa.objects.count()

        backfill_persone_commessa()

        self.assertEqual(PersonaCommessa.objects.count(), prima)

    def test_dry_run_non_scrive(self):
        report = backfill_persone_commessa(dry_run=True)

        self.assertEqual(PersonaCommessa.objects.count(), 0)
        self.assertEqual(sorted(report["aggiornate"]), ["55001", "55002"])

    def test_jobs_limita_il_backfill(self):
        backfill_persone_commessa(jobs=["55001"])

        self.assertTrue(self.t1.persone.exists())
        self.assertFalse(self.t2.persone.exists())

    def test_comando_job_singolo(self):
        out = io.StringIO()
        call_command("backfill_persone_commessa", "--job=55001", stdout=out)

        self.assertTrue(self.t1.persone.exists())
        self.assertFalse(self.t2.persone.exists())
        self.assertIn("55001", out.getvalue())

    def test_comando_job_non_trovato_in_workflow(self):
        err = io.StringIO()
        call_command("backfill_persone_commessa", "--job=00000", stderr=err)

        self.assertIn("non trovata", err.getvalue())

    def test_comando_dry_run(self):
        out = io.StringIO()
        call_command("backfill_persone_commessa", "--dry-run", stdout=out)

        self.assertEqual(PersonaCommessa.objects.count(), 0)
        self.assertIn("dry-run", out.getvalue())

    def test_abbina_un_cognome_a_un_utente_registrato(self):
        rossi = User.objects.create_user(
            "bpc_rossi", "bpc-rossi@b.it", "pw", first_name="Mario", last_name="Rossi"
        )

        backfill_persone_commessa()

        pm = self.t1.persone.get(ruolo="pm")
        self.assertEqual(pm.utente_id, rossi.pk)
        self.assertEqual(pm.nome_libero, "")
        # BIANCHI non ha un utente corrispondente: resta testo libero.
        pe = self.t1.persone.get(ruolo="pe")
        self.assertIsNone(pe.utente_id)
        self.assertEqual(pe.nome_libero, "BIANCHI")

    def test_da_risolvere_elenca_i_cognomi_non_abbinati(self):
        report = backfill_persone_commessa()

        self.assertEqual(
            {(v["job"], v["ruolo"], v["nome"]) for v in report["da_risolvere"]},
            {("55001", "pm", "ROSSI"), ("55001", "pe", "BIANCHI"), ("55002", "pm", "VERDI")},
        )

    def test_comando_segnala_da_risolvere(self):
        out = io.StringIO()
        call_command("backfill_persone_commessa", stdout=out)

        self.assertIn("Da risolvere", out.getvalue())
        self.assertIn("ROSSI", out.getvalue())

    def test_comando_risolvi_esistenti(self):
        backfill_persone_commessa()  # crea le righe a testo libero
        rossi = User.objects.create_user(
            "bpc_rossi2", "bpc-rossi2@b.it", "pw", first_name="Mario", last_name="Rossi"
        )

        out = io.StringIO()
        call_command("backfill_persone_commessa", "--risolvi-esistenti", stdout=out)

        pm = self.t1.persone.get(ruolo="pm")
        self.assertEqual(pm.utente_id, rossi.pk)
        self.assertIn("ROSSI", out.getvalue())


class TrovaUtentePerCognomeTests(TestCase):
    """trova_utente_per_cognome: abbinamento cognome -> utente registrato."""

    def test_una_sola_corrispondenza(self):
        u = User.objects.create_user("tup1", "tup1@b.it", "pw", last_name="Rossi")

        self.assertEqual(trova_utente_per_cognome("Rossi"), u)

    def test_case_insensitive(self):
        u = User.objects.create_user("tup2", "tup2@b.it", "pw", last_name="Rossi")

        self.assertEqual(trova_utente_per_cognome("ROSSI"), u)

    def test_nessuna_corrispondenza(self):
        self.assertIsNone(trova_utente_per_cognome("Sconosciuto"))

    def test_piu_corrispondenze_troppo_ambiguo(self):
        User.objects.create_user("tup3", "tup3@b.it", "pw", last_name="Rossi")
        User.objects.create_user("tup4", "tup4@b.it", "pw", last_name="Rossi")

        self.assertIsNone(trova_utente_per_cognome("Rossi"))


class RisolviPersoneLibereTests(TestCase):
    """risolvi_persone_libere: ri-abbina le PersonaCommessa già a testo libero."""

    def setUp(self):
        self.testata = Testata.objects.create(job="55010")
        self.non_abbinata = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PM, nome_libero="Rossi"
        )
        self.senza_utente = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.PE, nome_libero="Sconosciuto"
        )

    def test_risolve_quando_un_utente_ora_esiste(self):
        rossi = User.objects.create_user("rpl1", "rpl1@b.it", "pw", last_name="Rossi")

        report = risolvi_persone_libere()

        self.non_abbinata.refresh_from_db()
        self.assertEqual(self.non_abbinata.utente_id, rossi.pk)
        self.assertEqual(self.non_abbinata.nome_libero, "")
        self.assertEqual(
            {(v["job"], v["ruolo"], v["nome"]) for v in report["risolte"]},
            {("55010", "pm", "Rossi")},
        )

    def test_lascia_intatte_quelle_ancora_senza_corrispondenza(self):
        User.objects.create_user("rpl2", "rpl2@b.it", "pw", last_name="Rossi")

        report = risolvi_persone_libere()

        self.senza_utente.refresh_from_db()
        self.assertIsNone(self.senza_utente.utente_id)
        self.assertEqual(self.senza_utente.nome_libero, "Sconosciuto")
        self.assertEqual(
            {(v["job"], v["ruolo"], v["nome"]) for v in report["non_risolte"]},
            {("55010", "pe", "Sconosciuto")},
        )

    def test_dry_run_non_scrive(self):
        User.objects.create_user("rpl3", "rpl3@b.it", "pw", last_name="Rossi")

        risolvi_persone_libere(dry_run=True)

        self.non_abbinata.refresh_from_db()
        self.assertIsNone(self.non_abbinata.utente_id)
        self.assertEqual(self.non_abbinata.nome_libero, "Rossi")

    def test_jobs_limita_la_ricerca(self):
        User.objects.create_user("rpl4", "rpl4@b.it", "pw", last_name="Rossi")
        altra_testata = Testata.objects.create(job="55011")
        PersonaCommessa.objects.create(
            testata=altra_testata, ruolo=RuoloPersonaCommessa.PM, nome_libero="Rossi"
        )

        risolvi_persone_libere(jobs=["55011"])

        self.non_abbinata.refresh_from_db()
        self.assertIsNone(self.non_abbinata.utente_id)
        altra = PersonaCommessa.objects.get(testata=altra_testata)
        self.assertIsNotNone(altra.utente_id)


class UtentiCercaApiTests(TestCase):
    """GET /api/utenti/cerca/: autocomplete per PM/PE/QCI/WE."""

    def setUp(self):
        self.client = Client()
        self.utente = User.objects.create_user(
            "uca_utente",
            "uca@b.it",
            "pw",
            permesso=Permesso.WRITING,
            first_name="Mario",
            last_name="Rossi",
        )
        self.client.force_login(self.utente)
        User.objects.create_user(
            "uca_bianchi",
            "uca-bianchi@b.it",
            "pw",
            first_name="Anna",
            last_name="Bianchi",
        )
        User.objects.create_user(
            "uca_inattivo",
            "uca-inattivo@b.it",
            "pw",
            first_name="Fuori",
            last_name="Servizio",
            is_active=False,
        )

    def test_richiede_login(self):
        self.client.logout()
        risposta = self.client.get("/api/utenti/cerca/?q=Rossi")
        self.assertNotEqual(risposta.status_code, 200)

    def test_cerca_per_cognome(self):
        risposta = self.client.get("/api/utenti/cerca/?q=Rossi")
        utenti = risposta.json()["utenti"]
        self.assertEqual([u["username"] for u in utenti], ["uca_utente"])

    def test_cerca_per_username(self):
        risposta = self.client.get("/api/utenti/cerca/?q=uca_bianchi")
        utenti = risposta.json()["utenti"]
        self.assertEqual([u["username"] for u in utenti], ["uca_bianchi"])

    def test_meno_di_due_caratteri_restituisce_vuoto(self):
        risposta = self.client.get("/api/utenti/cerca/?q=R")
        self.assertEqual(risposta.json()["utenti"], [])

    def test_esclude_utenti_non_attivi(self):
        risposta = self.client.get("/api/utenti/cerca/?q=Servizio")
        self.assertEqual(risposta.json()["utenti"], [])

    def test_limita_a_dieci_risultati(self):
        for i in range(15):
            User.objects.create_user(f"uca_molti{i}", f"uca-molti{i}@b.it", "pw", last_name="Molti")
        risposta = self.client.get("/api/utenti/cerca/?q=Molti")
        self.assertEqual(len(risposta.json()["utenti"]), 10)


class PersonaCommessaRisolviApiTests(TestCase):
    """POST /api/persone-commessa/<pk>/risolvi/: collega a mano un nome libero a un utente."""

    def setUp(self):
        self.client = Client()
        self.utente_scrittura = User.objects.create_user(
            "pcr_scrittura", "pcr-scrittura@b.it", "pw", permesso=Permesso.WRITING
        )
        self.utente_lettura = User.objects.create_user(
            "pcr_lettura", "pcr-lettura@b.it", "pw", permesso=Permesso.READING
        )
        self.candidato = User.objects.create_user(
            "pcr_candidato", "pcr-candidato@b.it", "pw", last_name="Baldelli"
        )
        self.testata = Testata.objects.create(job="66020")
        self.persona = PersonaCommessa.objects.create(
            testata=self.testata, ruolo=RuoloPersonaCommessa.WE, nome_libero="Baldelli"
        )

    def _risolvi(self, persona_id, utente_id):
        return self.client.post(
            f"/api/persone-commessa/{persona_id}/risolvi/",
            data=json.dumps({"utente_id": utente_id}),
            content_type="application/json",
        )

    def test_richiede_login(self):
        risposta = self._risolvi(self.persona.pk, self.candidato.pk)
        self.assertNotEqual(risposta.status_code, 200)

    def test_richiede_permesso_di_scrittura(self):
        self.client.force_login(self.utente_lettura)

        risposta = self._risolvi(self.persona.pk, self.candidato.pk)

        self.assertEqual(risposta.status_code, 403)
        self.persona.refresh_from_db()
        self.assertIsNone(self.persona.utente_id)

    def test_collega_l_utente(self):
        self.client.force_login(self.utente_scrittura)

        risposta = self._risolvi(self.persona.pk, self.candidato.pk)

        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(risposta.json()["nome"], self.candidato.nome_completo)
        self.persona.refresh_from_db()
        self.assertEqual(self.persona.utente_id, self.candidato.pk)
        self.assertEqual(self.persona.nome_libero, "")

    def test_persona_inesistente(self):
        self.client.force_login(self.utente_scrittura)

        risposta = self._risolvi(999999, self.candidato.pk)

        self.assertEqual(risposta.status_code, 404)

    def test_utente_inesistente(self):
        self.client.force_login(self.utente_scrittura)

        risposta = self._risolvi(self.persona.pk, 999999)

        self.assertEqual(risposta.status_code, 404)

    def test_utente_id_mancante(self):
        self.client.force_login(self.utente_scrittura)

        risposta = self.client.post(
            f"/api/persone-commessa/{self.persona.pk}/risolvi/",
            data=json.dumps({}),
            content_type="application/json",
        )

        self.assertEqual(risposta.status_code, 400)


class SituazioneApiTestCase(TestCase):
    """Tests for batch situazione/documenti API endpoints."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "situazione_user",
            "situazione@brembanarolle.com",
            "pw",
            permesso=Permesso.READING,
        )
        self.client.force_login(self.user)
        self.fixtures = _make_fixtures_mock()
        self.job = self.fixtures["testata"].job

    def test_situazione_api_returns_documenti_and_revisioni(self):
        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["documenti"]), 2)
        revs_by_doc = data["revisioni_by_doc"]
        doc_qmdbi_id = str(self.fixtures["doc_qmdbi"].pk)
        doc_qcpa_id = str(self.fixtures["doc_qcpa"].pk)
        self.assertEqual(len(revs_by_doc[doc_qmdbi_id]), 1)
        self.assertEqual(len(revs_by_doc[doc_qcpa_id]), 3)

    def test_documenti_api_include_revisioni(self):
        response = self.client.get(f"/api/commesse/{self.job}/documenti/?include_revisioni=1")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["documenti"]), 2)
        doc_qcpa_id = str(self.fixtures["doc_qcpa"].pk)
        self.assertEqual(len(data["revisioni_by_doc"][doc_qcpa_id]), 3)

    def test_situazione_api_requires_auth(self):
        self.client.logout()
        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 401)

    def test_situazione_api_not_found(self):
        response = self.client.get("/api/commesse/INEXISTENT/situazione/")
        self.assertEqual(response.status_code, 404)

    def test_export_situazione_pdf(self):
        from datetime import date

        import fitz

        today = date.today()
        date_part = f"{today.day:02d}_{today.month:02d}_{today.year}"

        response = self.client.get(f"/api/commesse/{self.job}/situazione/export/?vista=orizzontale")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn(
            f'filename="situazione_documenti_orizzontale_{date_part}.pdf"',
            response["Content-Disposition"],
        )
        self.assertTrue(response.content.startswith(b"%PDF"))

        text = fitz.open(stream=response.content, filetype="pdf")[0].get_text()
        self.assertIn("Vendor Document", text)
        self.assertIn("Planning", text)
        self.assertIn("Dispatch", text)
        self.assertIn("25056-QMDBI", text)
        self.assertNotIn("Client Doc No", text)
        self.assertNotIn("Penalty", text)

        response_v = self.client.get(f"/api/commesse/{self.job}/situazione/export/?vista=verticale")
        self.assertEqual(response_v.status_code, 200)
        self.assertIn(
            f'filename="situazione_documenti_verticale_{date_part}.pdf"',
            response_v["Content-Disposition"],
        )
        text_v = fitz.open(stream=response_v.content, filetype="pdf")[0].get_text()
        self.assertIn("Vendor Document", text_v)

    def test_export_situazione_pdf_hides_empty_fixed_columns(self):
        from src.pdf import _active_fixed_doc_columns

        documenti = [
            {"vendor_doc": "DOC-1", "doc_title": "", "client_doc_no": ""},
            {"vendor_doc": "DOC-2", "doc_penalty": True, "item_no": "42"},
        ]
        active = _active_fixed_doc_columns(documenti)
        labels = [col["label"] for col in active]
        self.assertEqual(labels, ["Vendor Document", "Item", "Penalty"])

    def test_export_situazione_not_found(self):
        response = self.client.get("/api/commesse/INEXISTENT/situazione/export/")
        self.assertEqual(response.status_code, 404)

    def test_situazione_api_includes_ext_status_lettera(self):
        approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        doc = self.fixtures["doc_qmdbi"]
        rev = doc.revisioni.order_by("rev_no").first()
        rev.ext_status = approved
        rev.save(update_fields=["ext_status"])

        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        revs = response.json()["revisioni_by_doc"][str(doc.pk)]
        self.assertEqual(revs[0]["ext_status_lettera"], "A")
        self.assertEqual(revs[0]["ext_status_label"], "Approved")


class StatoEsternoLetteraTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "stato_lettera_user",
            "stato_lettera@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)

    def test_letter_for_status_name_aliases(self):
        self.assertEqual(letter_for_status_name("Approved"), "A")
        self.assertEqual(
            letter_for_status_name("Commented-To be resubmitted-Work can proceed"),
            "C",
        )
        self.assertEqual(
            letter_for_status_name("Commented - To be resubmitted - Work can proceed"),
            "C",
        )
        self.assertEqual(letter_for_status_name("Unknown Status"), "")

    def test_serialize_revisione_includes_lettera(self):
        testata = Testata.objects.create(job="LET01")
        doc = Documento.objects.create(testata=testata, item_no="001", vendor_doc="VD")
        stato = StatoEsterno.objects.create(nome="Final - As Built", lettera="F", colore="#FFC000")
        rev = Revisione.objects.create(documento=doc, rev_no=0, ext_status=stato)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_lettera"], "F")
        self.assertEqual(data["ext_status_label"], "Final - As Built")

    def test_list_stati_esterni_includes_lettera(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        items = list_stati_esterni()
        self.assertEqual(items[0]["lettera"], "A")

    def test_stati_esterni_api_create_and_patch_lettera(self):
        create = self.client.post(
            "/api/stati-esterni/",
            data=json.dumps({"nome": "For Information", "lettera": "z", "colore": "#FFC000"}),
            content_type="application/json",
        )
        self.assertEqual(create.status_code, 201)
        payload = create.json()["data"]
        self.assertEqual(payload["lettera"], "Z")
        pk = payload["id"]

        patch = self.client.patch(
            f"/api/stati-esterni/{pk}/",
            data=json.dumps({"lettera": "Z"}),
            content_type="application/json",
        )
        self.assertEqual(patch.status_code, 200)
        self.assertEqual(patch.json()["data"]["lettera"], "Z")


class RevisioneSbloccoTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.approved = StatoEsterno.objects.create(nome="Approved", colore="#00B050")
        self.rejected = StatoEsterno.objects.create(
            nome="Rejected - Work can not proceed", colore="#D61D09"
        )
        self.testata = Testata.objects.create(job="25012")
        self.doc = Documento.objects.create(
            testata=self.testata,
            vendor_doc="25012-01-4005",
            doc_title="Test doc",
        )
        self.writing_user = User.objects.create_user(
            "sbloc_user", "sbloc@example.com", "pw", permesso=Permesso.WRITING
        )
        self.reading_user = User.objects.create_user(
            "sbloc_read", "sbloc-read@example.com", "pw", permesso=Permesso.READING
        )

    def test_list_includes_latest_with_ext_status(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            rec_act_date="2025-06-01",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili("25012")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["revisione_id"], rev.pk)

    def test_list_excludes_when_successor_exists(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1)
        self.assertEqual(list_revisioni_sbloccabili("25012"), [])

    def test_list_excludes_latest_without_ext_status(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1, dis_plan_date="2025-02-01")
        self.assertEqual(list_revisioni_sbloccabili("25012"), [])

    def test_list_includes_any_ext_status(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            rec_act_date="2025-06-01",
            ext_status=self.rejected,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili("25012")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["revisione_id"], rev.pk)

    def test_sblocca_creates_next_revision(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            rec_act_date="2025-06-01",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        result = sblocca_revisione("25012", rev.pk, "2025-07-15")
        new_rev = Revisione.objects.get(pk=result["revisione"]["id"])
        self.assertEqual(new_rev.rev_no, 3)
        self.assertEqual(new_rev.dis_plan_date.isoformat(), "2025-07-15")
        self.assertTrue(new_rev.crea_nuova_rev)
        self.assertEqual(new_rev.ext_status_id, None)

    def test_sblocca_without_date(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        result = sblocca_revisione("25012", rev.pk)
        new_rev = Revisione.objects.get(pk=result["revisione"]["id"])
        self.assertIsNone(new_rev.dis_plan_date)

    def test_sblocca_rejects_non_latest(self):
        rev0 = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1)
        with self.assertRaises(ValueError):
            sblocca_revisione("25012", rev0.pk)

    def test_api_list_and_sblocca(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        self.client.force_login(self.reading_user)
        list_res = self.client.get("/api/commesse/25012/revisioni/sbloccabili/")
        self.assertEqual(list_res.status_code, 200)
        self.assertEqual(list_res.json()["count"], 1)

        self.client.force_login(self.reading_user)
        post_res = self.client.post(
            "/api/commesse/25012/revisioni/sblocca/",
            data=json.dumps({"revisione_id": rev.pk}),
            content_type="application/json",
        )
        self.assertEqual(post_res.status_code, 403)

        self.client.force_login(self.writing_user)
        post_res = self.client.post(
            "/api/commesse/25012/revisioni/sblocca/",
            data=json.dumps({"revisione_id": rev.pk, "dis_plan_date": "2025-08-01"}),
            content_type="application/json",
        )
        self.assertEqual(post_res.status_code, 200)
        data = post_res.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["count"], 0)
        self.assertEqual(Revisione.objects.filter(documento=self.doc, rev_no=2).count(), 1)


class ImportDocumentiExcelTests(TestCase):
    """Tests for the Excel import endpoint and the import template download."""

    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.writing_user = User.objects.create_user(
            "imp_writer", "w@example.com", "pw", permesso=Permesso.WRITING
        )
        self.reading_user = User.objects.create_user(
            "imp_reader", "r@example.com", "pw", permesso=Permesso.READING
        )
        Testata.objects.create(job="IMP01")

    def _make_xlsx(self, headers, rows):
        """Build an in-memory xlsx file and return it as bytes."""
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # ── import with date column ──────────────────────────────────────────────

    def test_import_with_date_sets_dis_plan_date(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title", "Planned send date (Rev. 0)"],
            [["001", "Doc test", "2026-03-15"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(body["created"], 1)
        self.assertEqual(body["errors"], [])
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNotNone(rev0)
        self.assertIsNotNone(rev0.dis_plan_date)
        self.assertEqual(rev0.dis_plan_date.isoformat(), "2026-03-15")

    def test_import_without_date_column_leaves_dis_plan_date_null(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title"],
            [["002", "Doc senza data"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["created"], 1)
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNone(rev0.dis_plan_date)

    def test_import_with_italian_date_format(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Planned send date (Rev. 0)"],
            [["003", "20/07/2026"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(body["created"], 1)
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertEqual(rev0.dis_plan_date.isoformat(), "2026-07-20")

    def test_import_with_invalid_date_creates_doc_and_reports_warning(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Planned send date (Rev. 0)"],
            [["004", "not-a-date"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        # Document is still created despite the invalid date
        self.assertEqual(body["created"], 1)
        self.assertEqual(len(body["errors"]), 1)
        self.assertIn("data non riconosciuta", body["errors"][0])
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNone(rev0.dis_plan_date)

    def test_import_with_contractor_doc_no(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title", "Contractor Doc N°"],
            [["005", "Doc contractor", "CTR-001"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["created"], 1)
        doc = Documento.objects.get(testata_id="IMP01", item_no="005")
        self.assertEqual(doc.contractor_doc_no, "CTR-001")

    # ── template download ────────────────────────────────────────────────────

    def test_template_download_returns_xlsx(self):
        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            resp.get("Content-Type", ""),
        )

    def test_template_download_contains_date_column(self):
        import io

        import openpyxl

        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Planned send date (Rev. 0)", headers)
        # All expected columns must be present
        for col in ["Item", "B&R Doc", "Contractor Doc N°", "Document title", "Notes"]:
            self.assertIn(col, headers)

    def test_template_download_reparto_has_dropdown_validation(self):
        import io

        import openpyxl
        from openpyxl.utils import get_column_letter

        from core.models import Reparto

        Reparto.objects.create(nome="Qualità e Controllo", acronimo="QC")
        Reparto.objects.create(nome="Project Management", acronimo="PM")

        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))

        self.assertIn("_Reparti", wb.sheetnames)
        ws_reparti = wb["_Reparti"]
        reparti_in_sheet = [ws_reparti.cell(row=r, column=1).value for r in range(1, 3)]
        self.assertEqual(
            sorted(reparti_in_sheet),
            ["Project Management", "Qualità e Controllo"],
        )

        ws = wb.active
        reparto_col = get_column_letter([cell.value for cell in ws[1]].index("Department") + 1)
        validations = [
            dv
            for dv in ws.data_validations.dataValidation
            if any(reparto_col in str(r) for r in dv.sqref.ranges)
        ]
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].type, "list")
        self.assertIn("_Reparti", validations[0].formula1)


class PasswordResetTests(TestCase):
    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.user = User.objects.create_user(
            username="reset_user",
            email="reset@brembanarolle.com",
            password="oldpass123",
        )

    def test_password_reset_page_renders(self):
        response = self.client.get("/password-reset/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Recupera password")

    def test_password_reset_sends_email_for_known_user(self):
        response = self.client.post(
            "/password-reset/",
            data={"email": "reset@brembanarolle.com"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, "/password-reset/done/")
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("reset@brembanarolle.com", mail.outbox[0].to)
        self.assertIn("/reset/", mail.outbox[0].body)
        self.assertRegex(mail.outbox[0].body, r"/reset/[^/\s]+/[^/\s]+/")

    def test_password_reset_unknown_email_shows_done_without_leak(self):
        response = self.client.post(
            "/password-reset/",
            data={"email": "unknown@brembanarolle.com"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, "/password-reset/done/")
        self.assertEqual(len(mail.outbox), 0)

    def test_password_reset_rejects_non_company_email(self):
        response = self.client.post(
            "/password-reset/",
            data={"email": "user@example.com"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "@brembanarolle.com")
        self.assertEqual(len(mail.outbox), 0)

    def test_password_reset_confirm_sets_new_password(self):
        uid = urlsafe_base64_encode(force_bytes(self.user.pk))
        token = default_token_generator.make_token(self.user)
        # Django redirects to .../set-password/ to avoid token leakage via Referer
        initial_url = f"/reset/{uid}/{token}/"
        self.client.get(initial_url)  # stores token in session, redirects to set-password
        confirm_url = f"/reset/{uid}/set-password/"
        response = self.client.post(
            confirm_url,
            data={
                "new_password1": "newsecurepass123",
                "new_password2": "newsecurepass123",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, "/reset/done/")
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("newsecurepass123"))
        self.assertIsNotNone(authenticate(username="reset_user", password="newsecurepass123"))

    def test_login_page_has_reset_link(self):
        response = self.client.get("/login/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'href="/password-reset/"')
        self.assertContains(response, "Recuperala")


class DateDisplayFormatTests(SimpleTestCase):
    """Dates are shown as ``10 jan 2026`` across the app."""

    def test_format_display_date(self):
        import datetime

        from .date_fmt import format_display_date

        self.assertEqual(format_display_date(datetime.date(2026, 1, 10)), "10 jan 2026")
        self.assertEqual(format_display_date("2026-03-05"), "5 mar 2026")
        self.assertEqual(format_display_date(None), "")
        self.assertEqual(format_display_date(""), "")

    def test_pdf_date_line(self):
        from src.pdf import _fmt_date

        self.assertEqual(_fmt_date("2026-01-10", "Milano"), "Milano, 10 jan 2026")
        self.assertEqual(_fmt_date("2026-01-10", ""), "10 jan 2026")


class RevisioneEtichettaTests(SimpleTestCase):
    """Lettera e numero sono due scritture dello stesso dato (revisione_label)."""

    def test_numero_convertito_in_lettera(self):
        self.assertEqual(numero_a_lettera(0), "A")
        self.assertEqual(numero_a_lettera(1), "B")
        self.assertEqual(numero_a_lettera(25), "Z")
        self.assertEqual(numero_a_lettera(26), "AA")
        self.assertEqual(numero_a_lettera(27), "AB")
        self.assertEqual(numero_a_lettera("3"), "D")
        self.assertEqual(numero_a_lettera(None), "")
        self.assertEqual(numero_a_lettera(-1), "")

    def test_lettera_convertita_in_numero(self):
        self.assertEqual(lettera_a_numero("A"), 0)
        self.assertEqual(lettera_a_numero("b"), 1)
        self.assertEqual(lettera_a_numero("Z"), 25)
        self.assertEqual(lettera_a_numero("AA"), 26)
        self.assertIsNone(lettera_a_numero(""))
        self.assertIsNone(lettera_a_numero("1"))
        self.assertIsNone(lettera_a_numero(None))

    def test_con_flag_attivo_si_vede_sempre_la_lettera(self):
        self.assertEqual(format_revisione_label(1, "A", True), "A")
        self.assertEqual(format_revisione_label(1, "a", True), "A")
        # Lettera non compilata: si ricava dal numero.
        self.assertEqual(format_revisione_label(0, "", True), "A")
        self.assertEqual(format_revisione_label(2, "", True), "C")
        self.assertEqual(format_revisione_label(26, None, True), "AA")
        # Lettera compilata con un numero: viene comunque convertita.
        self.assertEqual(format_revisione_label(None, "2", True), "C")
        self.assertEqual(format_revisione_label(None, "", True), "")

    def test_con_flag_spento_si_vede_sempre_il_numero(self):
        self.assertEqual(format_revisione_label(1, "A", False), "1")
        self.assertEqual(format_revisione_label(0, "", False), "0")
        # Numero non compilato: si ricava dalla lettera.
        self.assertEqual(format_revisione_label(None, "B", False), "1")
        self.assertEqual(format_revisione_label(None, "AA", False), "26")
        self.assertEqual(format_revisione_label(None, "", False), "")


class RevisioneLabelDisplayTests(TestCase):
    """Revision display depends on Testata.rev_let_flag, not on rev_let alone."""

    def test_list_documenti_ignores_rev_let_when_flag_false(self):
        t = Testata.objects.create(job="REVFLG1", rev_let_flag=False)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG1-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="C")

        payload = list_documenti("REVFLG1")
        self.assertEqual(len(payload), 1)
        self.assertEqual(payload[0]["latest_rev_display"], "2")
        self.assertEqual(payload[0]["latest_rev_let"], "C")
        self.assertEqual(payload[0]["latest_rev_no"], 2)

    def test_list_documenti_shows_letter_when_flag_true(self):
        t = Testata.objects.create(job="REVFLG2", rev_let_flag=True)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG2-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="C")

        payload = list_documenti("REVFLG2")
        self.assertEqual(payload[0]["latest_rev_display"], "C")

    def test_list_situazione_includes_rev_let_flag(self):
        Testata.objects.create(job="REVFLG3", rev_let_flag=True)
        data = list_situazione("REVFLG3")
        self.assertTrue(data["rev_let_flag"])

    def test_list_documenti_deriva_la_lettera_quando_manca(self):
        t = Testata.objects.create(job="REVFLG4", rev_let_flag=True)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG4-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="")

        self.assertEqual(list_documenti("REVFLG4")[0]["latest_rev_display"], "C")

    def test_list_documenti_deriva_il_numero_quando_manca(self):
        t = Testata.objects.create(job="REVFLG5", rev_let_flag=False)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG5-01")
        Revisione.objects.create(documento=doc, rev_no=None, rev_let="C")

        self.assertEqual(list_documenti("REVFLG5")[0]["latest_rev_display"], "2")

    def test_str_della_revisione_segue_il_flag(self):
        con_lettera = Testata.objects.create(job="REVFLG6", rev_let_flag=True)
        senza_lettera = Testata.objects.create(job="REVFLG7", rev_let_flag=False)
        doc_lettera = Documento.objects.create(testata=con_lettera, vendor_doc="REVFLG6-01")
        doc_numero = Documento.objects.create(testata=senza_lettera, vendor_doc="REVFLG7-01")
        rev_lettera = Revisione.objects.create(documento=doc_lettera, rev_no=1, rev_let="")
        rev_numero = Revisione.objects.create(documento=doc_numero, rev_no=1, rev_let="B")

        self.assertEqual(rev_lettera.etichetta(), "B")
        self.assertEqual(str(rev_lettera), "Rev B")
        self.assertEqual(rev_numero.etichetta(), "1")
        self.assertEqual(str(rev_numero), "Rev 1")


class RevisioneLabelExportTests(TestCase):
    """Excel, PDF e trasmittal stampano la revisione come dice l'archivio."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "revexp_user",
            "revexp@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="REVEXP1", client="ACME", rev_let_flag=True)
        self.doc = Documento.objects.create(
            testata=self.testata,
            item_no="001",
            vendor_doc="REVEXP1-01",
            doc_title="Data book index",
        )
        # Nessuna lettera in archivio: deve comunque uscire "B" (rev_no 1).
        self.rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            rev_let="",
            dis_plan_date=date(2026, 3, 15),
        )

    def _foglio(self, vista):
        import openpyxl

        response = self.client.get(
            f"/api/commesse/{self.testata.job}/situazione/export/?format=xlsx&vista={vista}"
        )
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def _colonna_rev(self):
        """Valori della colonna Rev. dell'Excel verticale."""
        ws = self._foglio("verticale")
        # Le colonne fisse hanno l'intestazione in riga 1 (unita con la riga 2).
        intestazioni = [cell.value for cell in ws[1]]
        col = intestazioni.index("Rev.") + 1
        return [ws.cell(row=r, column=col).value for r in range(3, ws.max_row + 1)]

    def _gruppi_rev_orizzontale(self):
        """Intestazioni dei gruppi revisione dell'Excel orizzontale."""
        ws = self._foglio("orizzontale")
        return [c.value for c in ws[1] if str(c.value or "").startswith("Rev.")]

    def test_export_xlsx_usa_la_lettera_col_flag_attivo(self):
        self.assertEqual(self._colonna_rev(), ["B"])
        # Unico gruppo, etichettato dalla revisione presente (rev_no 1).
        self.assertEqual(self._gruppi_rev_orizzontale(), ["Rev. B"])

    def test_export_xlsx_usa_il_numero_col_flag_spento(self):
        self.testata.rev_let_flag = False
        self.testata.save(update_fields=["rev_let_flag"])
        self.rev.rev_let = "B"
        self.rev.save(update_fields=["rev_let"])

        self.assertEqual(self._colonna_rev(), ["1"])
        self.assertEqual(self._gruppi_rev_orizzontale(), ["Rev. 1"])

    def test_intestazione_pdf_situazione_segue_il_flag(self):
        from src.pdf import _rev_group_label

        revs = [[{"rev_no": 0, "rev_let": ""}, {"rev_no": 1, "rev_let": ""}]]
        self.assertEqual(_rev_group_label(revs, 0, True), "Rev. A")
        self.assertEqual(_rev_group_label(revs, 1, True), "Rev. B")
        self.assertEqual(_rev_group_label(revs, 0, False), "Rev. 0")
        # Posizione senza revisioni: l'etichetta resta coerente col flag.
        self.assertEqual(_rev_group_label([[]], 2, True), "Rev. C")
        self.assertEqual(_rev_group_label([[]], 2, False), "Rev. 2")

    def test_righe_pdf_verticale_seguono_il_flag(self):
        from src.pdf import _situazione_verticale_flat_rows

        documenti = list_documenti(self.testata.job)
        rev_map = revisioni_by_doc_for_job(self.testata.job)
        righe = _situazione_verticale_flat_rows(documenti, rev_map, True)
        self.assertEqual([r["rev_label"] for r in righe], ["B"])
        righe_numero = _situazione_verticale_flat_rows(documenti, rev_map, False)
        self.assertEqual([r["rev_label"] for r in righe_numero], ["1"])

    def test_documenti_del_trasmittal_portano_la_revisione_giusta(self):
        from core.views import _trasmittal_documents
        from src.pdf import _cell_value_for_col

        documents = _trasmittal_documents(self.testata, [self.doc.pk])
        self.assertEqual(documents[0]["rev_label"], "B")
        self.assertEqual(_cell_value_for_col(documents[0], "rev_no"), "B")

        self.testata.rev_let_flag = False
        self.testata.save(update_fields=["rev_let_flag"])
        documents = _trasmittal_documents(self.testata, [self.doc.pk])
        self.assertEqual(documents[0]["rev_label"], "1")
        self.assertEqual(_cell_value_for_col(documents[0], "rev_no"), "1")

    def test_le_pagine_ricevono_il_flag_e_l_helper_condiviso(self):
        """Le tabelle costruite in JS formattano la revisione con la regola condivisa."""
        for url in (
            f"/commesse/{self.testata.job}/documenti/",
            f"/commesse/{self.testata.job}/situazione/",
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "function formatRevLabel(rev, revLetFlag)")
                self.assertContains(response, "REV_LET_FLAG = true")

    def test_export_emissione_xlsx_usa_la_lettera(self):
        import openpyxl

        response = self.client.get(f"/api/commesse/{self.testata.job}/emissione/export/")
        self.assertEqual(response.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        col = [cell.value for cell in ws[1]].index("Rev.") + 1
        self.assertEqual(
            [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)], ["B"]
        )


class CommessaPinTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "pin_user", "pin@example.com", "pw", permesso=Permesso.READING
        )
        self.other = User.objects.create_user(
            "pin_other", "pinother@example.com", "pw", permesso=Permesso.WRITING
        )
        self.jobs = []
        for i in range(10):
            self.jobs.append(Testata.objects.create(job=f"PIN-{i:02d}", client=f"C{i}"))

    def test_reading_user_can_pin_and_unpin(self):
        self.client.force_login(self.user)
        job = self.jobs[0].job
        pin_res = self.client.post(f"/api/commesse/{job}/pin/")
        self.assertEqual(pin_res.status_code, 200)
        self.assertTrue(pin_res.json()["data"]["pinned"])
        self.assertTrue(CommessaPin.objects.filter(user=self.user, testata_id=job).exists())

        list_res = self.client.get("/api/commesse/")
        self.assertEqual(list_res.status_code, 200)
        by_job = {c["job"]: c for c in list_res.json()["commesse"]}
        self.assertTrue(by_job[job]["pinned"])

        unpin_res = self.client.delete(f"/api/commesse/{job}/pin/")
        self.assertEqual(unpin_res.status_code, 200)
        self.assertFalse(unpin_res.json()["data"]["pinned"])
        self.assertFalse(CommessaPin.objects.filter(user=self.user, testata_id=job).exists())

    def test_pin_limit_is_eight(self):
        self.client.force_login(self.user)
        for t in self.jobs[:MAX_PINNED_COMMESSE]:
            res = self.client.post(f"/api/commesse/{t.job}/pin/")
            self.assertEqual(res.status_code, 200)
        overflow = self.client.post(f"/api/commesse/{self.jobs[8].job}/pin/")
        self.assertEqual(overflow.status_code, 400)
        self.assertIn("massimo", overflow.json()["error"].lower())
        self.assertEqual(CommessaPin.objects.filter(user=self.user).count(), 8)

    def test_home_fills_with_recent_when_fewer_than_eight_pinned(self):
        self.client.force_login(self.user)
        pinned_jobs = [self.jobs[0].job, self.jobs[1].job]
        for job in pinned_jobs:
            self.client.post(f"/api/commesse/{job}/pin/")

        res = self.client.get("/api/commesse/?home=1")
        self.assertEqual(res.status_code, 200)
        home = res.json()["commesse"]
        self.assertEqual(len(home), 8)
        self.assertEqual([c["job"] for c in home[:2]], pinned_jobs)
        self.assertTrue(all(c["pinned"] for c in home[:2]))
        self.assertTrue(all(not c["pinned"] for c in home[2:]))
        self.assertNotIn(self.jobs[0].job, [c["job"] for c in home[2:]])

    def test_home_shows_only_eight_when_eight_pinned(self):
        self.client.force_login(self.user)
        pinned = [t.job for t in self.jobs[:8]]
        for job in pinned:
            self.client.post(f"/api/commesse/{job}/pin/")

        res = self.client.get("/api/commesse/?home=1")
        home = res.json()["commesse"]
        self.assertEqual(len(home), 8)
        self.assertEqual([c["job"] for c in home], pinned)
        self.assertTrue(all(c["pinned"] for c in home))
        self.assertNotIn(self.jobs[8].job, [c["job"] for c in home])

    def test_pins_are_per_user(self):
        pin_commessa(self.user, self.jobs[0].job)
        self.client.force_login(self.other)
        res = self.client.get("/api/commesse/?home=1")
        home = res.json()["commesse"]
        by_job = {c["job"]: c for c in home}
        self.assertFalse(by_job.get(self.jobs[0].job, {}).get("pinned", False))
        self.assertEqual(len(list_home_commesse(self.other)), 8)


class SegnalazioniTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.reader = User.objects.create_user(
            "seg_reader",
            "seg-reader@example.com",
            "pw",
            permesso=Permesso.READING,
            first_name="Luca",
            last_name="Rossi",
        )
        self.writer = User.objects.create_user(
            "seg_writer",
            "seg-writer@example.com",
            "pw",
            permesso=Permesso.WRITING,
            first_name="Anna",
            last_name="Bianchi",
        )
        self.admin = User.objects.create_user(
            "seg_admin",
            "seg-admin@example.com",
            "pw",
            permesso=Permesso.ADMIN,
            first_name="Mario",
            last_name="Verdi",
        )
        self.voter = User.objects.create_user(
            "seg_voter",
            "seg-voter@example.com",
            "pw",
            permesso=Permesso.READING,
            first_name="Piero",
            last_name="Neri",
        )

    def _post(self, url, payload, user=None):
        if user:
            self.client.force_login(user)
        return self.client.post(url, data=json.dumps(payload), content_type="application/json")

    def _create(self, user, tipo="feature", titolo="Titolo", testo="Descrizione"):
        return self._post(
            "/api/segnalazioni/",
            {"tipo": tipo, "titolo": titolo, "testo": testo},
            user=user,
        )

    def test_page_requires_login(self):
        res = self.client.get("/segnalazioni/")
        self.assertEqual(res.status_code, 302)

    def test_page_ok_for_authenticated(self):
        self.client.force_login(self.reader)
        res = self.client.get("/segnalazioni/")
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Nuova proposta")
        self.assertContains(res, "Tutti i post")
        self.assertContains(res, "I miei post")
        self.assertContains(res, "Top 10 per voti")
        self.assertContains(res, "Nuove funzionalità")
        self.assertContains(res, "Segnalazioni")
        self.assertContains(res, "Solo aperti")
        self.assertContains(res, "Post chiusi")

    def test_home_contains_prompt_link(self):
        self.client.force_login(self.reader)
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Qualcosa da segnalare?")
        self.assertContains(res, 'href="/segnalazioni/"')

    def test_reading_user_can_create(self):
        res = self._create(self.reader, titolo="Nuova export", testo="Vorrei un export CSV.")
        self.assertEqual(res.status_code, 201)
        data = res.json()["data"]
        self.assertEqual(data["titolo"], "Nuova export")
        self.assertEqual(data["tipo"], TipoSegnalazione.FEATURE)
        self.assertEqual(data["autore_nome"], "Luca Rossi")
        self.assertEqual(data["autore_iniziale"], "L")
        self.assertIsNone(data["autore_avatar"])
        self.assertEqual(data["stato"], StatoSegnalazione.APERTO)
        self.assertEqual(data["score"], 0)
        self.assertEqual(data["mio_voto"], 0)
        self.assertNotIn("voti", data)
        self.assertNotIn("utente", data)

    def test_vote_unique_and_score(self):
        created = self._create(self.reader)
        pk = created.json()["data"]["id"]

        up = self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": 1}, user=self.writer)
        self.assertEqual(up.status_code, 200)
        self.assertEqual(up.json()["data"]["score"], 1)
        self.assertEqual(up.json()["data"]["up"], 1)
        self.assertEqual(up.json()["data"]["mio_voto"], 1)

        again = self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": 1}, user=self.writer)
        self.assertEqual(again.json()["data"]["score"], 1)
        self.assertEqual(SegnalazioneVoto.objects.filter(segnalazione_id=pk).count(), 1)

        down = self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": -1}, user=self.voter)
        self.assertEqual(down.json()["data"]["score"], 0)
        self.assertEqual(down.json()["data"]["down"], 1)

        switch = self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": -1}, user=self.writer)
        self.assertEqual(switch.json()["data"]["score"], -2)

        clear = self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": 0}, user=self.writer)
        self.assertEqual(clear.json()["data"]["mio_voto"], 0)
        self.assertEqual(clear.json()["data"]["score"], -1)

    def test_votes_are_anonymous_in_list(self):
        created = self._create(self.reader, titolo="Voto anonimo")
        pk = created.json()["data"]["id"]
        self._post(f"/api/segnalazioni/{pk}/voto/", {"valore": 1}, user=self.voter)

        self.client.force_login(self.reader)
        res = self.client.get("/api/segnalazioni/")
        self.assertEqual(res.status_code, 200)
        blob = res.content.decode()
        self.assertNotIn("seg_voter", blob)
        self.assertNotIn("Piero Neri", blob)
        item = res.json()["segnalazioni"][0]
        self.assertEqual(item["up"], 1)
        self.assertEqual(item["mio_voto"], 0)
        self.assertNotIn("voti", item)
        for key in item:
            self.assertNotIn("utente", key)

    def test_comment_on_open_and_closed(self):
        pk = self._create(self.reader).json()["data"]["id"]
        comment = self._post(
            f"/api/segnalazioni/{pk}/commenti/",
            {"testo": "Concordo."},
            user=self.writer,
        )
        self.assertEqual(comment.status_code, 201)
        self.assertEqual(comment.json()["data"]["commenti"][0]["autore_nome"], "Anna Bianchi")
        self.assertEqual(SegnalazioneCommento.objects.filter(segnalazione_id=pk).count(), 1)

        close = self._post(f"/api/segnalazioni/{pk}/chiudi/", {}, user=self.admin)
        self.assertEqual(close.status_code, 200)
        self.assertEqual(close.json()["data"]["stato"], StatoSegnalazione.CHIUSO)

        blocked = self._post(
            f"/api/segnalazioni/{pk}/commenti/",
            {"testo": "Troppo tardi."},
            user=self.writer,
        )
        self.assertEqual(blocked.status_code, 403)
        self.assertIn("chiuso", blocked.json()["error"].lower())

    def test_non_admin_cannot_close(self):
        pk = self._create(self.reader).json()["data"]["id"]
        res = self._post(f"/api/segnalazioni/{pk}/chiudi/", {}, user=self.writer)
        self.assertEqual(res.status_code, 403)
        self.assertEqual(res.json()["error"], "Permesso negato.")
        self.assertEqual(Segnalazione.objects.get(pk=pk).stato, StatoSegnalazione.APERTO)

        res2 = self._post(f"/api/segnalazioni/{pk}/riapri/", {}, user=self.reader)
        self.assertEqual(res2.status_code, 403)

    def test_list_orders_open_before_closed_by_score(self):
        low = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Basso",
            testo="x",
        )
        high = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.PROBLEMA,
            titolo="Alto",
            testo="x",
        )
        closed = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Chiuso alto",
            testo="x",
            stato=StatoSegnalazione.CHIUSO,
        )
        SegnalazioneVoto.objects.create(segnalazione=low, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=high, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=high, utente=self.voter, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.voter, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.admin, valore=1)

        self.client.force_login(self.reader)
        titles = [s["titolo"] for s in self.client.get("/api/segnalazioni/").json()["segnalazioni"]]
        self.assertEqual(titles, ["Alto", "Basso", "Chiuso alto"])

    def test_filter_by_tipo(self):
        self._create(self.reader, tipo="feature", titolo="Feat")
        self._create(self.reader, tipo="problema", titolo="Bug")
        self.client.force_login(self.reader)
        feats = self.client.get("/api/segnalazioni/?tipo=feature").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in feats], ["Feat"])
        bugs = self.client.get("/api/segnalazioni/?tipo=problema").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in bugs], ["Bug"])

    def test_filter_by_stato(self):
        self._create(self.reader, titolo="Aperto")
        closed_pk = self._create(self.reader, titolo="Chiuso").json()["data"]["id"]
        self._post(f"/api/segnalazioni/{closed_pk}/chiudi/", {}, user=self.admin)
        self.client.force_login(self.reader)
        aperti = self.client.get("/api/segnalazioni/?stato=aperto").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in aperti], ["Aperto"])
        chiusi = self.client.get("/api/segnalazioni/?stato=chiuso").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in chiusi], ["Chiuso"])

    def test_filter_mine(self):
        self._create(self.reader, titolo="Del reader")
        self._create(self.writer, titolo="Del writer")
        self.client.force_login(self.reader)
        mine = self.client.get("/api/segnalazioni/?mine=1").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in mine], ["Del reader"])
        self.client.force_login(self.writer)
        mine_w = self.client.get("/api/segnalazioni/?mine=1").json()["segnalazioni"]
        self.assertEqual([s["titolo"] for s in mine_w], ["Del writer"])

    def test_top_by_score(self):
        closed = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Chiuso top",
            testo="x",
            stato=StatoSegnalazione.CHIUSO,
        )
        mid = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Medio",
            testo="x",
        )
        high = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.PROBLEMA,
            titolo="Alto",
            testo="x",
        )
        Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Zero",
            testo="x",
        )
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.voter, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=closed, utente=self.admin, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=high, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=high, utente=self.voter, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=mid, utente=self.writer, valore=1)

        for i in range(8):
            Segnalazione.objects.create(
                autore=self.reader,
                tipo=TipoSegnalazione.FEATURE,
                titolo=f"Extra {i}",
                testo="x",
            )
        self.client.force_login(self.reader)
        top = self.client.get("/api/segnalazioni/?top=10").json()["segnalazioni"]
        self.assertEqual(len(top), 10)
        self.assertEqual([s["titolo"] for s in top[:3]], ["Chiuso top", "Alto", "Medio"])
        self.assertNotIn("Zero", [s["titolo"] for s in top])

    def test_list_ordered_by_created_at_desc(self):
        """L'elenco è cronologico (le più recenti in cima), i voti non contano."""
        votata = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Vecchia ma votata",
            testo="x",
        )
        recente = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.PROBLEMA,
            titolo="Recente senza voti",
            testo="x",
        )
        chiusa = Segnalazione.objects.create(
            autore=self.reader,
            tipo=TipoSegnalazione.FEATURE,
            titolo="Chiusa recentissima",
            testo="x",
            stato=StatoSegnalazione.CHIUSO,
        )
        base = timezone.now()
        for s, delta in ((votata, -3), (recente, -1), (chiusa, 0)):
            Segnalazione.objects.filter(pk=s.pk).update(created_at=base + timedelta(hours=delta))
        SegnalazioneVoto.objects.create(segnalazione=votata, utente=self.writer, valore=1)
        SegnalazioneVoto.objects.create(segnalazione=votata, utente=self.voter, valore=1)

        self.client.force_login(self.reader)
        titoli = [s["titolo"] for s in self.client.get("/api/segnalazioni/").json()["segnalazioni"]]
        self.assertEqual(
            titoli,
            ["Recente senza voti", "Vecchia ma votata", "Chiusa recentissima"],
        )

    def test_invalid_list_params(self):
        self.client.force_login(self.reader)
        self.assertEqual(self.client.get("/api/segnalazioni/?tipo=nope").status_code, 400)
        self.assertEqual(self.client.get("/api/segnalazioni/?stato=nope").status_code, 400)
        self.assertEqual(self.client.get("/api/segnalazioni/?top=abc").status_code, 400)
        self.assertEqual(self.client.get("/api/segnalazioni/?top=0").status_code, 400)
        self.assertEqual(self.client.get("/api/segnalazioni/?mine=yes").status_code, 400)


class NotificheTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.autore = User.objects.create_user(
            "notif_autore",
            "notif-autore@example.com",
            "pw",
            permesso=Permesso.READING,
            first_name="Luca",
            last_name="Rossi",
        )
        self.altro = User.objects.create_user(
            "notif_altro",
            "notif-altro@example.com",
            "pw",
            permesso=Permesso.WRITING,
            first_name="Anna",
            last_name="Bianchi",
        )
        self.terzo = User.objects.create_user(
            "notif_terzo",
            "notif-terzo@example.com",
            "pw",
            permesso=Permesso.ADMIN,
            first_name="Mario",
            last_name="Verdi",
        )

    def _crea_segnalazione(self, user, tipo="feature", titolo="Titolo", testo="Descrizione"):
        self.client.force_login(user)
        res = self.client.post(
            "/api/segnalazioni/",
            data=json.dumps({"tipo": tipo, "titolo": titolo, "testo": testo}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 201)
        return res.json()["data"]["id"]

    def test_nuova_feature_notifica_tutti_tranne_autore(self):
        self._crea_segnalazione(self.autore, tipo="feature")
        self.assertEqual(Notifica.objects.filter(destinatario=self.autore).count(), 0)
        testi = set(
            Notifica.objects.exclude(destinatario=self.autore).values_list("testo", flat=True)
        )
        self.assertEqual(testi, {"Luca Rossi ha proposto una nuova feature"})
        self.assertEqual(
            set(
                Notifica.objects.values_list("destinatario__username", flat=True),
            ),
            {"notif_altro", "notif_terzo"},
        )

    def test_nuovo_problema_ha_testo_dedicato(self):
        self._crea_segnalazione(self.autore, tipo="problema")
        testi = set(Notifica.objects.values_list("testo", flat=True))
        self.assertEqual(testi, {"Luca Rossi ha evidenziato un problema"})

    def test_utenti_disattivati_non_ricevono_notifiche(self):
        self.terzo.is_active = False
        self.terzo.save(update_fields=["is_active"])
        self._crea_segnalazione(self.autore)
        self.assertEqual(
            list(Notifica.objects.values_list("destinatario__username", flat=True)),
            ["notif_altro"],
        )

    def test_api_elenca_solo_le_proprie_notifiche(self):
        pk = self._crea_segnalazione(self.autore)
        self.client.force_login(self.altro)
        res = self.client.get("/api/notifiche/")
        self.assertEqual(res.status_code, 200)
        body = res.json()
        self.assertEqual(body["count"], 1)
        item = body["notifiche"][0]
        self.assertEqual(item["testo"], "Luca Rossi ha proposto una nuova feature")
        self.assertEqual(item["segnalazione_id"], pk)
        self.assertEqual(item["tipo"], TipoSegnalazione.FEATURE)
        self.assertIsNotNone(item["created_at"])

        self.client.force_login(self.autore)
        self.assertEqual(self.client.get("/api/notifiche/").json()["count"], 0)

    def test_api_richiede_autenticazione(self):
        self.assertEqual(self.client.get("/api/notifiche/").status_code, 401)
        self.assertEqual(self.client.post("/api/notifiche/lette/").status_code, 401)

    def test_notifica_sparisce_dopo_la_visualizzazione(self):
        self._crea_segnalazione(self.autore)
        self.client.force_login(self.altro)
        ids = [n["id"] for n in self.client.get("/api/notifiche/").json()["notifiche"]]

        lette = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({"ids": ids}),
            content_type="application/json",
        )
        self.assertEqual(lette.status_code, 200)
        self.assertEqual(lette.json()["lette"], 1)
        self.assertEqual(lette.json()["count"], 0)
        self.assertEqual(self.client.get("/api/notifiche/").json()["notifiche"], [])
        self.assertIsNotNone(Notifica.objects.get(pk=ids[0]).letta_il)

        # Le notifiche degli altri utenti restano da leggere
        self.client.force_login(self.terzo)
        self.assertEqual(self.client.get("/api/notifiche/").json()["count"], 1)

    def test_lette_senza_ids_marca_tutte(self):
        self._crea_segnalazione(self.autore, titolo="Prima")
        self._crea_segnalazione(self.terzo, titolo="Seconda")
        self.client.force_login(self.altro)
        self.assertEqual(self.client.get("/api/notifiche/").json()["count"], 2)

        res = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["lette"], 2)
        self.assertEqual(self.client.get("/api/notifiche/").json()["count"], 0)

    def test_lette_marca_solo_gli_id_richiesti(self):
        self._crea_segnalazione(self.autore, titolo="Prima")
        self._crea_segnalazione(self.terzo, titolo="Seconda")
        self.client.force_login(self.altro)
        notifiche = self.client.get("/api/notifiche/").json()["notifiche"]
        self.assertEqual(len(notifiche), 2)

        res = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({"ids": [notifiche[0]["id"]]}),
            content_type="application/json",
        )
        self.assertEqual(res.json()["lette"], 1)
        rimaste = self.client.get("/api/notifiche/").json()["notifiche"]
        self.assertEqual([n["id"] for n in rimaste], [notifiche[1]["id"]])

    def test_lette_rifiuta_payload_non_valido(self):
        self.client.force_login(self.altro)
        bad_type = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({"ids": "1"}),
            content_type="application/json",
        )
        self.assertEqual(bad_type.status_code, 400)
        bad_item = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({"ids": ["abc"]}),
            content_type="application/json",
        )
        self.assertEqual(bad_item.status_code, 400)
        bad_json = self.client.post(
            "/api/notifiche/lette/",
            data="non-json",
            content_type="application/json",
        )
        self.assertEqual(bad_json.status_code, 400)

    def test_lette_ignora_notifiche_di_altri_utenti(self):
        self._crea_segnalazione(self.autore)
        altrui = Notifica.objects.get(destinatario=self.terzo)
        self.client.force_login(self.altro)
        res = self.client.post(
            "/api/notifiche/lette/",
            data=json.dumps({"ids": [altrui.pk]}),
            content_type="application/json",
        )
        self.assertEqual(res.json()["lette"], 0)
        self.assertIsNone(Notifica.objects.get(pk=altrui.pk).letta_il)

    def test_notifiche_rimosse_con_la_segnalazione(self):
        pk = self._crea_segnalazione(self.autore)
        Segnalazione.objects.filter(pk=pk).delete()
        self.assertEqual(Notifica.objects.count(), 0)

    def test_service_helpers(self):
        self._crea_segnalazione(self.autore)
        self.assertEqual(count_notifiche(self.altro), 1)
        self.assertEqual(len(list_notifiche(self.altro)), 1)
        self.assertEqual(segna_lette(self.altro), 1)
        self.assertEqual(count_notifiche(self.altro), 0)
        self.assertEqual(list_notifiche(self.altro), [])
        with self.assertRaises(ValueError):
            segna_lette(self.terzo, ["x"])

    def test_campanella_visibile_nelle_pagine_autenticate(self):
        self.client.force_login(self.altro)
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'id="notif-btn"')
        self.assertContains(res, 'id="notif-panel"')
        self.assertContains(res, 'aria-label="Notifiche"')

    def test_campanella_assente_per_anonimi(self):
        res = self.client.get("/login/")
        self.assertEqual(res.status_code, 200)
        self.assertNotContains(res, 'id="notif-btn"')


class ScaricaDatiGrezziTestCase(TestCase):
    """Pagina "Scarica" ed export Excel dei dati grezzi di una commessa."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "grezzi_reader",
            "grezzi-reader@example.com",
            "pw",
            permesso=Permesso.READING,
        )
        self.testata = Testata.objects.create(job="GR01", client="ACME Spa")
        self.altra = Testata.objects.create(job="GR02", client="Altro Cliente")

        self.indirizzo = IndirSped.objects.create(
            testata=self.testata,
            consignee="ACME Spa",
            city="Bergamo",
            country="IT",
        )
        self.doc = Documento.objects.create(
            testata=self.testata,
            item_no="001",
            vendor_doc="GR01-QMDBI",
            client_doc_no="CL-001",
            client_doc_class="CLASSE-A",
            contractor_doc_no="CTR-001",
            doc_title="Data book index",
        )
        self.stato = StatoEsterno.objects.create(nome="Approved", lettera="A")
        self.rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_plan_date=date(2026, 3, 15),
            int_status="da_emettere",
            ext_status=self.stato,
        )

        # Dati della seconda commessa: non devono finire nell'export di GR01.
        self.doc_altra = Documento.objects.create(
            testata=self.altra,
            item_no="999",
            doc_title="Documento di un'altra commessa",
        )

    def _scarica(self, job="GR01", tabelle="documenti"):
        return self.client.get(f"/api/commesse/{job}/dati-grezzi/?tabelle={tabelle}")

    def _workbook(self, response):
        import io

        import openpyxl

        return openpyxl.load_workbook(io.BytesIO(response.content))

    def _righe(self, ws):
        return [list(row) for row in ws.iter_rows(values_only=True)]

    # ── Pagina HTML ──────────────────────────────────────────────────────────

    def test_pagina_richiede_autenticazione(self):
        res = self.client.get("/scarica/")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login/", res["Location"])

    def test_pagina_elenca_tutte_le_tabelle_e_il_seleziona_tutte(self):
        self.client.force_login(self.user)
        res = self.client.get("/scarica/")
        self.assertEqual(res.status_code, 200)
        self.assertTemplateUsed(res, "core/scarica.html")
        self.assertContains(res, 'id="tabelle-all"')
        self.assertContains(res, "Seleziona tutte")
        self.assertContains(res, 'id="commessa-search"')
        for tabella in list_tabelle_grezze():
            self.assertContains(res, f'value="{tabella["key"]}"')
            self.assertContains(res, tabella["label"])

    def test_navbar_mostra_scarica_tra_cerca_e_impostazioni(self):
        self.client.force_login(self.user)
        for url in ("/", "/commesse/", "/segnalazioni/", "/scarica/"):
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, 200)
                html = res.content.decode()
                self.assertIn('href="/scarica/"', html)
                self.assertIn("Scarica", html)
                # La voce sta fra il pulsante "Cerca" e il link "Impostazioni".
                self.assertLess(html.index("openSpotlight()"), html.index('href="/scarica/"'))
                self.assertLess(html.index('href="/scarica/"'), html.index('href="/admin/"'))

    # ── Export ───────────────────────────────────────────────────────────────

    def test_export_richiede_autenticazione(self):
        self.assertEqual(self._scarica().status_code, 401)

    def test_export_restituisce_un_xlsx_con_nome_file_dedicato(self):
        self.client.force_login(self.user)
        res = self._scarica()
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", res["Content-Disposition"])
        self.assertIn("GR01_dati_grezzi_", res["Content-Disposition"])
        self.assertIn(".xlsx", res["Content-Disposition"])

    def test_un_foglio_per_ogni_tabella_richiesta(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="documenti,revisioni")
        wb = self._workbook(res)
        self.assertEqual(wb.sheetnames, ["Documenti", "Revisioni"])

    def test_i_fogli_seguono_lordine_canonico_non_quello_richiesto(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni,testate,documenti"))
        self.assertEqual(wb.sheetnames, ["Commessa", "Documenti", "Revisioni"])

    def test_seleziona_tutte_scarica_ogni_tabella(self):
        self.client.force_login(self.user)
        keys = ",".join(t["key"] for t in list_tabelle_grezze())
        wb = self._workbook(self._scarica(tabelle=keys))
        self.assertEqual(wb.sheetnames, [t["label"] for t in list_tabelle_grezze()])

    def test_header_con_i_nomi_colonna_e_dato_grezzo(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        righe = self._righe(wb["Documenti"])
        headers = righe[0]
        self.assertEqual(headers[0], "id")
        self.assertIn("testata_id", headers)
        self.assertIn("doc_title", headers)
        self.assertEqual(len(righe), 2)
        riga = dict(zip(headers, righe[1]))
        self.assertEqual(riga["id"], self.doc.pk)
        self.assertEqual(riga["testata_id"], "GR01")
        self.assertEqual(riga["item_no"], "001")
        self.assertEqual(riga["doc_title"], "Data book index")
        self.assertFalse(riga["doc_penalty"])

    def test_solo_header_in_grassetto_e_nessuna_colorazione(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        ws = wb["Documenti"]
        for cell in ws[1]:
            self.assertTrue(cell.font.bold)
        for cell in ws[2]:
            self.assertFalse(cell.font.bold)
            self.assertEqual(cell.fill.fill_type, None)

    def test_export_contiene_solo_i_dati_della_commessa_scelta(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        ids = [r[0] for r in self._righe(wb["Documenti"])]
        self.assertNotIn(self.doc_altra.pk, ids)

        wb_altra = self._workbook(self._scarica(job="GR02", tabelle="documenti"))
        righe = self._righe(wb_altra["Documenti"])
        self.assertEqual(len(righe), 2)
        self.assertEqual(dict(zip(righe[0], righe[1]))["id"], self.doc_altra.pk)

    def test_revisioni_e_indirizzi_esportano_le_righe_collegate(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="testate,indirizzi_spedizione,revisioni"))

        commessa = self._righe(wb["Commessa"])
        self.assertEqual(dict(zip(commessa[0], commessa[1]))["job"], "GR01")

        indirizzi = self._righe(wb["Indirizzi spedizione"])
        self.assertEqual(dict(zip(indirizzi[0], indirizzi[1]))["city"], "Bergamo")

        revisioni = self._righe(wb["Revisioni"])
        riga_rev = dict(zip(revisioni[0], revisioni[1]))
        self.assertEqual(riga_rev["rev_no"], 0)
        self.assertEqual(riga_rev["int_status"], "da_emettere")

    def test_revisioni_portano_i_dati_del_documento_e_la_lettera_dello_stato(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        righe = self._righe(wb["Revisioni"])
        headers = righe[0]
        self.assertEqual(
            headers,
            [
                "client_doc_no",
                "client_doc_class",
                "contractor_doc_no",
                "vendor_doc",
                "item_no",
                "rev_no",
                "rev_let",
                "dis_plan_date",
                "dis_act_date",
                "rec_plan_date",
                "rec_act_date",
                "int_status",
                "ext_status",
                "crea_nuova_rev",
            ],
        )
        riga = dict(zip(headers, righe[1]))
        self.assertEqual(riga["client_doc_no"], "CL-001")
        self.assertEqual(riga["client_doc_class"], "CLASSE-A")
        self.assertEqual(riga["contractor_doc_no"], "CTR-001")
        self.assertEqual(riga["vendor_doc"], "GR01-QMDBI")
        self.assertEqual(riga["item_no"], "001")
        self.assertEqual(riga["ext_status"], "A")

    def test_revisione_senza_risposta_cliente_ha_la_lettera_vuota(self):
        Revisione.objects.filter(pk=self.rev.pk).update(ext_status=None)
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        righe = self._righe(wb["Revisioni"])
        # openpyxl rilegge una cella vuota come None.
        self.assertIn(dict(zip(righe[0], righe[1]))["ext_status"], ("", None))

    def test_le_tabelle_esportabili_sono_solo_le_quattro_utili(self):
        self.assertEqual(
            [t["key"] for t in list_tabelle_grezze()],
            ["testate", "indirizzi_spedizione", "documenti", "revisioni"],
        )
        for key in ("revisioni_file_link", "trasmittal", "trasmittal_revisioni"):
            with self.subTest(key=key), self.assertRaises(ValueError):
                resolve_tabelle_grezze([key])

    def test_tabella_vuota_esporta_il_solo_header(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(job="GR02", tabelle="indirizzi_spedizione"))
        righe = self._righe(wb["Indirizzi spedizione"])
        self.assertEqual(len(righe), 1)
        self.assertIn("consignee", righe[0])

    def test_le_date_restano_date_e_i_timestamp_perdono_il_fuso(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        revisioni = self._righe(wb["Revisioni"])
        self.assertEqual(
            dict(zip(revisioni[0], revisioni[1]))["dis_plan_date"],
            datetime(2026, 3, 15),
        )
        aware = timezone.make_aware(datetime(2026, 3, 15, 8, 30))
        self.assertEqual(_cella_grezza(aware), timezone.localtime(aware).replace(tzinfo=None))
        self.assertIsNone(_cella_grezza(aware).tzinfo)

    # ── Errori ───────────────────────────────────────────────────────────────

    def test_senza_tabelle_selezionate_400(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="")
        self.assertEqual(res.status_code, 400)
        self.assertIn("almeno una tabella", res.json()["error"])

    def test_tabella_sconosciuta_400(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="documenti,pippo")
        self.assertEqual(res.status_code, 400)
        self.assertIn("pippo", res.json()["error"])

    def test_commessa_inesistente_404(self):
        self.client.force_login(self.user)
        res = self._scarica(job="NOPE")
        self.assertEqual(res.status_code, 404)
        self.assertEqual(res.json()["error"], "Commessa non trovata.")

    # ── Service ──────────────────────────────────────────────────────────────

    def test_resolve_tabelle_deduplica_e_ordina(self):
        tabelle = resolve_tabelle_grezze(["revisioni", "documenti", "revisioni"])
        self.assertEqual([t.key for t in tabelle], ["documenti", "revisioni"])

    def test_resolve_tabelle_rifiuta_selezioni_non_valide(self):
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze([])
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze(["   "])
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze(["documenti", "sconosciuta"])

    def test_build_workbook_su_commessa_inesistente(self):
        with self.assertRaises(Testata.DoesNotExist):
            build_dati_grezzi_workbook("NOPE", ["documenti"])


class ColoriRisposteClienteTests(SimpleTestCase):
    """Colori delle celle legate alle risposte del cliente (viste situazione)."""

    # ── Conversioni ──────────────────────────────────────────────────────────

    def test_hex_to_rgb_accetta_le_forme_valide(self):
        self.assertEqual(hex_to_rgb("#00B050"), (0, 176, 80))
        self.assertEqual(hex_to_rgb("00b050"), (0, 176, 80))
        self.assertEqual(hex_to_rgb("  #FFC000 "), (255, 192, 0))
        self.assertEqual(hex_to_rgb("#FFF"), (255, 255, 255))

    def test_hex_to_rgb_rifiuta_i_valori_non_validi(self):
        for value in ["", None, "verde", "#12", "#12345", "#GGGGGG"]:
            self.assertIsNone(hex_to_rgb(value), value)

    def test_rgb_to_hex_normalizza_in_maiuscolo(self):
        self.assertEqual(rgb_to_hex((0, 176, 80)), "#00B050")
        self.assertEqual(rgb_to_hex((255.4, -3, 300)), "#FF00FF")

    # ── Scelta del testo ─────────────────────────────────────────────────────

    def test_testo_nero_su_colori_chiari(self):
        # Bianco e quasi bianco: il testo bianco sparirebbe.
        for colore in ["#FFFFFF", "#FEFEFE", "#FFFF00", "#FFC000", "#9E9E9E"]:
            self.assertEqual(readable_text_color(colore), TEXT_DARK, colore)

    def test_testo_bianco_su_colori_scuri(self):
        # Nero e grigi scuri: il testo nero sparirebbe.
        for colore in ["#000000", "#111111", "#404040", "#0070C0", "#D61D09"]:
            self.assertEqual(readable_text_color(colore), TEXT_LIGHT, colore)

    def test_testo_di_default_quando_il_colore_manca(self):
        self.assertEqual(readable_text_color(""), readable_text_color(FALLBACK_BG))

    # ── Colori della cella ───────────────────────────────────────────────────

    def test_fondo_uguale_al_colore_configurato(self):
        # Uniformità con la legenda: il colore inserito non viene sbiadito.
        for colore in ["#00B050", "#FFC000", "#FFFFFF", "#000000", "#D61D09"]:
            self.assertEqual(cell_colors(colore)["bg"], colore, colore)

    def test_bianco_e_nero_restano_leggibili(self):
        self.assertEqual(cell_colors("#FFFFFF"), {"bg": "#FFFFFF", "fg": TEXT_DARK})
        self.assertEqual(cell_colors("#000000"), {"bg": "#000000", "fg": TEXT_LIGHT})

    def test_mezzi_toni_spostati_quel_tanto_che_basta(self):
        # Su #797979 nessuno dei due testi arriva al contrasto minimo: il fondo
        # si sposta di pochi punti, restando lo stesso colore a vista.
        cella = cell_colors("#797979")
        self.assertNotEqual(cella["bg"], "#797979")
        self.assertGreaterEqual(contrast_ratio(cella["bg"], cella["fg"]), MIN_CONTRAST)
        scarto = max(abs(a - b) for a, b in zip(hex_to_rgb("#797979"), hex_to_rgb(cella["bg"])))
        self.assertLessEqual(scarto, 25)

    def test_contrasto_minimo_garantito_su_tutta_la_gamma(self):
        passo = 51  # campiona 6 livelli per canale
        for r in range(0, 256, passo):
            for g in range(0, 256, passo):
                for b in range(0, 256, passo):
                    colore = rgb_to_hex((r, g, b))
                    cella = cell_colors(colore)
                    self.assertGreaterEqual(
                        contrast_ratio(cella["bg"], cella["fg"]),
                        MIN_CONTRAST,
                        colore,
                    )
                    scarto = max(abs(a - b) for a, b in zip((r, g, b), hex_to_rgb(cella["bg"])))
                    self.assertLessEqual(scarto, 25, colore)

    def test_colore_mancante_o_invalido_usa_il_fallback(self):
        atteso = {"bg": FALLBACK_BG, "fg": readable_text_color(FALLBACK_BG)}
        for colore in ["", None, "  ", "non-un-colore"]:
            self.assertEqual(cell_colors(colore), atteso, colore)


class ColoriRisposteClienteApiTests(TestCase):
    """I colori pronti per le celle viaggiano nei payload delle viste."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "colori_user",
            "colori@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="COL01")
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="COL01-01-0001"
        )
        self.approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        self.final = StatoEsterno.objects.create(
            nome="Final - As Built", lettera="F", colore="#FFFFFF"
        )

    def test_serialize_revisione_include_i_colori_della_cella(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_colore"], "#00B050")
        self.assertEqual(data["ext_status_bg"], "#00B050")
        self.assertEqual(data["ext_status_fg"], TEXT_DARK)

    def test_serialize_revisione_su_stato_bianco_usa_testo_nero(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.final)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_bg"], "#FFFFFF")
        self.assertEqual(data["ext_status_fg"], TEXT_DARK)

    def test_serialize_revisione_senza_risposta_non_ha_colori(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_bg"], "")
        self.assertEqual(data["ext_status_fg"], "")

    def test_situazione_api_include_i_colori_della_cella(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        response = self.client.get(f"/api/commesse/{self.testata.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        revs = response.json()["revisioni_by_doc"][str(self.doc.pk)]
        self.assertEqual(revs[0]["ext_status_bg"], "#00B050")
        self.assertEqual(revs[0]["ext_status_fg"], TEXT_DARK)

    def test_revisioni_sbloccabili_includono_i_colori_della_cella(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-06-01",
            ext_status=self.final,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili(self.testata.job)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["ext_status_bg"], "#FFFFFF")
        self.assertEqual(items[0]["ext_status_fg"], TEXT_DARK)

    def test_pdf_colora_la_cella_vendor_con_il_colore_pieno(self):
        from src.pdf import _status_cell_rgb

        self.assertEqual(_status_cell_rgb("#00B050"), ((0, 176, 80), hex_to_rgb(TEXT_DARK)))
        self.assertEqual(_status_cell_rgb("#D61D09"), ((214, 29, 9), hex_to_rgb(TEXT_LIGHT)))
        self.assertIsNone(_status_cell_rgb(""))


class SituazioneOrizzontaleXlsxTests(TestCase):
    """L'Excel della vista orizzontale riproduce la tabella a schermo."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "sitxls_user",
            "sitxls@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="SITXLS1", rev_let_flag=False)
        self.commented = StatoEsterno.objects.create(
            nome="Commented", lettera="C", colore="#D61D09"
        )
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="SITXLS1-01", doc_title="Index"
        )
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date=date(2026, 1, 10),
            rec_act_date=date(2026, 1, 20),
            ext_status=self.commented,
        )
        Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            dis_act_date=date(2026, 2, 1),
            rec_plan_date=date(2026, 3, 1),
        )
        Documento.objects.create(testata=self.testata, item_no="002", vendor_doc="SITXLS1-02")

    def _foglio(self):
        import openpyxl

        response = self.client.get(
            f"/api/commesse/{self.testata.job}/situazione/export/?format=xlsx&vista=orizzontale"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'filename="situazione_documenti_orizzontale_SITXLS1.xlsx"',
            response["Content-Disposition"],
        )
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def _riga(self, ws, vendor_doc):
        col = [c.value for c in ws[1]].index("B&R Doc") + 1
        return next(r for r in range(3, ws.max_row + 1) if ws.cell(r, col).value == vendor_doc)

    def test_una_riga_per_documento_con_gruppi_per_revisione(self):
        ws = self._foglio()
        intestazioni = [c.value for c in ws[1]]
        self.assertEqual(
            intestazioni[:6],
            ["Client Doc N°", "Client Doc Class", "Contractor Doc N°", "B&R Doc", "Title", "Item"],
        )
        self.assertEqual(
            [v for v in intestazioni if v], intestazioni[:6] + ["Planning", "Rev. 0", "Rev. 1"]
        )
        self.assertEqual(
            [c.value for c in ws[2]][6:],
            ["Submission date", "Receipt date"] + ["Dispatch", "Received", "Status"] * 2,
        )
        # Due righe di intestazione + una riga per ciascuno dei due documenti.
        self.assertEqual(ws.max_row, 4)

    def test_valori_di_planning_e_revisioni(self):
        from .date_fmt import format_display_date

        ws = self._foglio()
        row = self._riga(ws, "SITXLS1-01")
        col = [c.value for c in ws[1]].index("Rev. 0") + 1
        self.assertEqual(
            [ws.cell(row, c).value for c in range(col, col + 6)],
            [
                format_display_date(date(2026, 1, 10)),
                format_display_date(date(2026, 1, 20)),
                "C",
                format_display_date(date(2026, 2, 1)),
                None,
                None,
            ],
        )
        # Ultima revisione inviata: resta pianificato solo il rientro.
        self.assertEqual(ws.cell(row, col - 2).value, None)
        self.assertEqual(ws.cell(row, col - 1).value, format_display_date(date(2026, 3, 1)))

    def test_b_r_doc_e_status_hanno_il_colore_della_risposta(self):
        ws = self._foglio()
        row = self._riga(ws, "SITXLS1-01")
        intestazioni = [c.value for c in ws[1]]
        vendor = ws.cell(row, intestazioni.index("B&R Doc") + 1)
        status_rev0 = ws.cell(row, intestazioni.index("Rev. 0") + 3)
        status_rev1 = ws.cell(row, intestazioni.index("Rev. 1") + 3)
        for cell in (vendor, status_rev0):
            self.assertEqual(cell.fill.fgColor.rgb[-6:], "D61D09")
            self.assertEqual(cell.font.color.rgb[-6:], TEXT_LIGHT.lstrip("#"))
        self.assertTrue(status_rev0.font.bold)
        self.assertIsNone(status_rev1.fill.fill_type)

        senza_revisioni = ws.cell(self._riga(ws, "SITXLS1-02"), intestazioni.index("B&R Doc") + 1)
        self.assertIsNone(senza_revisioni.fill.fill_type)

    def test_la_pagina_apre_il_menu_excel_pdf_in_ogni_vista(self):
        response = self.client.get(f"/commesse/{self.testata.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="btn-export" type="button" onclick="toggleExportMenu(')
        self.assertNotContains(response, "exportSituazione(")


class LegendaStatusPdfTests(TestCase):
    """La legenda STATUS dell'header PDF segue le risposte messe a sistema."""

    def _colori_disegnati(self, vista="orizzontale", **kwargs):
        """Colori passati a ``set_fill_color`` generando il PDF situazione."""
        from src.pdf import SituazioneDocumentiPDF, genera_situazione_documenti_pdf

        originale = SituazioneDocumentiPDF.set_fill_color
        chiamate = []

        def spia(pdf, *args):
            chiamate.append(tuple(args))
            return originale(pdf, *args)

        with patch.object(SituazioneDocumentiPDF, "set_fill_color", spia):
            genera_situazione_documenti_pdf(
                {"job": "LEG01", "client": "ACME"},
                documenti=[],
                revisioni_by_doc={},
                vista=vista,
                **kwargs,
            )
        return chiamate

    # ── Voci della legenda ───────────────────────────────────────────────────

    def test_legenda_usa_lettera_nome_e_colore_configurati(self):
        StatoEsterno.objects.create(nome="Rejected", lettera="R", colore="#8B0000")
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        self.assertEqual(
            legenda_stati_esterni(),
            [
                {"lettera": "A", "nome": "Approved", "colore": "#0070C0"},
                {"lettera": "R", "nome": "Rejected", "colore": "#8B0000"},
            ],
        )

    def test_lettera_dedotta_dal_nome_quando_non_e_configurata(self):
        StatoEsterno.objects.create(nome="Final - As Built", colore="#FFC000")
        self.assertEqual(legenda_stati_esterni()[0]["lettera"], "F")

    def test_colore_mancante_ricade_sul_default_della_lettera(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="")
        self.assertEqual(legenda_stati_esterni()[0]["colore"], DEFAULT_STATUS_COLORS["A"])

    def test_risposta_senza_lettera_riconoscibile_resta_fuori(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        StatoEsterno.objects.create(nome="Stato inventato", colore="#123456")
        self.assertEqual([voce["lettera"] for voce in legenda_stati_esterni()], ["A"])

    def test_senza_risposte_a_sistema_resta_la_legenda_canonica(self):
        self.assertEqual(legenda_stati_esterni(), legenda_default())
        self.assertEqual(
            [voce["lettera"] for voce in legenda_default()],
            ["A", "C", "F", "I", "O", "R", "S", "Z"],
        )

    # ── Righe disegnate nel PDF ──────────────────────────────────────────────

    def test_le_voci_del_pdf_usano_il_fondo_delle_celle(self):
        from src.pdf import _status_legend_entries

        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#FFFFFF")
        self.assertEqual(_status_legend_entries(), [("A", "Approved", (255, 255, 255))])

    def test_voce_con_colore_non_valido_usa_il_fallback(self):
        from src.pdf import _status_legend_entries

        voci = [{"lettera": "A", "nome": "Approved", "colore": "non-un-colore"}]
        self.assertEqual(_status_legend_entries(voci)[0][2], hex_to_rgb(FALLBACK_BG))

    def test_le_voci_si_fermano_alle_righe_disponibili(self):
        from src.pdf import _LEGEND_MAX_ROWS, _status_legend_entries

        voci = [
            {"lettera": chr(ord("A") + i), "nome": f"Stato {i}", "colore": "#00B050"}
            for i in range(_LEGEND_MAX_ROWS + 3)
        ]
        self.assertEqual(len(_status_legend_entries(voci)), _LEGEND_MAX_ROWS)

    def test_le_righe_restano_dentro_il_riquadro_status(self):
        from src.pdf import _HEADER_BOTTOM_PT, _LEGEND_MAX_ROWS, _legend_row_ys

        otto = _legend_row_ys(8)
        # Otto voci = layout storico del foglio di riferimento.
        self.assertEqual(len(otto), 8)
        self.assertAlmostEqual(otto[0], 49.5)
        self.assertAlmostEqual(otto[1] - otto[0], 9.7)
        for count in range(1, _LEGEND_MAX_ROWS + 1):
            ys = _legend_row_ys(count)
            self.assertEqual(len(ys), count)
            self.assertLess(ys[-1], _HEADER_BOTTOM_PT, count)

    # ── PDF generato ─────────────────────────────────────────────────────────

    def test_il_pdf_orizzontale_usa_i_colori_a_sistema(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        colori = self._colori_disegnati(vista="orizzontale")
        self.assertIn((0, 112, 192), colori)
        # Il vecchio verde fisso non viene più disegnato.
        self.assertNotIn((0, 176, 80), colori)

    def test_il_pdf_verticale_usa_i_colori_a_sistema(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        colori = self._colori_disegnati(vista="verticale")
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)

    def test_cambiare_colore_a_sistema_cambia_la_legenda(self):
        approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        self.assertIn((0, 176, 80), self._colori_disegnati())
        approved.colore = "#0070C0"
        approved.save(update_fields=["colore"])
        colori = self._colori_disegnati()
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)

    def test_legenda_passata_a_mano_ha_la_precedenza(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        voci = [{"lettera": "A", "nome": "Approved", "colore": "#0070C0"}]
        colori = self._colori_disegnati(status_legend=voci)
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)


class StatoInternoEffettivoTests(TestCase):
    """Lo stato interno mostrato quando in archivio è rimasto vuoto."""

    def setUp(self):
        self.testata = Testata.objects.create(job="INT01")
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="INT01-01-0001"
        )
        self.approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")

    def test_risposta_cliente_senza_stato_salvato_e_ricevuta(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status"], "")
        self.assertEqual(data["int_status_eff"], "ricevuto")
        self.assertEqual(data["int_status_eff_label"], "Ricevuto")

    def test_data_di_invio_senza_stato_salvato_e_inviata(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, dis_act_date=date(2025, 1, 1))
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "inviato_al_cliente")

    def test_revisione_senza_nulla_resta_da_inviare(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0)
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "")
        self.assertEqual(data["int_status_eff_label"], "Da inviare")

    def test_stato_salvato_ha_la_precedenza(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="in_lavorazione",
            ext_status=self.approved,
        )
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "in_lavorazione")

    def test_documento_con_risposta_cliente_non_e_da_emettere(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        doc = list_documenti(self.testata.job)[0]
        self.assertEqual(doc["latest_int_status"], "ricevuto")
        self.assertEqual(doc["latest_int_status_label"], "Ricevuto")

    def test_etichetta_export_segue_lo_stato_effettivo(self):
        from src.pdf import _int_status_export_label

        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        self.assertEqual(_int_status_export_label(serialize_revisione(rev)), "Ricevuto")
        self.assertEqual(_int_status_export_label(None), "Da inviare")
        self.assertEqual(_int_status_export_label({}), "Da inviare")


# ── Update dati da Business Central ───────────────────────────────────────────


class _FakeBusinessCentral:
    """Connettore Business Central finto con dati preparati per commessa.

    Un valore ``Exception`` fra i dati viene sollevato al posto della risposta,
    così da simulare una query fallita su una singola commessa.
    """

    def __init__(self, dati_per_job, conn=True, codici_sito=None):
        self.dati_per_job = dati_per_job
        self.conn = "connessione-finta" if conn else None
        self.jobs_richiesti = []
        self.chiusa = False
        # sync_business_central invoca anche la sync del sito costruttivo
        # nella stessa passata: di norma nessun test di questa classe se ne
        # occupa, quindi "nessun sito trovato" è il comportamento neutro di
        # default — un test puntuale può passare codici_sito per verificare
        # anche quella parte.
        self.codici_sito = codici_sito or {}

    def dati_commessa(self, job):
        self.jobs_richiesti.append(job)
        dati = self.dati_per_job.get(job, {})
        if isinstance(dati, Exception):
            raise dati
        return dict(dati)

    def get_commessa_codice_sito(self, job):
        return self.codici_sito.get(job)

    def close(self):
        self.chiusa = True


class SincronizzazioneBusinessCentralTests(TestCase):
    """Controllo giornaliero di congruenza fra Business Central e le commesse."""

    def setUp(self):
        self.testata = Testata.objects.create(
            job="26010",
            client="Cliente Vecchio",
            po_no="PO-1",
            job_detail="Descrizione vecchia",
            delivery_date=date(2026, 1, 10),
        )

    def _attiva_bc(self, dati_per_job, conn=True):
        """Sostituisce connettore e lettura BC con un doppio di test."""
        fake = _FakeBusinessCentral(dati_per_job, conn=conn)
        connessione = patch("core.services.bc_sync._apri_connessione", return_value=fake)
        lettura = patch(
            "core.services.bc_sync.fetch_from_bc",
            side_effect=lambda job, bc=None: bc.dati_commessa(job),
        )
        connessione.start()
        lettura.start()
        self.addCleanup(connessione.stop)
        self.addCleanup(lettura.stop)
        return fake

    @staticmethod
    def _dati_bc(**overrides):
        dati = {
            "job": "26010",
            "client": "Cliente Nuovo",
            "job_detail": "Descrizione nuova",
            "po_no": "PO-2",
            "delivery_date": "2026-03-01",
        }
        dati.update(overrides)
        return dati

    def test_aggiorna_i_campi_disallineati(self):
        self._attiva_bc({"26010": self._dati_bc()})

        report = sincronizza_commesse()

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Nuovo")
        self.assertEqual(self.testata.po_no, "PO-2")
        self.assertEqual(self.testata.job_detail, "Descrizione nuova")
        self.assertEqual(self.testata.delivery_date, date(2026, 3, 1))
        self.assertEqual(report["controllate"], 1)
        self.assertEqual(report["aggiornate"], 1)
        self.assertEqual(report["non_trovate"], [])
        self.assertEqual(report["errori"], [])
        self.assertEqual(len(report["aggiornamenti"]), 4)

    def test_accetta_una_data_timestamp_da_bc(self):
        self._attiva_bc({"26010": self._dati_bc(delivery_date=pd.Timestamp("2026-03-01"))})

        sincronizza_commesse()

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.delivery_date, date(2026, 3, 1))

    def test_ogni_modifica_viene_registrata(self):
        self._attiva_bc({"26010": self._dati_bc(job_detail="Descrizione vecchia", po_no="PO-1")})

        sincronizza_commesse()

        registrati = {
            (a.campo, a.valore_precedente, a.valore_nuovo)
            for a in AggiornamentoBC.objects.filter(testata=self.testata)
        }
        self.assertEqual(
            registrati,
            {
                ("client", "Cliente Vecchio", "Cliente Nuovo"),
                ("delivery_date", "2026-01-10", "2026-03-01"),
            },
        )

    def test_valori_vuoti_in_bc_non_sovrascrivono(self):
        self._attiva_bc(
            {
                "26010": {
                    "job": "26010",
                    "client": "   ",
                    "po_no": None,
                    "job_detail": float("nan"),
                    "delivery_date": "",
                }
            }
        )

        report = sincronizza_commesse()

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Vecchio")
        self.assertEqual(self.testata.po_no, "PO-1")
        self.assertEqual(self.testata.job_detail, "Descrizione vecchia")
        self.assertEqual(self.testata.delivery_date, date(2026, 1, 10))
        self.assertEqual(report["aggiornate"], 0)
        self.assertFalse(AggiornamentoBC.objects.exists())

    def test_dati_gia_allineati_non_producono_modifiche(self):
        self._attiva_bc(
            {
                "26010": {
                    "job": "26010",
                    "client": "Cliente Vecchio",
                    "po_no": "PO-1",
                    "job_detail": "Descrizione vecchia",
                    "delivery_date": "2026-01-10",
                }
            }
        )

        report = sincronizza_commesse()

        self.assertEqual(report["controllate"], 1)
        self.assertEqual(report["aggiornate"], 0)
        self.assertEqual(report["aggiornamenti"], [])
        self.assertFalse(AggiornamentoBC.objects.exists())

    def test_dry_run_mostra_le_differenze_senza_salvarle(self):
        self._attiva_bc({"26010": self._dati_bc()})

        report = sincronizza_commesse(dry_run=True)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Vecchio")
        self.assertFalse(AggiornamentoBC.objects.exists())
        self.assertTrue(report["dry_run"])
        self.assertEqual(report["aggiornate"], 1)
        self.assertEqual(len(report["aggiornamenti"]), 4)

    def test_le_commesse_chiuse_sono_escluse(self):
        Testata.objects.create(job="25001", actual_delivery_date=date(2025, 12, 31))
        fake = self._attiva_bc({"26010": self._dati_bc()})

        sincronizza_commesse()

        self.assertEqual(fake.jobs_richiesti, ["26010"])

    def test_le_commesse_chiuse_si_possono_includere(self):
        Testata.objects.create(job="25001", actual_delivery_date=date(2025, 12, 31))
        fake = self._attiva_bc({"26010": self._dati_bc()})

        sincronizza_commesse(includi_chiuse=True)

        self.assertEqual(fake.jobs_richiesti, ["25001", "26010"])

    def test_si_puo_controllare_una_sola_commessa(self):
        Testata.objects.create(job="26011")
        fake = self._attiva_bc({"26010": self._dati_bc()})

        sincronizza_commesse(jobs=["26010"])

        self.assertEqual(fake.jobs_richiesti, ["26010"])

    def test_commessa_assente_in_bc_viene_segnalata(self):
        self._attiva_bc({})

        report = sincronizza_commesse()

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Vecchio")
        self.assertEqual(report["non_trovate"], ["26010"])
        self.assertEqual(report["aggiornate"], 0)

    def test_un_errore_non_blocca_le_altre_commesse(self):
        Testata.objects.create(job="26011", client="Altro Cliente")
        self._attiva_bc(
            {
                "26010": RuntimeError("query fallita"),
                "26011": {"job": "26011", "client": "Cliente Nuovo"},
            }
        )

        report = sincronizza_commesse()

        self.assertEqual(report["controllate"], 2)
        self.assertEqual(report["aggiornate"], 1)
        self.assertEqual(report["errori"], [{"job": "26010", "errore": "query fallita"}])
        self.assertEqual(Testata.objects.get(job="26011").client, "Cliente Nuovo")

    def test_connessione_non_disponibile_solleva_errore(self):
        fake = self._attiva_bc({"26010": self._dati_bc()}, conn=False)

        with self.assertRaises(BusinessCentralNonDisponibile):
            sincronizza_commesse()

        self.assertTrue(fake.chiusa)
        self.assertFalse(AggiornamentoBC.objects.exists())

    def test_la_connessione_viene_chiusa_a_fine_controllo(self):
        fake = self._attiva_bc({"26010": self._dati_bc()})

        sincronizza_commesse()

        self.assertTrue(fake.chiusa)

    def test_valore_bc_troppo_lungo_viene_troncato_e_resta_allineato(self):
        descrizione = "X" * 400
        self._attiva_bc({"26010": self._dati_bc(job_detail=descrizione)})

        sincronizza_commesse()
        self.testata.refresh_from_db()
        self.assertEqual(self.testata.job_detail, "X" * 300)

        # Secondo giro: il valore troncato non deve risultare di nuovo diverso.
        report = sincronizza_commesse()
        self.assertEqual(report["aggiornate"], 0)
        self.assertEqual(AggiornamentoBC.objects.filter(campo="job_detail").count(), 1)

    def test_confronta_ignora_i_campi_non_sincronizzati(self):
        differenze = confronta_commessa(self.testata, {"job": "99999", "requisition": "BID-9"})

        self.assertEqual(differenze, [])

    def test_lista_aggiornamenti_per_la_pagina_archivio(self):
        self._attiva_bc({"26010": self._dati_bc()})
        sincronizza_commesse()

        voci = {v["campo"]: v for v in list_aggiornamenti("26010")}

        self.assertEqual(voci["client"]["etichetta"], "Cliente")
        self.assertEqual(voci["client"]["nuovo"], "Cliente Nuovo")
        # Le date sono già pronte per l'interfaccia (10 jan 2026 → 1 mar 2026).
        self.assertEqual(voci["delivery_date"]["precedente"], "10 jan 2026")
        self.assertEqual(voci["delivery_date"]["nuovo"], "1 mar 2026")

    def test_fetch_from_bc_riusa_il_connettore_aperto(self):
        bc = Mock()
        bc.get_commessa_anagrafica.return_value = pd.DataFrame(
            [{"commessa": "26010", "descrizione": "Descrizione BC", "cliente": "Cliente BC"}]
        )
        bc.get_commessa_commerciale.return_value = pd.DataFrame(
            [{"commessa": "26010", "po_cliente": "PO-BC", "data_consegna": "2026-03-01"}]
        )

        dati = fetch_from_bc("26010", bc=bc)

        self.assertEqual(dati["client"], "Cliente BC")
        self.assertEqual(dati["job_detail"], "Descrizione BC")
        self.assertEqual(dati["po_no"], "PO-BC")
        self.assertEqual(dati["delivery_date"], "2026-03-01")
        bc.close.assert_not_called()


class _FakeBusinessCentralSito:
    """Connettore finto per get_commessa_codice_sito, per commessa.

    Un valore ``Exception`` fra i dati viene sollevato al posto della
    risposta, per simulare una query fallita su una singola commessa.
    """

    def __init__(self, codici_per_job, conn=True):
        self.codici_per_job = codici_per_job
        self.conn = "connessione-finta" if conn else None
        self.chiusa = False

    def get_commessa_codice_sito(self, job):
        codice = self.codici_per_job.get(job)
        if isinstance(codice, Exception):
            raise codice
        return codice

    def close(self):
        self.chiusa = True


class SincronizzazioneSitoCostruttivoTests(TestCase):
    """Backfill di Testata.sito_costruttivo dal codice sito Business Central."""

    def setUp(self):
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.pd = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        self.senza_sito = Testata.objects.create(job="99050")

    def _attiva_bc(self, codici_per_job, conn=True):
        fake = _FakeBusinessCentralSito(codici_per_job, conn=conn)
        patcher = patch("core.services.bc_sync._apri_connessione", return_value=fake)
        patcher.start()
        self.addCleanup(patcher.stop)
        return fake

    def test_imposta_il_sito_dal_codice_bc(self):
        self._attiva_bc({"99050": "1"})

        report = sincronizza_sito_costruttivo()

        self.senza_sito.refresh_from_db()
        self.assertEqual(self.senza_sito.sito_costruttivo_id, self.bg.pk)
        self.assertEqual(report["aggiornate"], 1)
        self.assertEqual(report["aggiornamenti"], [{"job": "99050", "stabilimento": "Valbrembo"}])

    def test_registra_l_aggiornamento_per_il_pannello_archivio(self):
        self._attiva_bc({"99050": "2"})

        sincronizza_sito_costruttivo()

        voci = {v["campo"]: v for v in list_aggiornamenti("99050")}
        self.assertEqual(voci["sito_costruttivo"]["etichetta"], "Sito costruttivo")
        self.assertEqual(voci["sito_costruttivo"]["nuovo"], "Albignasego")

    def test_non_sovrascrive_un_sito_gia_impostato(self):
        self.senza_sito.sito_costruttivo = self.pd
        self.senza_sito.save(update_fields=["sito_costruttivo"])
        fake = self._attiva_bc({"99050": "1"})

        report = sincronizza_sito_costruttivo()

        self.senza_sito.refresh_from_db()
        self.assertEqual(self.senza_sito.sito_costruttivo_id, self.pd.pk)
        self.assertEqual(report["controllate"], 0)
        self.assertEqual(fake.chiusa, False)  # mai aperta: nessuna commessa da controllare

    def test_commessa_non_trovata_o_senza_sito_in_bc(self):
        self._attiva_bc({"99050": None})

        report = sincronizza_sito_costruttivo()

        self.assertEqual(report["non_trovate"], ["99050"])
        self.assertEqual(report["aggiornate"], 0)
        self.senza_sito.refresh_from_db()
        self.assertIsNone(self.senza_sito.sito_costruttivo_id)

    def test_codice_sito_senza_stabilimento_corrispondente(self):
        self._attiva_bc({"99050": "999"})

        report = sincronizza_sito_costruttivo()

        self.assertEqual(report["senza_stabilimento"], [{"job": "99050", "codice_sito": "999"}])
        self.senza_sito.refresh_from_db()
        self.assertIsNone(self.senza_sito.sito_costruttivo_id)

    def test_errore_su_una_commessa_non_ferma_le_altre(self):
        altra = Testata.objects.create(job="99051")
        self._attiva_bc({"99050": RuntimeError("query fallita"), "99051": "2"})

        report = sincronizza_sito_costruttivo()

        self.assertEqual(len(report["errori"]), 1)
        self.assertEqual(report["errori"][0]["job"], "99050")
        altra.refresh_from_db()
        self.assertEqual(altra.sito_costruttivo_id, self.pd.pk)

    def test_dry_run_non_scrive_nulla(self):
        self._attiva_bc({"99050": "1"})

        report = sincronizza_sito_costruttivo(dry_run=True)

        self.assertEqual(report["aggiornate"], 1)
        self.senza_sito.refresh_from_db()
        self.assertIsNone(self.senza_sito.sito_costruttivo_id)

    def test_limita_a_un_job(self):
        Testata.objects.create(job="99052")
        self._attiva_bc({"99050": "1", "99052": "2"})

        report = sincronizza_sito_costruttivo(jobs=["99050"])

        self.assertEqual(report["controllate"], 1)
        self.assertEqual(Testata.objects.get(job="99052").sito_costruttivo_id, None)

    def test_connessione_non_disponibile(self):
        self._attiva_bc({"99050": "1"}, conn=False)

        with self.assertRaises(BusinessCentralNonDisponibile):
            sincronizza_sito_costruttivo()


class ComandoSyncBusinessCentralTests(TestCase):
    """Il comando ``sync_business_central``, pensato per l'esecuzione giornaliera."""

    def setUp(self):
        self.testata = Testata.objects.create(job="26010", client="Cliente Vecchio")

    def _attiva_bc(self, dati_per_job, conn=True, codici_sito=None):
        fake = _FakeBusinessCentral(dati_per_job, conn=conn, codici_sito=codici_sito)
        connessione = patch("core.services.bc_sync._apri_connessione", return_value=fake)
        lettura = patch(
            "core.services.bc_sync.fetch_from_bc",
            side_effect=lambda job, bc=None: bc.dati_commessa(job),
        )
        connessione.start()
        lettura.start()
        self.addCleanup(connessione.stop)
        self.addCleanup(lettura.stop)
        return fake

    def test_il_comando_aggiorna_e_riepiloga(self):
        self._attiva_bc({"26010": {"job": "26010", "client": "Cliente Nuovo"}})
        out = io.StringIO()

        call_command("sync_business_central", stdout=out)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Nuovo")
        output = out.getvalue()
        self.assertIn("26010: Cliente: Cliente Vecchio → Cliente Nuovo", output)
        self.assertIn("Controllate 1 commesse, aggiornate 1", output)

    def test_il_comando_aggiorna_anche_il_sito_costruttivo(self):
        bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self._attiva_bc(
            {"26010": {"job": "26010", "client": "Cliente Nuovo"}},
            codici_sito={"26010": "1"},
        )
        out = io.StringIO()

        call_command("sync_business_central", stdout=out)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.sito_costruttivo_id, bg.pk)
        output = out.getvalue()
        self.assertIn("26010: sito costruttivo -> Valbrembo", output)
        self.assertIn("Sito costruttivo — controllate 1 commesse, aggiornate 1", output)

    def test_il_comando_non_sovrascrive_un_sito_gia_impostato(self):
        pd_stabilimento = Stabilimento.objects.create(nome="Albignasego", sigla="PD", codice_bc=2)
        self.testata.sito_costruttivo = pd_stabilimento
        self.testata.save(update_fields=["sito_costruttivo"])
        self._attiva_bc(
            {"26010": {"job": "26010", "client": "Cliente Nuovo"}},
            codici_sito={"26010": "1"},  # BG: diverso, ma il campo è già impostato
        )
        out = io.StringIO()

        call_command("sync_business_central", stdout=out)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.sito_costruttivo_id, pd_stabilimento.pk)

    def test_il_comando_in_dry_run_non_salva(self):
        self._attiva_bc({"26010": {"job": "26010", "client": "Cliente Nuovo"}})
        out = io.StringIO()

        call_command("sync_business_central", "--dry-run", stdout=out)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.client, "Cliente Vecchio")
        self.assertIn("[dry-run]", out.getvalue())

    def test_il_comando_accetta_una_singola_commessa(self):
        Testata.objects.create(job="26011")
        fake = self._attiva_bc({"26010": {"job": "26010", "client": "Cliente Nuovo"}})

        call_command("sync_business_central", "--job", "26010", stdout=io.StringIO())

        self.assertEqual(fake.jobs_richiesti, ["26010"])

    def test_il_comando_segnala_una_commessa_inesistente(self):
        self._attiva_bc({})
        err = io.StringIO()

        call_command("sync_business_central", "--job", "99999", stderr=err)

        self.assertIn("non trovata", err.getvalue())

    def test_il_comando_segnala_business_central_non_raggiungibile(self):
        self._attiva_bc({"26010": {"job": "26010"}}, conn=False)
        err = io.StringIO()

        call_command("sync_business_central", stderr=err)

        self.assertIn("Connessione a Business Central non disponibile", err.getvalue())


class ComandoBackfillSitoCostruttivoTests(TestCase):
    """Il comando ``backfill_sito_costruttivo``, eseguito una tantum a mano."""

    def setUp(self):
        self.bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        self.testata = Testata.objects.create(job="99060")

    def _attiva_bc(self, codici_per_job, conn=True):
        fake = _FakeBusinessCentralSito(codici_per_job, conn=conn)
        patcher = patch("core.services.bc_sync._apri_connessione", return_value=fake)
        patcher.start()
        self.addCleanup(patcher.stop)
        return fake

    def test_il_comando_aggiorna_e_riepiloga(self):
        self._attiva_bc({"99060": "1"})
        out = io.StringIO()

        call_command("backfill_sito_costruttivo", stdout=out)

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.sito_costruttivo_id, self.bg.pk)
        output = out.getvalue()
        self.assertIn("99060: sito costruttivo -> Valbrembo", output)
        self.assertIn("aggiornate 1", output)

    def test_il_comando_in_dry_run_non_salva(self):
        self._attiva_bc({"99060": "1"})
        out = io.StringIO()

        call_command("backfill_sito_costruttivo", "--dry-run", stdout=out)

        self.testata.refresh_from_db()
        self.assertIsNone(self.testata.sito_costruttivo_id)
        self.assertIn("[dry-run]", out.getvalue())

    def test_il_comando_accetta_una_singola_commessa(self):
        Testata.objects.create(job="99061")
        self._attiva_bc({"99060": "1", "99061": "1"})

        call_command("backfill_sito_costruttivo", "--job", "99060", stdout=io.StringIO())

        self.testata.refresh_from_db()
        self.assertEqual(self.testata.sito_costruttivo_id, self.bg.pk)
        self.assertIsNone(Testata.objects.get(job="99061").sito_costruttivo_id)

    def test_il_comando_segnala_una_commessa_inesistente(self):
        self._attiva_bc({})
        err = io.StringIO()

        call_command("backfill_sito_costruttivo", "--job", "99999", stderr=err)

        self.assertIn("non trovata", err.getvalue())

    def test_il_comando_segnala_business_central_non_raggiungibile(self):
        self._attiva_bc({"99060": "1"}, conn=False)
        err = io.StringIO()

        call_command("backfill_sito_costruttivo", stderr=err)

        self.assertIn("Connessione a Business Central non disponibile", err.getvalue())


class SchedulerBusinessCentralTests(TestCase):
    """Lo scheduler interno che lancia la sincronizzazione BC una volta al giorno."""

    def setUp(self):
        self.oggi = timezone.localtime().date()

    @staticmethod
    def _momento(giorno, ora, minuto):
        """Istante locale, con il fuso della configurazione."""
        return timezone.make_aware(datetime.combine(giorno, time(ora, minuto)))

    def _sync_finta(self, **report):
        """Sostituisce la sincronizzazione vera con un doppio di test."""
        dati = {
            "controllate": 3,
            "aggiornate": 1,
            "non_trovate": [],
            "errori": [],
        }
        dati.update(report)
        finta = patch("core.services.scheduler.sincronizza_commesse", return_value=dati)
        mock = finta.start()
        self.addCleanup(finta.stop)
        return mock

    def _stato(self):
        return EsecuzioneSchedulata.objects.get(nome=scheduler.NOME_JOB)

    def test_non_esegue_prima_dell_orario(self):
        adesso = self._momento(self.oggi, 16, 59)

        self.assertFalse(scheduler._deve_eseguire(None, adesso))

    def test_esegue_dopo_l_orario(self):
        adesso = self._momento(self.oggi, 17, 1)

        self.assertTrue(scheduler._deve_eseguire(None, adesso))

    def test_non_riesegue_nello_stesso_giorno(self):
        ultima = self._momento(self.oggi, 17, 2)
        adesso = self._momento(self.oggi, 18, 0)

        self.assertFalse(scheduler._deve_eseguire(ultima, adesso))

    def test_recupera_l_esecuzione_persa(self):
        """Server spento alle 17:00 e riacceso alle 19:00: la sync parte comunque."""
        ultima = self._momento(self.oggi - timedelta(days=1), 17, 0)
        adesso = self._momento(self.oggi, 19, 0)

        self.assertTrue(scheduler._deve_eseguire(ultima, adesso))

    def test_esegue_la_sync_e_registra_l_esito(self):
        sync = self._sync_finta()
        adesso = self._momento(self.oggi, 17, 1)

        self.assertTrue(scheduler.esegui_se_dovuto(adesso=adesso))

        sync.assert_called_once_with()
        stato = self._stato()
        self.assertEqual(timezone.localtime(stato.ultima_esecuzione), adesso)
        self.assertIn("controllate 3", stato.esito)
        self.assertIn("aggiornate 1", stato.esito)

    def test_un_secondo_giro_non_riesegue(self):
        sync = self._sync_finta()
        adesso = self._momento(self.oggi, 17, 1)

        scheduler.esegui_se_dovuto(adesso=adesso)
        eseguito = scheduler.esegui_se_dovuto(adesso=self._momento(self.oggi, 17, 6))

        self.assertFalse(eseguito)
        self.assertEqual(sync.call_count, 1)

    def test_business_central_non_disponibile_non_propaga(self):
        finta = patch(
            "core.services.scheduler.sincronizza_commesse",
            side_effect=BusinessCentralNonDisponibile("ERP irraggiungibile"),
        )
        sync = finta.start()
        self.addCleanup(finta.stop)
        adesso = self._momento(self.oggi, 17, 1)

        self.assertTrue(scheduler.esegui_se_dovuto(adesso=adesso))

        self.assertIn("errore", self._stato().esito)
        # Un tentativo al giorno: dopo un errore si ritenta domani, non subito.
        self.assertFalse(scheduler.esegui_se_dovuto(adesso=self._momento(self.oggi, 17, 6)))
        self.assertEqual(sync.call_count, 1)

    def test_un_errore_imprevisto_resta_nel_giro(self):
        finta = patch(
            "core.services.scheduler.sincronizza_commesse",
            side_effect=RuntimeError("driver ODBC assente"),
        )
        finta.start()
        self.addCleanup(finta.stop)

        scheduler.esegui_se_dovuto(adesso=self._momento(self.oggi, 17, 1))

        self.assertIn("driver ODBC assente", self._stato().esito)

    def test_orario_malformato_ricade_sul_predefinito(self):
        with self.settings(BC_SYNC_ORARIO="mezzogiorno"):
            self.assertEqual(scheduler._ora_schedulata(), scheduler.ORARIO_PREDEFINITO)

    def test_lo_scheduler_non_parte_durante_i_test(self):
        with patch.object(sys, "argv", ["manage.py", "test"]):
            self.assertFalse(scheduler._processo_adatto())

    def test_lo_scheduler_parte_sotto_il_server_wsgi(self):
        with patch.object(sys, "argv", ["gunicorn", "config.wsgi:application"]):
            self.assertTrue(scheduler._processo_adatto())


class ArchivioAggiornamentiBCViewTests(TestCase):
    """La pagina Informazioni archivio mostra le modifiche arrivate da BC."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "bc_user",
            "bc@brembanarolle.com",
            "pw",
            permesso=Permesso.READING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="26010", client="Cliente Nuovo")

    def test_la_pagina_elenca_gli_aggiornamenti(self):
        AggiornamentoBC.objects.create(
            testata=self.testata,
            campo="client",
            valore_precedente="Cliente Vecchio",
            valore_nuovo="Cliente Nuovo",
        )

        response = self.client.get(f"/commesse/{self.testata.job}/archivio/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Aggiornamenti da Business Central")
        self.assertContains(response, "Cliente Vecchio")

    def test_senza_aggiornamenti_mostra_lo_stato_allineato(self):
        response = self.client.get(f"/commesse/{self.testata.job}/archivio/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nessun aggiornamento da Business Central")


def _dai_firma(utente, nome="firme/prova.png"):
    """L'utente con un'immagine di firma.

    Basta il nome del file: non serve leggerlo. I test che caricano davvero
    l'immagine usano ``_png`` e una cartella media temporanea.
    """
    utente.firma = nome
    utente.save(update_fields=["firma"])
    return utente


def _png(dimensione=(160, 48), modo="RGBA", formato="PNG"):
    """Un'immagine vera, come la carica un utente."""
    buffer = io.BytesIO()
    Image.new(modo, dimensione, (20, 40, 90, 0) if modo == "RGBA" else (255, 255, 255)).save(
        buffer, format=formato
    )
    return buffer.getvalue()


class FirmaUtenteTests(TestCase):
    """Validatore dell'immagine di firma e upload/rimozione dal profilo."""

    def setUp(self):
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        impostazioni = override_settings(MEDIA_ROOT=media.name)
        impostazioni.enable()
        self.addCleanup(impostazioni.disable)

        self.client = Client()
        self.writer = User.objects.create_user(
            "firma_writer",
            "firma-writer@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
            first_name="Mario",
            last_name="Rossi",
        )
        _dai_firma(self.writer)
        self.client.force_login(self.writer)

    def _profilo(self, **campi):
        # Il modulo del profilo manda tutti i campi: uno che manca si svuota.
        dati = {
            "action": "update_profile",
            "first_name": "Mario",
            "last_name": "Rossi",
            "email": self.writer.email,
        }
        return self.client.post("/profilo/", {**dati, **campi})

    def test_upload_di_una_immagine_di_firma_valida(self):
        risposta = self._profilo(firma=SimpleUploadedFile("firma.png", _png(), "image/png"))

        self.assertEqual(risposta.status_code, 302)
        self.writer.refresh_from_db()
        self.assertTrue(self.writer.firma.name.startswith("firme/"))
        self.assertNotEqual(self.writer.firma.name, "firme/prova.png")
        # Il file si legge e si chiude subito: su Windows un file aperto non si
        # cancella, e il prossimo upload sostituisce quello vecchio.
        with self.writer.firma.open("rb") as file:
            contenuto = file.read()
        with Image.open(io.BytesIO(contenuto)) as immagine:
            self.assertEqual((immagine.format, immagine.mode), ("PNG", "RGBA"))
        # Anche un JPEG va bene.
        risposta = self._profilo(
            firma=SimpleUploadedFile("firma.jpg", _png(modo="RGB", formato="JPEG"), "image/jpeg")
        )
        self.assertEqual(risposta.status_code, 302)

    def test_upload_rifiuta_un_file_che_non_e_un_immagine(self):
        for nome, contenuto, messaggio in (
            # L'apostrofo nell'HTML è &#x27;: si cerca il resto del messaggio.
            ("firma.png", b"non sono un'immagine", "immagine valida: carica un PNG"),
            ("firma.gif", _png(modo="RGB", formato="GIF"), "Formato non ammesso"),
        ):
            with self.subTest(nome=nome):
                risposta = self._profilo(firma=SimpleUploadedFile(nome, contenuto, "image/png"))

                self.assertEqual(risposta.status_code, 200)
                self.assertContains(risposta, messaggio)
                # L'immagine che c'era resta.
                self.writer.refresh_from_db()
                self.assertEqual(self.writer.firma.name, "firme/prova.png")

    def test_upload_rifiuta_un_file_oltre_il_limite(self):
        pesante = _png() + b"\0" * FIRMA_MAX_BYTE

        risposta = self._profilo(firma=SimpleUploadedFile("firma.png", pesante, "image/png"))

        self.assertContains(risposta, "al massimo 1 MB")

    def test_l_immagine_di_firma_si_toglie_dal_profilo(self):
        self._profilo(firma=SimpleUploadedFile("firma.png", _png(), "image/png"))

        self._profilo(rimuovi_firma="1")

        self.writer.refresh_from_db()
        self.assertFalse(self.writer.firma)
        pagina = self.client.get("/profilo/").content.decode()
        self.assertIn("Nessuna immagine di firma caricata.", pagina)
        self.assertNotIn("firma digitale", pagina.lower())
