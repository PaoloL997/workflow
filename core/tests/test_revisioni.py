import io
import json
from datetime import date

import openpyxl
import pandas as pd
from django.test import Client, SimpleTestCase, TestCase

from ..models import (
    Documento,
    Permesso,
    Revisione,
    StatoEsterno,
    Testata,
)
from ..services.commesse import (
    esegui_emissione,
    esegui_ricezione,
    list_documenti,
    list_situazione,
    list_stati_esterni,
    revisioni_by_doc_for_job,
    serialize_revisione,
)
from ..services.revisione_anomalie import (
    audit_commessa,
    audit_commessa_summary,
    audit_revisione,
    classifica_revisione,
    ignora_anomalie_revisione,
    serialize_anomalie_gruppi,
)
from ..services.revisione_label import (
    format_revisione_label,
    lettera_a_numero,
    numero_a_lettera,
    numero_revisione,
    valida_sequenza_revisioni,
)
from ..services.revisione_sblocco import list_revisioni_sbloccabili, sblocca_revisione
from ..services.revisioni_cleanup import drop_orphan_revisioni, find_orphan_indices
from ..services.stato_esterno_codes import letter_for_status_name

# ── Shared fixture factory (unit tests only — uses temp dir) ──────────────────
from ._fixtures import User


class RevisioniCleanupTests(TestCase):
    def test_single_orphan_after_approved_with_only_dis_plan(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "C",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": "2025-02-01",
                    "RecPlanDate": "2025-02-10",
                    "RecActDate": "2025-02-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 2,
                    "DisPlanDate": "2025-03-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        orphans = find_orphan_indices(df)
        self.assertEqual(orphans, {2})
        cleaned = drop_orphan_revisioni(df)
        self.assertEqual(len(cleaned), 2)

    def test_not_orphan_when_previous_status_needs_new_rev(self):
        # Status C has crea_nuova_rev=True by default → keep last placeholder.
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "C",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())

    def test_not_orphan_when_dis_act_date_present(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": "2025-02-01",
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())

    def test_not_orphan_without_dis_plan_date(self):
        df = pd.DataFrame(
            [
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 0,
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "JOB-01",
                    "RevNo": 1,
                    "DisPlanDate": None,
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        self.assertEqual(find_orphan_indices(df), set())


class RevisioneAnomalieTests(TestCase):
    def setUp(self):
        self.approved = StatoEsterno.objects.create(nome="Approved", colore="#00B050")
        self.rejected = StatoEsterno.objects.create(
            nome="Rejected - Work can not proceed", colore="#D61D09"
        )
        self.testata = Testata.objects.create(job="25012", time_cli_doc_rev=21)
        self.doc = Documento.objects.create(
            testata=self.testata,
            vendor_doc="25012-01-4005",
            doc_title="Test doc",
        )

    def test_empty_revision_not_anomalous(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=1, dis_plan_date="2026-03-01")
        self.assertEqual(audit_revisione(rev), [])

    def test_risposta_senza_ricezione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        codes = [a.codice for a in audit_revisione(rev, time_cli=21)]
        self.assertIn("RISPOSTA_SENZA_RICEZIONE", codes)
        self.assertIn("DISPATCH_SENZA_RIENTRO_PREV", codes)

    def test_inviato_senza_dispatch(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=4,
            int_status="inviato_al_cliente",
        )
        codes = [a.codice for a in audit_revisione(rev)]
        self.assertEqual(codes, ["INVIATO_SENZA_DISPATCH"])

    def test_empty_after_approved_not_anomalous(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date="2025-01-01",
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        rev1 = Revisione.objects.create(documento=self.doc, rev_no=1)
        self.assertEqual(audit_revisione(rev1), [])
        self.assertEqual(audit_commessa("25012"), [])

    def test_post_ricezione_pending_emission_not_anomalous(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date="2025-01-01",
            rec_act_date="2025-01-10",
            ext_status=self.rejected,
            int_status="ricevuto",
        )
        rev1 = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            dis_plan_date="2025-01-24",
        )
        self.assertEqual(audit_revisione(rev1), [])
        self.assertEqual(audit_commessa("25012"), [])

    def test_int_status_incongruente_ricevuto_senza_data(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="ricevuto",
        )
        codes = [a.codice for a in audit_revisione(rev)]
        self.assertIn("INT_STATUS_INCONGRUENTE", codes)

    def test_classifica_conclusi_fixes_risposta_senza_ricezione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        result = classifica_revisione(
            "25012",
            rev.pk,
            "conclusi",
            {"rec_act_date": "2026-03-10"},
        )
        self.assertEqual(result["revisione"]["int_status"], "ricevuto")
        self.assertEqual(result["revisione"]["rec_act_date"], "2026-03-10")
        rev.refresh_from_db()
        self.assertEqual(audit_revisione(rev, time_cli=21), [])

    def test_classifica_da_emettere_clears_inviato_senza_dispatch(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=4,
            int_status="inviato_al_cliente",
        )
        classifica_revisione("25012", rev.pk, "da_emettere")
        rev.refresh_from_db()
        self.assertEqual(rev.int_status, "")
        self.assertEqual(audit_revisione(rev), [])

    def test_gruppi_marks_latest_revision(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="ricevuto",
            rec_act_date="2025-01-10",
            dis_act_date="2025-01-01",
            ext_status=self.approved,
        )
        rev1 = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            int_status="inviato_al_cliente",
        )
        gruppi = serialize_anomalie_gruppi("25012", audit_commessa("25012"))
        self.assertEqual(len(gruppi), 1)
        self.assertEqual(gruppi[0]["revisione_id"], rev1.pk)
        self.assertTrue(gruppi[0]["is_latest"])
        self.assertEqual(gruppi[0]["suggest_bucket"], "da_ricevere")
        self.assertEqual(gruppi[0]["situazione_attuale"]["label"], "Da ricevere")

    def test_ignora_anomalie_esclude_revisione(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            dis_act_date="2026-02-26",
            ext_status=self.approved,
        )
        self.assertEqual(len(audit_commessa("25012")), 2)
        ignora_anomalie_revisione("25012", rev.pk)
        rev.refresh_from_db()
        self.assertTrue(rev.ignora_anomalie)
        self.assertEqual(audit_commessa("25012"), [])
        self.assertEqual(audit_commessa_summary("25012")["count"], 0)


class EseguiRicezioneTests(TestCase):
    def setUp(self):
        self.commented = StatoEsterno.objects.create(
            nome="Commented - To be issued as Final", colore="#f59b1d", crea_nuova_rev=True
        )
        self.testata = Testata.objects.create(job="25089", time_ven_doc_rev=14)
        self.doc = Documento.objects.create(
            testata=self.testata,
            vendor_doc="25089-01-515",
            doc_title="Test doc",
        )
        self.rev0 = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date="2026-08-26",
            int_status="inviato_al_cliente",
        )

    def _entry(self):
        return {
            "doc_id": self.doc.pk,
            "rec_act_date": "2026-09-10",
            "ext_status_id": self.commented.pk,
            "crea_nuova_revisione": True,
        }

    def test_registra_ricezione_crea_nuova_revisione(self):
        esegui_ricezione([self._entry()])
        self.rev0.refresh_from_db()
        self.assertEqual(self.rev0.rec_act_date, date(2026, 9, 10))
        self.assertEqual(self.rev0.int_status, "ricevuto")
        revs = list(self.doc.revisioni.order_by("rev_no"))
        self.assertEqual(len(revs), 2)
        self.assertEqual(revs[1].rev_no, 1)
        self.assertIsNone(revs[1].rec_act_date)

    def test_invio_duplicato_su_bozza_stantia_non_corrompe_nuova_revisione(self):
        """Un doc_id già evaso, reinviato per bozza non ripulita lato client
        (vedi ricezione_detail.html._draft), non deve toccare la revisione
        appena creata: non è mai stata spedita, quindi non può essere
        "rientrata"."""
        entry = self._entry()
        esegui_ricezione([entry])
        esegui_ricezione([entry])  # invio duplicato/stantio dello stesso doc_id

        revs = list(self.doc.revisioni.order_by("rev_no"))
        self.assertEqual(len(revs), 2, "non deve essere creata una terza revisione fantasma")
        nuova_rev = revs[1]
        self.assertIsNone(nuova_rev.rec_act_date)
        self.assertIsNone(nuova_rev.ext_status_id)
        self.assertEqual(nuova_rev.int_status, "")

    def test_errore_su_una_voce_annulla_anche_le_voci_precedenti(self):
        """Tutto o niente: con una risposta cliente inesistente nella seconda
        voce, la prima (valida) non deve restare salvata come "ricevuta"
        mentre l'utente vede un errore."""
        altro_doc = Documento.objects.create(testata=self.testata, vendor_doc="25089-01-516")
        Revisione.objects.create(
            documento=altro_doc,
            rev_no=0,
            dis_act_date="2026-08-26",
            int_status="inviato_al_cliente",
        )
        voce_rotta = {**self._entry(), "doc_id": altro_doc.pk, "ext_status_id": 999999}

        with self.assertRaises(ValueError):
            esegui_ricezione([self._entry(), voce_rotta])

        self.rev0.refresh_from_db()
        self.assertIsNone(self.rev0.rec_act_date)
        self.assertEqual(self.rev0.int_status, "inviato_al_cliente")
        self.assertEqual(self.doc.revisioni.count(), 1)


class RevisioneSbloccoTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.approved = StatoEsterno.objects.create(nome="Approved", colore="#00B050")
        self.rejected = StatoEsterno.objects.create(
            nome="Rejected - Work can not proceed", colore="#D61D09"
        )
        self.testata = Testata.objects.create(job="25012")
        self.doc = Documento.objects.create(
            testata=self.testata,
            vendor_doc="25012-01-4005",
            doc_title="Test doc",
        )
        self.writing_user = User.objects.create_user(
            "sbloc_user", "sbloc@example.com", "pw", permesso=Permesso.WRITING
        )
        self.reading_user = User.objects.create_user(
            "sbloc_read", "sbloc-read@example.com", "pw", permesso=Permesso.READING
        )

    def test_list_includes_latest_with_ext_status(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            rec_act_date="2025-06-01",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili("25012")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["revisione_id"], rev.pk)

    def test_list_excludes_when_successor_exists(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1)
        self.assertEqual(list_revisioni_sbloccabili("25012"), [])

    def test_list_excludes_latest_without_ext_status(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-01-10",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1, dis_plan_date="2025-02-01")
        self.assertEqual(list_revisioni_sbloccabili("25012"), [])

    def test_list_includes_any_ext_status(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            rec_act_date="2025-06-01",
            ext_status=self.rejected,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili("25012")
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["revisione_id"], rev.pk)

    def test_sblocca_creates_next_revision(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=2,
            rec_act_date="2025-06-01",
            ext_status=self.approved,
            int_status="ricevuto",
        )
        result = sblocca_revisione("25012", rev.pk, "2025-07-15")
        new_rev = Revisione.objects.get(pk=result["revisione"]["id"])
        self.assertEqual(new_rev.rev_no, 3)
        self.assertEqual(new_rev.dis_plan_date.isoformat(), "2025-07-15")
        self.assertTrue(new_rev.crea_nuova_rev)
        self.assertEqual(new_rev.ext_status_id, None)

    def test_sblocca_without_date(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        result = sblocca_revisione("25012", rev.pk)
        new_rev = Revisione.objects.get(pk=result["revisione"]["id"])
        self.assertIsNone(new_rev.dis_plan_date)

    def test_sblocca_rejects_non_latest(self):
        rev0 = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        Revisione.objects.create(documento=self.doc, rev_no=1)
        with self.assertRaises(ValueError):
            sblocca_revisione("25012", rev0.pk)

    def test_api_list_and_sblocca(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            ext_status=self.approved,
            int_status="ricevuto",
        )
        self.client.force_login(self.reading_user)
        list_res = self.client.get("/api/commesse/25012/revisioni/sbloccabili/")
        self.assertEqual(list_res.status_code, 200)
        self.assertEqual(list_res.json()["count"], 1)

        self.client.force_login(self.reading_user)
        post_res = self.client.post(
            "/api/commesse/25012/revisioni/sblocca/",
            data=json.dumps({"revisione_id": rev.pk}),
            content_type="application/json",
        )
        self.assertEqual(post_res.status_code, 403)

        self.client.force_login(self.writing_user)
        post_res = self.client.post(
            "/api/commesse/25012/revisioni/sblocca/",
            data=json.dumps({"revisione_id": rev.pk, "dis_plan_date": "2025-08-01"}),
            content_type="application/json",
        )
        self.assertEqual(post_res.status_code, 200)
        data = post_res.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["count"], 0)
        self.assertEqual(Revisione.objects.filter(documento=self.doc, rev_no=2).count(), 1)


class RevisioneEtichettaTests(SimpleTestCase):
    """Lettera e numero sono due scritture dello stesso dato (revisione_label)."""

    def test_numero_convertito_in_lettera(self):
        self.assertEqual(numero_a_lettera(0), "A")
        self.assertEqual(numero_a_lettera(1), "B")
        self.assertEqual(numero_a_lettera(25), "Z")
        self.assertEqual(numero_a_lettera(26), "AA")
        self.assertEqual(numero_a_lettera(27), "AB")
        self.assertEqual(numero_a_lettera("3"), "D")
        self.assertEqual(numero_a_lettera(None), "")
        self.assertEqual(numero_a_lettera(-1), "")

    def test_lettera_convertita_in_numero(self):
        self.assertEqual(lettera_a_numero("A"), 0)
        self.assertEqual(lettera_a_numero("b"), 1)
        self.assertEqual(lettera_a_numero("Z"), 25)
        self.assertEqual(lettera_a_numero("AA"), 26)
        self.assertIsNone(lettera_a_numero(""))
        self.assertIsNone(lettera_a_numero("1"))
        self.assertIsNone(lettera_a_numero(None))

    def test_con_flag_attivo_si_vede_sempre_la_lettera(self):
        self.assertEqual(format_revisione_label(1, "A", True), "A")
        self.assertEqual(format_revisione_label(1, "a", True), "A")
        # Lettera non compilata: si ricava dal numero.
        self.assertEqual(format_revisione_label(0, "", True), "A")
        self.assertEqual(format_revisione_label(2, "", True), "C")
        self.assertEqual(format_revisione_label(26, None, True), "AA")
        # Lettera compilata con un numero: viene comunque convertita.
        self.assertEqual(format_revisione_label(None, "2", True), "C")
        self.assertEqual(format_revisione_label(None, "", True), "")

    def test_con_flag_spento_si_vede_sempre_il_numero(self):
        self.assertEqual(format_revisione_label(1, "A", False), "1")
        self.assertEqual(format_revisione_label(0, "", False), "0")
        # Numero non compilato: si ricava dalla lettera.
        self.assertEqual(format_revisione_label(None, "B", False), "1")
        self.assertEqual(format_revisione_label(None, "AA", False), "26")
        self.assertEqual(format_revisione_label(None, "", False), "")

    SEQUENZA = ["1", "2", "3", "D", "E", "F", "4", "5", "6"]

    def test_sequenza_personalizzata_parte_dalla_prima_revisione(self):
        etichette = [format_revisione_label(n, "", False, self.SEQUENZA) for n in range(9)]
        self.assertEqual(etichette, self.SEQUENZA)

    def test_oltre_la_sequenza_prosegue_dall_ultima_etichetta(self):
        self.assertEqual(format_revisione_label(9, "", False, self.SEQUENZA), "7")
        self.assertEqual(format_revisione_label(11, "", False, self.SEQUENZA), "9")
        casi = [
            (["0", "A", "B"], 3, "C"),
            (["A", "Z"], 2, "AA"),
            (["a", "b"], 2, "c"),
            (["R08", "R09"], 2, "R10"),
            (["1A"], 1, "1B"),
            # Ultima etichetta che non finisce né con numeri né con lettere.
            (["A-"], 2, "2"),
        ]
        for sequenza, rev_no, atteso in casi:
            with self.subTest(sequenza=sequenza, rev_no=rev_no):
                self.assertEqual(format_revisione_label(rev_no, "", False, sequenza), atteso)

    def test_sequenza_vince_sul_flag_e_ignora_la_lettera_salvata(self):
        self.assertEqual(format_revisione_label(1, "", True, ["x", "y"]), "y")
        self.assertEqual(format_revisione_label(1, "Z", False, self.SEQUENZA), "2")

    def test_sequenza_senza_numero_usa_l_etichetta_salvata(self):
        self.assertEqual(format_revisione_label(None, "e", False, ["d", "E"]), "E")
        self.assertEqual(format_revisione_label(None, "X", False, self.SEQUENZA), "X")
        self.assertEqual(format_revisione_label(None, "", False, self.SEQUENZA), "")

    def test_sequenza_vuota_lascia_numero_e_lettera(self):
        self.assertEqual(format_revisione_label(2, "", True, []), "C")
        self.assertEqual(format_revisione_label(2, "", False, None), "2")

    def test_numero_revisione_segue_l_etichetta_mostrata(self):
        self.assertEqual(numero_revisione(1, "C", True), 2)
        self.assertEqual(numero_revisione(1, "C", False), 1)
        self.assertEqual(numero_revisione(None, "D", False, self.SEQUENZA), 3)
        self.assertIsNone(numero_revisione(None, "X", False, self.SEQUENZA))

    def test_validazione_della_sequenza(self):
        self.assertEqual(valida_sequenza_revisioni(" 1, 2 ,D "), ["1", "2", "D"])
        self.assertEqual(valida_sequenza_revisioni([1, "b"]), ["1", "b"])
        self.assertEqual(valida_sequenza_revisioni(None), [])
        self.assertEqual(valida_sequenza_revisioni("  "), [])
        for valore in (
            ["1", "1"],
            ["a", "A"],
            ["1", ""],
            "1,2,",
            ["12345678901"],
            ["1,2"],
            [True],
            {"a": 1},
        ):
            with self.subTest(valore=valore), self.assertRaises(ValueError):
                valida_sequenza_revisioni(valore)


class RevisioneLabelDisplayTests(TestCase):
    """Revision display depends on Testata.rev_let_flag, not on rev_let alone."""

    def test_list_documenti_ignores_rev_let_when_flag_false(self):
        t = Testata.objects.create(job="REVFLG1", rev_let_flag=False)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG1-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="C")

        payload = list_documenti("REVFLG1")
        self.assertEqual(len(payload), 1)
        self.assertEqual(payload[0]["latest_rev_display"], "2")
        self.assertEqual(payload[0]["latest_rev_let"], "C")
        self.assertEqual(payload[0]["latest_rev_no"], 2)

    def test_list_documenti_shows_letter_when_flag_true(self):
        t = Testata.objects.create(job="REVFLG2", rev_let_flag=True)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG2-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="C")

        payload = list_documenti("REVFLG2")
        self.assertEqual(payload[0]["latest_rev_display"], "C")

    def test_list_situazione_includes_rev_let_flag(self):
        Testata.objects.create(job="REVFLG3", rev_let_flag=True)
        data = list_situazione("REVFLG3")
        self.assertTrue(data["rev_let_flag"])

    def test_list_documenti_deriva_la_lettera_quando_manca(self):
        t = Testata.objects.create(job="REVFLG4", rev_let_flag=True)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG4-01")
        Revisione.objects.create(documento=doc, rev_no=2, rev_let="")

        self.assertEqual(list_documenti("REVFLG4")[0]["latest_rev_display"], "C")

    def test_list_documenti_deriva_il_numero_quando_manca(self):
        t = Testata.objects.create(job="REVFLG5", rev_let_flag=False)
        doc = Documento.objects.create(testata=t, vendor_doc="REVFLG5-01")
        Revisione.objects.create(documento=doc, rev_no=None, rev_let="C")

        self.assertEqual(list_documenti("REVFLG5")[0]["latest_rev_display"], "2")

    def test_str_della_revisione_segue_il_flag(self):
        con_lettera = Testata.objects.create(job="REVFLG6", rev_let_flag=True)
        senza_lettera = Testata.objects.create(job="REVFLG7", rev_let_flag=False)
        doc_lettera = Documento.objects.create(testata=con_lettera, vendor_doc="REVFLG6-01")
        doc_numero = Documento.objects.create(testata=senza_lettera, vendor_doc="REVFLG7-01")
        rev_lettera = Revisione.objects.create(documento=doc_lettera, rev_no=1, rev_let="")
        rev_numero = Revisione.objects.create(documento=doc_numero, rev_no=1, rev_let="B")

        self.assertEqual(rev_lettera.etichetta(), "B")
        self.assertEqual(str(rev_lettera), "Rev B")
        self.assertEqual(rev_numero.etichetta(), "1")
        self.assertEqual(str(rev_numero), "Rev 1")

    def test_sequenza_personalizzata_in_elenco_situazione_e_str(self):
        t = Testata.objects.create(job="REVSEQ1", rev_sequenza=["1", "2", "D"])
        doc = Documento.objects.create(testata=t, vendor_doc="REVSEQ1-01")
        Revisione.objects.create(documento=doc, rev_no=1, rev_let="B")
        rev = Revisione.objects.create(documento=doc, rev_no=2, rev_let="")

        self.assertEqual(list_documenti("REVSEQ1")[0]["latest_rev_display"], "D")
        self.assertEqual(list_situazione("REVSEQ1")["rev_sequenza"], ["1", "2", "D"])
        self.assertEqual(str(rev), "Rev D")


class RevisioneLabelExportTests(TestCase):
    """Excel, PDF e trasmittal stampano la revisione come dice l'archivio."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "revexp_user",
            "revexp@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="REVEXP1", client="ACME", rev_let_flag=True)
        self.doc = Documento.objects.create(
            testata=self.testata,
            item_no="001",
            vendor_doc="REVEXP1-01",
            doc_title="Data book index",
        )
        # Nessuna lettera in archivio: deve comunque uscire "B" (rev_no 1).
        self.rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            rev_let="",
            dis_plan_date=date(2026, 3, 15),
        )

    def _foglio(self, vista):

        response = self.client.get(
            f"/api/commesse/{self.testata.job}/situazione/export/?format=xlsx&vista={vista}"
        )
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def _colonna_rev(self):
        """Valori della colonna Rev. dell'Excel verticale."""
        ws = self._foglio("verticale")
        # Le colonne fisse hanno l'intestazione in riga 1 (unita con la riga 2).
        intestazioni = [cell.value for cell in ws[1]]
        col = intestazioni.index("Rev.") + 1
        return [ws.cell(row=r, column=col).value for r in range(3, ws.max_row + 1)]

    def _gruppi_rev_orizzontale(self):
        """Intestazioni dei gruppi revisione dell'Excel orizzontale."""
        ws = self._foglio("orizzontale")
        return [c.value for c in ws[1] if str(c.value or "").startswith("Rev.")]

    def test_export_xlsx_usa_la_lettera_col_flag_attivo(self):
        self.assertEqual(self._colonna_rev(), ["B"])
        # Unico gruppo, etichettato dalla revisione presente (rev_no 1).
        self.assertEqual(self._gruppi_rev_orizzontale(), ["Rev. B"])

    def test_export_xlsx_usa_il_numero_col_flag_spento(self):
        self.testata.rev_let_flag = False
        self.testata.save(update_fields=["rev_let_flag"])
        self.rev.rev_let = "B"
        self.rev.save(update_fields=["rev_let"])

        self.assertEqual(self._colonna_rev(), ["1"])
        self.assertEqual(self._gruppi_rev_orizzontale(), ["Rev. 1"])

    def test_intestazione_pdf_situazione_segue_il_flag(self):
        from src.pdf import _rev_group_label

        revs = [[{"rev_no": 0, "rev_let": ""}, {"rev_no": 1, "rev_let": ""}]]
        self.assertEqual(_rev_group_label(revs, 0, True), "Rev. A")
        self.assertEqual(_rev_group_label(revs, 1, True), "Rev. B")
        self.assertEqual(_rev_group_label(revs, 0, False), "Rev. 0")
        # Posizione senza revisioni: l'etichetta resta coerente col flag.
        self.assertEqual(_rev_group_label([[]], 2, True), "Rev. C")
        self.assertEqual(_rev_group_label([[]], 2, False), "Rev. 2")
        # ...e con la sequenza personalizzata, anche oltre la sua fine.
        sequenza = ["1", "2", "D"]
        self.assertEqual(_rev_group_label(revs, 1, False, sequenza), "Rev. 2")
        self.assertEqual(_rev_group_label([[]], 2, True, sequenza), "Rev. D")
        self.assertEqual(_rev_group_label([[]], 4, False, sequenza), "Rev. F")

    def test_righe_pdf_verticale_seguono_il_flag(self):
        from src.pdf import _situazione_verticale_flat_rows

        documenti = list_documenti(self.testata.job)
        rev_map = revisioni_by_doc_for_job(self.testata.job)
        righe = _situazione_verticale_flat_rows(documenti, rev_map, True)
        self.assertEqual([r["rev_label"] for r in righe], ["B"])
        righe_numero = _situazione_verticale_flat_rows(documenti, rev_map, False)
        self.assertEqual([r["rev_label"] for r in righe_numero], ["1"])
        righe_sequenza = _situazione_verticale_flat_rows(documenti, rev_map, False, ["X", "Y"])
        self.assertEqual([r["rev_label"] for r in righe_sequenza], ["Y"])

    def _usa_sequenza(self, sequenza):
        self.testata.rev_let_flag = False
        self.testata.rev_sequenza = sequenza
        self.testata.save(update_fields=["rev_let_flag", "rev_sequenza"])

    def test_export_xlsx_usa_la_sequenza_personalizzata(self):
        self._usa_sequenza(["P0", "P1"])

        self.assertEqual(self._colonna_rev(), ["P1"])
        self.assertEqual(self._gruppi_rev_orizzontale(), ["Rev. P1"])

    def test_export_pdf_situazione_con_sequenza_personalizzata(self):
        self._usa_sequenza(["P0", "P1"])

        for vista in ("orizzontale", "verticale"):
            with self.subTest(vista=vista):
                response = self.client.get(
                    f"/api/commesse/{self.testata.job}/situazione/export/?vista={vista}"
                )
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.content.startswith(b"%PDF"))

    def test_trasmittal_usa_la_sequenza_personalizzata(self):
        from core.views import _trasmittal_documents

        self._usa_sequenza(["1", "2"])

        documents = _trasmittal_documents(self.testata, [self.doc.pk])
        self.assertEqual(documents[0]["rev_label"], "2")

    def test_documenti_del_trasmittal_portano_la_revisione_giusta(self):
        from core.views import _trasmittal_documents
        from src.pdf import _cell_value_for_col

        documents = _trasmittal_documents(self.testata, [self.doc.pk])
        self.assertEqual(documents[0]["rev_label"], "B")
        self.assertEqual(_cell_value_for_col(documents[0], "rev_no"), "B")

        self.testata.rev_let_flag = False
        self.testata.save(update_fields=["rev_let_flag"])
        documents = _trasmittal_documents(self.testata, [self.doc.pk])
        self.assertEqual(documents[0]["rev_label"], "1")
        self.assertEqual(_cell_value_for_col(documents[0], "rev_no"), "1")

    def test_le_pagine_ricevono_il_flag_e_l_helper_condiviso(self):
        """Le tabelle costruite in JS formattano la revisione con la regola condivisa."""
        for url in (
            f"/commesse/{self.testata.job}/documenti/",
            f"/commesse/{self.testata.job}/situazione/",
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "function formatRevLabel(rev, revLetFlag, revSeq)")
                self.assertContains(response, "REV_LET_FLAG = true")
                self.assertContains(response, 'id="rev-seq-pagina"')
                self._assert_helper_revisione_integri(response)

    def _assert_helper_revisione_integri(self, response):
        """Lo script degli helper in base.html non è spezzato da un </script> interno."""
        html = response.content.decode()
        inizio = html.index("function revNumeroInLettera")
        fine = html.index("function bindBackdropClose")
        self.assertNotIn("</script>", html[inizio:fine])

    def test_export_emissione_xlsx_usa_la_lettera(self):

        response = self.client.get(f"/api/commesse/{self.testata.job}/emissione/export/")
        self.assertEqual(response.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        col = [cell.value for cell in ws[1]].index("Rev.") + 1
        self.assertEqual(
            [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)], ["B"]
        )


class RevisioniPersonalizzateArchivioTests(TestCase):
    """Informazioni archivio → Revisioni: numero, lettera o sequenza personalizzata."""

    def setUp(self):
        self.testata = Testata.objects.create(job="26020")
        self.writer = User.objects.create_user(
            "writer_rev", "writer_rev@example.com", "pw", permesso=Permesso.WRITING
        )
        self.client = Client()
        self.client.force_login(self.writer)

    def _put(self, **dati):
        return self.client.put(
            f"/api/commesse/{self.testata.job}/", data=dati, content_type="application/json"
        )

    def test_salva_la_sequenza_ripulita(self):
        response = self._put(rev_let_flag=False, rev_sequenza=[" 1", "2 ", "D"])

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"]["rev_sequenza"], ["1", "2", "D"])
        self.testata.refresh_from_db()
        self.assertEqual(self.testata.rev_sequenza, ["1", "2", "D"])
        self.assertEqual(self.testata.rev_modalita, "personalizzata")

    def test_lettera_e_sequenza_insieme_rifiutate(self):
        response = self._put(rev_let_flag=True, rev_sequenza=["1", "2"])

        self.assertEqual(response.status_code, 422)
        self.assertIn("rev_sequenza", response.json()["error"])
        self.testata.refresh_from_db()
        self.assertFalse(self.testata.rev_let_flag)
        self.assertEqual(self.testata.rev_sequenza, [])

    def test_etichette_ripetute_rifiutate(self):
        response = self._put(rev_sequenza=["1", "D", "d"])

        self.assertEqual(response.status_code, 422)
        self.assertIn("ripetuta", response.json()["error"]["rev_sequenza"][0])

    def test_passare_alle_lettere_azzera_la_sequenza(self):
        self.testata.rev_sequenza = ["1", "2"]
        self.testata.save(update_fields=["rev_sequenza"])

        response = self._put(rev_let_flag=True, rev_sequenza=[])

        self.assertEqual(response.status_code, 200)
        self.testata.refresh_from_db()
        self.assertEqual(self.testata.rev_modalita, "lettera")

    def test_salvare_i_tempi_non_tocca_le_revisioni(self):
        self.testata.rev_sequenza = ["1", "2"]
        self.testata.save(update_fields=["rev_sequenza"])

        response = self._put(time_cli_doc_rev=10)

        self.assertEqual(response.status_code, 200)
        self.testata.refresh_from_db()
        self.assertEqual(self.testata.rev_sequenza, ["1", "2"])

    def test_modalita(self):
        self.assertEqual(Testata(job="X").rev_modalita, "numero")
        self.assertEqual(Testata(job="X", rev_let_flag=True).rev_modalita, "lettera")
        self.assertEqual(Testata(job="X", rev_sequenza=["1"]).rev_modalita, "personalizzata")
        self.assertEqual(Testata(job="X", rev_sequenza=["1"]).etichetta_prima_revisione, "1")

    def test_la_pagina_mostra_la_sezione_revisioni(self):
        self.testata.rev_sequenza = ["1", "2", "D"]
        self.testata.save(update_fields=["rev_sequenza"])

        response = self.client.get(f"/commesse/{self.testata.job}/archivio/")

        self.assertContains(response, 'id="arch-rev-seq"')
        self.assertContains(response, 'value="1, 2, D"')
        self.assertContains(response, 'value="personalizzata" checked')
        self.assertNotContains(response, 'id="arch-rev-let"')
        RevisioneLabelExportTests._assert_helper_revisione_integri(self, response)


class UltimaRevisioneTests(TestCase):
    """Una regola sola per "ultima revisione", con ``rev_no`` NULL incluso.

    PostgreSQL mette i NULL per primi in ordinamento decrescente: un
    ``order_by("-rev_no")`` prendeva la revisione senza numero (import Access)
    come ultima, mentre il codice in Python la trattava come revisione 0.
    """

    def setUp(self):
        self.testata = Testata.objects.create(job="26100", time_cli_doc_rev=10)
        self.doc = Documento.objects.create(testata=self.testata, vendor_doc="26100-A")
        self.senza_numero = Revisione.objects.create(documento=self.doc, rev_no=None)
        self.rev0 = Revisione.objects.create(documento=self.doc, rev_no=0)
        self.rev1 = Revisione.objects.create(documento=self.doc, rev_no=1)

    def test_la_revisione_senza_numero_non_e_mai_l_ultima(self):
        from core.services.ultima_revisione import (
            ultima_revisione,
            ultima_revisione_per_documento_id,
        )

        self.assertEqual(ultima_revisione(self.doc), self.rev1)
        self.assertEqual(ultima_revisione_per_documento_id(self.doc.pk), self.rev1)

    def test_dal_prefetch_stessa_risposta_e_nessuna_query(self):
        from core.services.ultima_revisione import revisione_zero, ultima_revisione

        doc = Documento.objects.prefetch_related("revisioni").get(pk=self.doc.pk)

        with self.assertNumQueries(0):
            self.assertEqual(ultima_revisione(doc), self.rev1)
            self.assertEqual(revisione_zero(doc), self.rev0)

    def test_l_emissione_scrive_sull_ultima_revisione_vera(self):
        esegui_emissione([self.doc.pk], "2026-09-02")

        self.rev1.refresh_from_db()
        self.senza_numero.refresh_from_db()
        self.assertEqual(self.rev1.int_status, "inviato_al_cliente")
        self.assertEqual(self.rev1.dis_act_date, date(2026, 9, 2))
        self.assertIsNone(self.senza_numero.dis_act_date)

    def test_sblocco_riconosce_l_ultima_revisione_vera(self):
        risposta = StatoEsterno.objects.create(nome="Commented", lettera="C", colore="#f59b1d")
        self.rev1.ext_status = risposta
        self.rev1.save(update_fields=["ext_status"])

        esito = sblocca_revisione(self.testata.job, self.rev1.pk)

        self.assertEqual(esito["revisione"]["rev_no"], 2)
        with self.assertRaises(ValueError):
            sblocca_revisione(self.testata.job, self.senza_numero.pk)

    def test_la_lista_documenti_costa_un_numero_fisso_di_query(self):
        for n in range(2, 6):
            doc = Documento.objects.create(testata=self.testata, vendor_doc=f"26100-{n}")
            Revisione.objects.create(documento=doc, rev_no=0)
            Revisione.objects.create(documento=doc, rev_no=1)

        # Testata, acronimi reparto, documenti, revisioni (prefetch): 4 query
        # per 5 documenti, non 2 in più per ciascuno.
        with self.assertNumQueries(4):
            documenti = list_documenti(self.testata.job)

        self.assertEqual(len(documenti), 5)
        self.assertEqual(documenti[0]["latest_rev_no"], 1)
        self.assertEqual(documenti[0]["latest_rev_id"], self.rev1.pk)


class StatoEsternoLetteraTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "stato_lettera_user",
            "stato_lettera@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)

    def test_letter_for_status_name_aliases(self):
        self.assertEqual(letter_for_status_name("Approved"), "A")
        self.assertEqual(
            letter_for_status_name("Commented-To be resubmitted-Work can proceed"),
            "C",
        )
        self.assertEqual(
            letter_for_status_name("Commented - To be resubmitted - Work can proceed"),
            "C",
        )
        self.assertEqual(letter_for_status_name("Unknown Status"), "")

    def test_serialize_revisione_includes_lettera(self):
        testata = Testata.objects.create(job="LET01")
        doc = Documento.objects.create(testata=testata, item_no="001", vendor_doc="VD")
        stato = StatoEsterno.objects.create(nome="Final - As Built", lettera="F", colore="#FFC000")
        rev = Revisione.objects.create(documento=doc, rev_no=0, ext_status=stato)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_lettera"], "F")
        self.assertEqual(data["ext_status_label"], "Final - As Built")

    def test_list_stati_esterni_includes_lettera(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        items = list_stati_esterni()
        self.assertEqual(items[0]["lettera"], "A")

    def test_stati_esterni_api_create_and_patch_lettera(self):
        create = self.client.post(
            "/api/stati-esterni/",
            data=json.dumps({"nome": "For Information", "lettera": "z", "colore": "#FFC000"}),
            content_type="application/json",
        )
        self.assertEqual(create.status_code, 201)
        payload = create.json()["data"]
        self.assertEqual(payload["lettera"], "Z")
        pk = payload["id"]

        patch = self.client.patch(
            f"/api/stati-esterni/{pk}/",
            data=json.dumps({"lettera": "Z"}),
            content_type="application/json",
        )
        self.assertEqual(patch.status_code, 200)
        self.assertEqual(patch.json()["data"]["lettera"], "Z")


class StatoInternoEffettivoTests(TestCase):
    """Lo stato interno mostrato quando in archivio è rimasto vuoto."""

    def setUp(self):
        self.testata = Testata.objects.create(job="INT01")
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="INT01-01-0001"
        )
        self.approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")

    def test_risposta_cliente_senza_stato_salvato_e_ricevuta(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status"], "")
        self.assertEqual(data["int_status_eff"], "ricevuto")
        self.assertEqual(data["int_status_eff_label"], "Ricevuto")

    def test_data_di_invio_senza_stato_salvato_e_inviata(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, dis_act_date=date(2025, 1, 1))
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "inviato_al_cliente")

    def test_revisione_senza_nulla_resta_da_inviare(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0)
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "")
        self.assertEqual(data["int_status_eff_label"], "Da inviare")

    def test_stato_salvato_ha_la_precedenza(self):
        rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            int_status="in_lavorazione",
            ext_status=self.approved,
        )
        data = serialize_revisione(rev)
        self.assertEqual(data["int_status_eff"], "in_lavorazione")

    def test_documento_con_risposta_cliente_non_e_da_emettere(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        doc = list_documenti(self.testata.job)[0]
        self.assertEqual(doc["latest_int_status"], "ricevuto")
        self.assertEqual(doc["latest_int_status_label"], "Ricevuto")

    def test_etichetta_export_segue_lo_stato_effettivo(self):
        from src.pdf import _int_status_export_label

        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        self.assertEqual(_int_status_export_label(serialize_revisione(rev)), "Ricevuto")
        self.assertEqual(_int_status_export_label(None), "Da inviare")
        self.assertEqual(_int_status_export_label({}), "Da inviare")
