import io
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings

from ..models import (
    Documento,
    Permesso,
    Revisione,
    Stabilimento,
    StatoEsterno,
    Testata,
)
from ..services.import_old import importa_commessa_da_access

# ── Shared fixture factory (unit tests only — uses temp dir) ──────────────────
from ._fixtures import User, _costruisci_xlsm_organizzazione, _FakeBusinessCentralSito


class ImportOldTests(TestCase):
    def _frames(self):
        return {
            "testata": pd.DataFrame(
                [
                    {
                        "Job": "99999",
                        "Client": "Cliente Test",
                        "POno": "PO-1",
                        "JobDetail": "Dettaglio",
                        "DeliveryDate": None,
                        "DeliveryTerm": "",
                        "Requisition": "",
                        "TimeCliDocRev": None,
                        "TimeVenDocRev": None,
                        "RevLetFlag": False,
                    }
                ]
            ),
            "indirsped": pd.DataFrame(),
            "dettaglio": pd.DataFrame(
                [
                    {
                        "ItemNo": "1",
                        "VendorDoc": "99999-01",
                        "ClientDocNo": "",
                        "ContractorDocNo": "CTR-99",
                        "ClientDocClass": "",
                        "DocTitle": "Titolo",
                        "DocPenalty": False,
                        "DocPayment": False,
                        "RevGen": False,
                        "Reparto": 4,
                        "Remarks": "",
                    }
                ]
            ),
            "revisioni": pd.DataFrame(
                [
                    {
                        "VendorDoc": "99999-01",
                        "RevNo": 0,
                        "RevLet": "",
                        "DisPlanDate": None,
                        "DisActDate": "2025-01-01",
                        "RecPlanDate": None,
                        "RecActDate": "2025-01-10",
                        "Status": "A",
                    },
                    {
                        "VendorDoc": "99999-01",
                        "RevNo": 1,
                        "RevLet": "",
                        "DisPlanDate": None,
                        "DisActDate": None,
                        "RecPlanDate": None,
                        "RecActDate": None,
                        "Status": None,
                    },
                ]
            ),
        }

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_includes_all_revisioni_and_reports_orphans(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        result = importa_commessa_da_access("99999")
        self.assertEqual(result["documenti"], 1)
        self.assertEqual(result["revisioni"], 2)
        self.assertEqual(result["revisioni_orfane_escluse"], 0)
        self.assertEqual(Revisione.objects.count(), 2)
        self.assertIn("anomalie", result)
        # Rev 0 has DisActDate + RecActDate → ricevuto; rev 1 has no dates → Da inviare
        self.assertEqual(Revisione.objects.filter(rev_no=0).get().int_status, "ricevuto")
        self.assertEqual(Revisione.objects.filter(rev_no=1).get().int_status, "")
        doc = Documento.objects.get(vendor_doc="99999-01")
        self.assertEqual(doc.contractor_doc_no, "CTR-99")

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_sets_ricevuto_when_client_answered_without_dates(self, mock_fetch):
        StatoEsterno.objects.create(nome="Approved", colore="#00B050", crea_nuova_rev=False)
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "A",
                    "DisPlanDate": None,
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": "A",
                }
            ]
        )
        mock_fetch.return_value = frames
        importa_commessa_da_access("99999")
        rev = Revisione.objects.get()
        self.assertEqual(rev.int_status, "ricevuto")
        self.assertIsNotNone(rev.ext_status)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_accetta_requisition_lunga(self, mock_fetch):
        """Regressione: Access può avere Requisition oltre i 100 caratteri
        storici del campo (es. due bid concatenate); il limite è 300."""
        frames = self._frames()
        requisition_lunga = "TECHNICAL REQUISITION FOR TRIM COOLERS 21- HA-104/204 " * 3
        self.assertGreater(len(requisition_lunga), 100)
        frames["testata"].loc[0, "Requisition"] = requisition_lunga
        mock_fetch.return_value = frames
        importa_commessa_da_access("99999")
        testata = Testata.objects.get(job="99999")
        self.assertEqual(testata.requisition, requisition_lunga.strip())

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_excludes_orphan_after_approved(self, mock_fetch):
        StatoEsterno.objects.create(nome="Approved", colore="#00B050", crea_nuova_rev=False)
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "",
                    "DisPlanDate": "2025-01-01",
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": "2025-01-10",
                    "RecActDate": "2025-01-10",
                    "Status": "A",
                },
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 1,
                    "RevLet": "",
                    "DisPlanDate": "2025-02-01",
                    "DisActDate": None,
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                },
            ]
        )
        mock_fetch.return_value = frames
        result = importa_commessa_da_access("99999")
        self.assertEqual(result["revisioni_orfane_escluse"], 1)
        self.assertEqual(result["revisioni"], 1)
        self.assertEqual(Revisione.objects.count(), 1)
        self.assertEqual(Revisione.objects.get().rev_no, 0)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_sets_inviato_when_only_dis_act_date(self, mock_fetch):
        frames = self._frames()
        frames["revisioni"] = pd.DataFrame(
            [
                {
                    "VendorDoc": "99999-01",
                    "RevNo": 0,
                    "RevLet": "A",
                    "DisPlanDate": None,
                    "DisActDate": "2025-01-01",
                    "RecPlanDate": None,
                    "RecActDate": None,
                    "Status": None,
                }
            ]
        )
        mock_fetch.return_value = frames
        importa_commessa_da_access("99999")
        rev = Revisione.objects.get()
        self.assertEqual(rev.int_status, "inviato_al_cliente")
        self.assertEqual(rev.dis_act_date.isoformat(), "2025-01-01")
        self.assertIsNone(rev.rec_act_date)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_import_rejects_existing_job(self, mock_fetch):
        Testata.objects.create(job="99999")
        with self.assertRaises(ValueError):
            importa_commessa_da_access("99999")
        mock_fetch.assert_not_called()

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_legge_persone_da_excel(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        xlsm_path = Path(media.name) / "organizzazione.xlsx"
        _costruisci_xlsm_organizzazione(
            xlsm_path,
            [{"Job": "99999", "PM": "ROSSI", "PE": "BIANCHI/VERDI", "WE": None, "QCI": "N/A"}],
        )
        # ROSSI ha un utente registrato con quel cognome: va abbinato. BIANCHI
        # e VERDI no: restano testo libero, da risolvere a mano.
        rossi = User.objects.create_user(
            "rossi_test", "rossi@b.it", "pw", first_name="Mario", last_name="Rossi"
        )

        with override_settings(ORGANIZZAZIONE_COMMESSE_XLSM_PATH=str(xlsm_path)):
            with self.captureOnCommitCallbacks(execute=True):
                importa_commessa_da_access("99999")

        testata = Testata.objects.get(job="99999")
        pm = testata.persone.get(ruolo="pm")
        self.assertEqual(pm.utente_id, rossi.pk)
        self.assertEqual(pm.nome_libero, "")

        liberi = testata.persone.filter(ruolo="pe")
        self.assertEqual(sorted(liberi.values_list("nome_libero", flat=True)), ["BIANCHI", "VERDI"])
        self.assertTrue(all(p.utente_id is None for p in liberi))

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_file_excel_assente_non_blocca_import(self, mock_fetch):
        mock_fetch.return_value = self._frames()

        with override_settings(
            ORGANIZZAZIONE_COMMESSE_XLSM_PATH=r"Z:\percorso\che\non\esiste.xlsm"
        ):
            with self.captureOnCommitCallbacks(execute=True):
                result = importa_commessa_da_access("99999")

        self.assertEqual(result["documenti"], 1)
        testata = Testata.objects.get(job="99999")
        self.assertEqual(testata.persone.count(), 0)

    @patch("core.services.import_old.fetch_commessa_frames")
    def test_importa_commessa_da_access_risolve_sito_costruttivo_da_bc(self, mock_fetch):
        mock_fetch.return_value = self._frames()
        bg = Stabilimento.objects.create(nome="Valbrembo", sigla="BG", codice_bc=1)
        fake = _FakeBusinessCentralSito({"99999": "1"})

        with patch("core.services.bc_sync._apri_connessione", return_value=fake):
            with self.captureOnCommitCallbacks(execute=True):
                importa_commessa_da_access("99999")

        testata = Testata.objects.get(job="99999")
        self.assertEqual(testata.sito_costruttivo_id, bg.pk)


class ImportDocumentiExcelTests(TestCase):
    """Tests for the Excel import endpoint and the import template download."""

    def setUp(self):
        self.client = Client()
        User = get_user_model()
        self.writing_user = User.objects.create_user(
            "imp_writer", "w@example.com", "pw", permesso=Permesso.WRITING
        )
        self.reading_user = User.objects.create_user(
            "imp_reader", "r@example.com", "pw", permesso=Permesso.READING
        )
        Testata.objects.create(job="IMP01")

    def _make_xlsx(self, headers, rows):
        """Build an in-memory xlsx file and return it as bytes."""

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # ── import with date column ──────────────────────────────────────────────

    def test_import_with_date_sets_dis_plan_date(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title", "Planned send date (Rev. 0)"],
            [["001", "Doc test", "2026-03-15"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(body["created"], 1)
        self.assertEqual(body["errors"], [])
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNotNone(rev0)
        self.assertIsNotNone(rev0.dis_plan_date)
        self.assertEqual(rev0.dis_plan_date.isoformat(), "2026-03-15")

    def test_import_without_date_column_leaves_dis_plan_date_null(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title"],
            [["002", "Doc senza data"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["created"], 1)
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNone(rev0.dis_plan_date)

    def test_import_with_italian_date_format(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Planned send date (Rev. 0)"],
            [["003", "20/07/2026"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(body["created"], 1)
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertEqual(rev0.dis_plan_date.isoformat(), "2026-07-20")

    def test_import_with_invalid_date_creates_doc_and_reports_warning(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Planned send date (Rev. 0)"],
            [["004", "not-a-date"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        # Document is still created despite the invalid date
        self.assertEqual(body["created"], 1)
        self.assertEqual(len(body["errors"]), 1)
        self.assertIn("data non riconosciuta", body["errors"][0])
        rev0 = Revisione.objects.filter(documento__testata_id="IMP01", rev_no=0).first()
        self.assertIsNone(rev0.dis_plan_date)

    def test_import_with_contractor_doc_no(self):
        self.client.force_login(self.writing_user)
        xlsx = self._make_xlsx(
            ["Item", "Document title", "Contractor Doc N°"],
            [["005", "Doc contractor", "CTR-001"]],
        )
        resp = self.client.post(
            "/api/commesse/IMP01/import-excel/",
            data={"file": xlsx},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["created"], 1)
        doc = Documento.objects.get(testata_id="IMP01", item_no="005")
        self.assertEqual(doc.contractor_doc_no, "CTR-001")

    # ── template download ────────────────────────────────────────────────────

    def test_template_download_returns_xlsx(self):
        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            resp.get("Content-Type", ""),
        )

    def test_template_download_contains_date_column(self):

        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Planned send date (Rev. 0)", headers)
        # All expected columns must be present
        for col in ["Item", "B&R Doc", "Contractor Doc N°", "Document title", "Notes"]:
            self.assertIn(col, headers)

    def test_template_download_reparto_has_dropdown_validation(self):

        from openpyxl.utils import get_column_letter

        from core.models import Reparto

        Reparto.objects.create(nome="Qualità e Controllo", acronimo="QC")
        Reparto.objects.create(nome="Project Management", acronimo="PM")

        self.client.force_login(self.reading_user)
        resp = self.client.get("/api/import-documenti-template/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))

        self.assertIn("_Reparti", wb.sheetnames)
        ws_reparti = wb["_Reparti"]
        reparti_in_sheet = [ws_reparti.cell(row=r, column=1).value for r in range(1, 3)]
        self.assertEqual(
            sorted(reparti_in_sheet),
            ["Project Management", "Qualità e Controllo"],
        )

        ws = wb.active
        reparto_col = get_column_letter([cell.value for cell in ws[1]].index("Department") + 1)
        validations = [
            dv
            for dv in ws.data_validations.dataValidation
            if any(reparto_col in str(r) for r in dv.sqref.ranges)
        ]
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].type, "list")
        self.assertIn("_Reparti", validations[0].formula1)
