import io
from datetime import date, datetime

import openpyxl
from django.test import Client, TestCase
from django.utils import timezone

from ..models import (
    Documento,
    IndirSped,
    Permesso,
    Revisione,
    StatoEsterno,
    Testata,
)
from ..services.export_grezzo import _cell as _cella_grezza
from ..services.export_grezzo import build_workbook as build_dati_grezzi_workbook
from ..services.export_grezzo import list_tabelle as list_tabelle_grezze
from ..services.export_grezzo import resolve_tabelle as resolve_tabelle_grezze

# ── Shared fixture factory (unit tests only — uses temp dir) ──────────────────
from ._fixtures import User


class ScaricaDatiGrezziTestCase(TestCase):
    """Pagina "Scarica" ed export Excel dei dati grezzi di una commessa."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "grezzi_reader",
            "grezzi-reader@example.com",
            "pw",
            permesso=Permesso.READING,
        )
        self.testata = Testata.objects.create(job="GR01", client="ACME Spa")
        self.altra = Testata.objects.create(job="GR02", client="Altro Cliente")

        self.indirizzo = IndirSped.objects.create(
            testata=self.testata,
            consignee="ACME Spa",
            city="Bergamo",
            country="IT",
        )
        self.doc = Documento.objects.create(
            testata=self.testata,
            item_no="001",
            vendor_doc="GR01-QMDBI",
            client_doc_no="CL-001",
            client_doc_class="CLASSE-A",
            contractor_doc_no="CTR-001",
            doc_title="Data book index",
        )
        self.stato = StatoEsterno.objects.create(nome="Approved", lettera="A")
        self.rev = Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_plan_date=date(2026, 3, 15),
            int_status="da_emettere",
            ext_status=self.stato,
        )

        # Dati della seconda commessa: non devono finire nell'export di GR01.
        self.doc_altra = Documento.objects.create(
            testata=self.altra,
            item_no="999",
            doc_title="Documento di un'altra commessa",
        )

    def _scarica(self, job="GR01", tabelle="documenti"):
        return self.client.get(f"/api/commesse/{job}/dati-grezzi/?tabelle={tabelle}")

    def _workbook(self, response):

        return openpyxl.load_workbook(io.BytesIO(response.content))

    def _righe(self, ws):
        return [list(row) for row in ws.iter_rows(values_only=True)]

    # ── Pagina HTML ──────────────────────────────────────────────────────────

    def test_pagina_richiede_autenticazione(self):
        res = self.client.get("/scarica/")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login/", res["Location"])

    def test_pagina_elenca_tutte_le_tabelle_e_il_seleziona_tutte(self):
        self.client.force_login(self.user)
        res = self.client.get("/scarica/")
        self.assertEqual(res.status_code, 200)
        self.assertTemplateUsed(res, "core/scarica.html")
        self.assertContains(res, 'id="tabelle-all"')
        self.assertContains(res, "Seleziona tutte")
        self.assertContains(res, 'id="commessa-search"')
        for tabella in list_tabelle_grezze():
            self.assertContains(res, f'value="{tabella["key"]}"')
            self.assertContains(res, tabella["label"])

    def test_navbar_mostra_scarica_tra_cerca_e_impostazioni(self):
        self.client.force_login(self.user)
        for url in ("/", "/commesse/", "/segnalazioni/", "/scarica/"):
            with self.subTest(url=url):
                res = self.client.get(url)
                self.assertEqual(res.status_code, 200)
                html = res.content.decode()
                self.assertIn('href="/scarica/"', html)
                self.assertIn("Scarica", html)
                # La voce sta fra il pulsante "Cerca" e il link "Impostazioni".
                self.assertLess(html.index("openSpotlight()"), html.index('href="/scarica/"'))
                self.assertLess(html.index('href="/scarica/"'), html.index('href="/admin/"'))

    # ── Export ───────────────────────────────────────────────────────────────

    def test_export_richiede_autenticazione(self):
        self.assertEqual(self._scarica().status_code, 401)

    def test_export_restituisce_un_xlsx_con_nome_file_dedicato(self):
        self.client.force_login(self.user)
        res = self._scarica()
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", res["Content-Disposition"])
        self.assertIn("GR01_dati_grezzi_", res["Content-Disposition"])
        self.assertIn(".xlsx", res["Content-Disposition"])

    def test_un_foglio_per_ogni_tabella_richiesta(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="documenti,revisioni")
        wb = self._workbook(res)
        self.assertEqual(wb.sheetnames, ["Documenti", "Revisioni"])

    def test_i_fogli_seguono_lordine_canonico_non_quello_richiesto(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni,testate,documenti"))
        self.assertEqual(wb.sheetnames, ["Commessa", "Documenti", "Revisioni"])

    def test_seleziona_tutte_scarica_ogni_tabella(self):
        self.client.force_login(self.user)
        keys = ",".join(t["key"] for t in list_tabelle_grezze())
        wb = self._workbook(self._scarica(tabelle=keys))
        self.assertEqual(wb.sheetnames, [t["label"] for t in list_tabelle_grezze()])

    def test_header_con_i_nomi_colonna_e_dato_grezzo(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        righe = self._righe(wb["Documenti"])
        headers = righe[0]
        self.assertEqual(headers[0], "id")
        self.assertIn("testata_id", headers)
        self.assertIn("doc_title", headers)
        self.assertEqual(len(righe), 2)
        riga = dict(zip(headers, righe[1]))
        self.assertEqual(riga["id"], self.doc.pk)
        self.assertEqual(riga["testata_id"], "GR01")
        self.assertEqual(riga["item_no"], "001")
        self.assertEqual(riga["doc_title"], "Data book index")
        self.assertFalse(riga["doc_penalty"])

    def test_solo_header_in_grassetto_e_nessuna_colorazione(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        ws = wb["Documenti"]
        for cell in ws[1]:
            self.assertTrue(cell.font.bold)
        for cell in ws[2]:
            self.assertFalse(cell.font.bold)
            self.assertEqual(cell.fill.fill_type, None)

    def test_export_contiene_solo_i_dati_della_commessa_scelta(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="documenti"))
        ids = [r[0] for r in self._righe(wb["Documenti"])]
        self.assertNotIn(self.doc_altra.pk, ids)

        wb_altra = self._workbook(self._scarica(job="GR02", tabelle="documenti"))
        righe = self._righe(wb_altra["Documenti"])
        self.assertEqual(len(righe), 2)
        self.assertEqual(dict(zip(righe[0], righe[1]))["id"], self.doc_altra.pk)

    def test_revisioni_e_indirizzi_esportano_le_righe_collegate(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="testate,indirizzi_spedizione,revisioni"))

        commessa = self._righe(wb["Commessa"])
        self.assertEqual(dict(zip(commessa[0], commessa[1]))["job"], "GR01")

        indirizzi = self._righe(wb["Indirizzi spedizione"])
        self.assertEqual(dict(zip(indirizzi[0], indirizzi[1]))["city"], "Bergamo")

        revisioni = self._righe(wb["Revisioni"])
        riga_rev = dict(zip(revisioni[0], revisioni[1]))
        self.assertEqual(riga_rev["rev_no"], 0)
        self.assertEqual(riga_rev["int_status"], "da_emettere")

    def test_revisioni_portano_i_dati_del_documento_e_la_lettera_dello_stato(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        righe = self._righe(wb["Revisioni"])
        headers = righe[0]
        self.assertEqual(
            headers,
            [
                "client_doc_no",
                "client_doc_class",
                "contractor_doc_no",
                "vendor_doc",
                "item_no",
                "rev_no",
                "rev_let",
                "dis_plan_date",
                "dis_act_date",
                "rec_plan_date",
                "rec_act_date",
                "int_status",
                "ext_status",
                "crea_nuova_rev",
            ],
        )
        riga = dict(zip(headers, righe[1]))
        self.assertEqual(riga["client_doc_no"], "CL-001")
        self.assertEqual(riga["client_doc_class"], "CLASSE-A")
        self.assertEqual(riga["contractor_doc_no"], "CTR-001")
        self.assertEqual(riga["vendor_doc"], "GR01-QMDBI")
        self.assertEqual(riga["item_no"], "001")
        self.assertEqual(riga["ext_status"], "A")

    def test_revisione_senza_risposta_cliente_ha_la_lettera_vuota(self):
        Revisione.objects.filter(pk=self.rev.pk).update(ext_status=None)
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        righe = self._righe(wb["Revisioni"])
        # openpyxl rilegge una cella vuota come None.
        self.assertIn(dict(zip(righe[0], righe[1]))["ext_status"], ("", None))

    def test_le_tabelle_esportabili_sono_solo_le_quattro_utili(self):
        self.assertEqual(
            [t["key"] for t in list_tabelle_grezze()],
            ["testate", "indirizzi_spedizione", "documenti", "revisioni"],
        )
        for key in ("revisioni_file_link", "trasmittal", "trasmittal_revisioni"):
            with self.subTest(key=key), self.assertRaises(ValueError):
                resolve_tabelle_grezze([key])

    def test_tabella_vuota_esporta_il_solo_header(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(job="GR02", tabelle="indirizzi_spedizione"))
        righe = self._righe(wb["Indirizzi spedizione"])
        self.assertEqual(len(righe), 1)
        self.assertIn("consignee", righe[0])

    def test_le_date_restano_date_e_i_timestamp_perdono_il_fuso(self):
        self.client.force_login(self.user)
        wb = self._workbook(self._scarica(tabelle="revisioni"))
        revisioni = self._righe(wb["Revisioni"])
        self.assertEqual(
            dict(zip(revisioni[0], revisioni[1]))["dis_plan_date"],
            datetime(2026, 3, 15),
        )
        # Formato italiano e colonna abbastanza larga, altrimenti Excel mostra ########.
        ws = wb["Revisioni"]
        colonna = revisioni[0].index("dis_plan_date") + 1
        self.assertEqual(ws.cell(row=2, column=colonna).number_format, "DD/MM/YYYY")
        lettera = ws.cell(row=1, column=colonna).column_letter
        self.assertGreaterEqual(ws.column_dimensions[lettera].width, 12)
        aware = timezone.make_aware(datetime(2026, 3, 15, 8, 30))
        self.assertEqual(_cella_grezza(aware), timezone.localtime(aware).replace(tzinfo=None))
        self.assertIsNone(_cella_grezza(aware).tzinfo)
        self.assertEqual(_cella_grezza(["1", "2", "D"]), "1,2,D")

    # ── Errori ───────────────────────────────────────────────────────────────

    def test_senza_tabelle_selezionate_400(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="")
        self.assertEqual(res.status_code, 400)
        self.assertIn("almeno una tabella", res.json()["error"])

    def test_tabella_sconosciuta_400(self):
        self.client.force_login(self.user)
        res = self._scarica(tabelle="documenti,pippo")
        self.assertEqual(res.status_code, 400)
        self.assertIn("pippo", res.json()["error"])

    def test_commessa_inesistente_404(self):
        self.client.force_login(self.user)
        res = self._scarica(job="NOPE")
        self.assertEqual(res.status_code, 404)
        self.assertEqual(res.json()["error"], "Commessa non trovata.")

    # ── Service ──────────────────────────────────────────────────────────────

    def test_resolve_tabelle_deduplica_e_ordina(self):
        tabelle = resolve_tabelle_grezze(["revisioni", "documenti", "revisioni"])
        self.assertEqual([t.key for t in tabelle], ["documenti", "revisioni"])

    def test_resolve_tabelle_rifiuta_selezioni_non_valide(self):
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze([])
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze(["   "])
        with self.assertRaises(ValueError):
            resolve_tabelle_grezze(["documenti", "sconosciuta"])

    def test_build_workbook_su_commessa_inesistente(self):
        with self.assertRaises(Testata.DoesNotExist):
            build_dati_grezzi_workbook("NOPE", ["documenti"])
