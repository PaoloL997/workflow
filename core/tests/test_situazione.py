import io
import json
from datetime import date
from unittest.mock import patch

import openpyxl
from django.test import Client, SimpleTestCase, TestCase

from ..models import (
    Documento,
    Permesso,
    Revisione,
    StatoEsterno,
    Testata,
)
from ..services.commesse import (
    filtra_situazione,
    serialize_revisione,
)
from ..services.revisione_sblocco import list_revisioni_sbloccabili
from ..services.situazione_cella_vendor import colori_cella_vendor
from ..services.stato_esterno_codes import DEFAULT_STATUS_COLORS
from ..services.stato_esterno_colori import (
    DA_INVIARE_BG,
    FALLBACK_BG,
    INVIATO_BG,
    MIN_CONTRAST,
    TEXT_DARK,
    TEXT_LIGHT,
    cell_colors,
    contrast_ratio,
    hex_to_rgb,
    readable_text_color,
    rgb_to_hex,
)
from ..services.stato_esterno_legenda import legenda_default, legenda_stati_esterni

# ── Shared fixture factory (unit tests only — uses temp dir) ──────────────────
from ._fixtures import User, _make_fixtures_mock


class SituazioneApiTestCase(TestCase):
    """Tests for batch situazione/documenti API endpoints."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "situazione_user",
            "situazione@brembanarolle.com",
            "pw",
            permesso=Permesso.READING,
        )
        self.client.force_login(self.user)
        self.fixtures = _make_fixtures_mock()
        self.job = self.fixtures["testata"].job

    def test_situazione_api_returns_documenti_and_revisioni(self):
        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["documenti"]), 2)
        revs_by_doc = data["revisioni_by_doc"]
        doc_qmdbi_id = str(self.fixtures["doc_qmdbi"].pk)
        doc_qcpa_id = str(self.fixtures["doc_qcpa"].pk)
        self.assertEqual(len(revs_by_doc[doc_qmdbi_id]), 1)
        self.assertEqual(len(revs_by_doc[doc_qcpa_id]), 3)

    def test_documenti_api_include_revisioni(self):
        response = self.client.get(f"/api/commesse/{self.job}/documenti/?include_revisioni=1")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["documenti"]), 2)
        doc_qcpa_id = str(self.fixtures["doc_qcpa"].pk)
        self.assertEqual(len(data["revisioni_by_doc"][doc_qcpa_id]), 3)

    def test_situazione_api_requires_auth(self):
        self.client.logout()
        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 401)

    def test_situazione_api_not_found(self):
        response = self.client.get("/api/commesse/INEXISTENT/situazione/")
        self.assertEqual(response.status_code, 404)

    def test_export_situazione_pdf(self):
        """Il PDF della situazione: risposta, nome file e dati passati al generatore.

        Senza un parser PDF fra le dipendenze non si legge il testo del file:
        si verifica invece cosa riceve ``genera_situazione_documenti_pdf``
        (documenti e vista), lasciando il generatore vero a produrre il PDF.
        """
        from src.pdf import genera_situazione_documenti_pdf

        today = date.today()
        date_part = f"{today.day:02d}_{today.month:02d}_{today.year}"

        for vista in ("orizzontale", "verticale"):
            with (
                self.subTest(vista=vista),
                patch(
                    "core.views.genera_situazione_documenti_pdf",
                    wraps=genera_situazione_documenti_pdf,
                ) as generatore,
            ):
                response = self.client.get(
                    f"/api/commesse/{self.job}/situazione/export/?vista={vista}"
                )

                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/pdf")
                self.assertIn(
                    f'filename="{self.job}_document_status_{date_part}.pdf"',
                    response["Content-Disposition"],
                )
                self.assertTrue(response.content.startswith(b"%PDF"))
                chiamata = generatore.call_args
                self.assertEqual(chiamata.kwargs["vista"], vista)
                vendor_docs = {d["vendor_doc"] for d in chiamata.kwargs["documenti"]}
                self.assertIn("25056-QMDBI", vendor_docs)

    def test_export_situazione_pdf_hides_empty_fixed_columns(self):
        from src.pdf import _active_fixed_doc_columns

        documenti = [
            {"vendor_doc": "DOC-1", "doc_title": "", "client_doc_no": ""},
            {"vendor_doc": "DOC-2", "doc_penalty": True, "item_no": "42"},
        ]
        active = _active_fixed_doc_columns(documenti)
        labels = [col["label"] for col in active]
        self.assertEqual(labels, ["Vendor Document", "Item", "Penalty"])

    def test_export_situazione_not_found(self):
        response = self.client.get("/api/commesse/INEXISTENT/situazione/export/")
        self.assertEqual(response.status_code, 404)

    def test_situazione_api_includes_ext_status_lettera(self):
        approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        doc = self.fixtures["doc_qmdbi"]
        rev = doc.revisioni.order_by("rev_no").first()
        rev.ext_status = approved
        rev.save(update_fields=["ext_status"])

        response = self.client.get(f"/api/commesse/{self.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        revs = response.json()["revisioni_by_doc"][str(doc.pk)]
        self.assertEqual(revs[0]["ext_status_lettera"], "A")
        self.assertEqual(revs[0]["ext_status_label"], "Approved")


class ColoriRisposteClienteTests(SimpleTestCase):
    """Colori delle celle legate alle risposte del cliente (viste situazione)."""

    # ── Conversioni ──────────────────────────────────────────────────────────

    def test_hex_to_rgb_accetta_le_forme_valide(self):
        self.assertEqual(hex_to_rgb("#00B050"), (0, 176, 80))
        self.assertEqual(hex_to_rgb("00b050"), (0, 176, 80))
        self.assertEqual(hex_to_rgb("  #FFC000 "), (255, 192, 0))
        self.assertEqual(hex_to_rgb("#FFF"), (255, 255, 255))

    def test_hex_to_rgb_rifiuta_i_valori_non_validi(self):
        for value in ["", None, "verde", "#12", "#12345", "#GGGGGG"]:
            self.assertIsNone(hex_to_rgb(value), value)

    def test_rgb_to_hex_normalizza_in_maiuscolo(self):
        self.assertEqual(rgb_to_hex((0, 176, 80)), "#00B050")
        self.assertEqual(rgb_to_hex((255.4, -3, 300)), "#FF00FF")

    # ── Scelta del testo ─────────────────────────────────────────────────────

    def test_testo_nero_su_colori_chiari(self):
        # Bianco e quasi bianco: il testo bianco sparirebbe.
        for colore in ["#FFFFFF", "#FEFEFE", "#FFFF00", "#FFC000", "#9E9E9E"]:
            self.assertEqual(readable_text_color(colore), TEXT_DARK, colore)

    def test_testo_bianco_su_colori_scuri(self):
        # Nero e grigi scuri: il testo nero sparirebbe.
        for colore in ["#000000", "#111111", "#404040", "#0070C0", "#D61D09"]:
            self.assertEqual(readable_text_color(colore), TEXT_LIGHT, colore)

    def test_testo_di_default_quando_il_colore_manca(self):
        self.assertEqual(readable_text_color(""), readable_text_color(FALLBACK_BG))

    # ── Colori della cella ───────────────────────────────────────────────────

    def test_fondo_uguale_al_colore_configurato(self):
        # Uniformità con la legenda: il colore inserito non viene sbiadito.
        for colore in ["#00B050", "#FFC000", "#FFFFFF", "#000000", "#D61D09"]:
            self.assertEqual(cell_colors(colore)["bg"], colore, colore)

    def test_bianco_e_nero_restano_leggibili(self):
        self.assertEqual(cell_colors("#FFFFFF"), {"bg": "#FFFFFF", "fg": TEXT_DARK})
        self.assertEqual(cell_colors("#000000"), {"bg": "#000000", "fg": TEXT_LIGHT})

    def test_mezzi_toni_spostati_quel_tanto_che_basta(self):
        # Su #797979 nessuno dei due testi arriva al contrasto minimo: il fondo
        # si sposta di pochi punti, restando lo stesso colore a vista.
        cella = cell_colors("#797979")
        self.assertNotEqual(cella["bg"], "#797979")
        self.assertGreaterEqual(contrast_ratio(cella["bg"], cella["fg"]), MIN_CONTRAST)
        scarto = max(abs(a - b) for a, b in zip(hex_to_rgb("#797979"), hex_to_rgb(cella["bg"])))
        self.assertLessEqual(scarto, 25)

    def test_contrasto_minimo_garantito_su_tutta_la_gamma(self):
        passo = 51  # campiona 6 livelli per canale
        for r in range(0, 256, passo):
            for g in range(0, 256, passo):
                for b in range(0, 256, passo):
                    colore = rgb_to_hex((r, g, b))
                    cella = cell_colors(colore)
                    self.assertGreaterEqual(
                        contrast_ratio(cella["bg"], cella["fg"]),
                        MIN_CONTRAST,
                        colore,
                    )
                    scarto = max(abs(a - b) for a, b in zip((r, g, b), hex_to_rgb(cella["bg"])))
                    self.assertLessEqual(scarto, 25, colore)

    def test_colore_mancante_o_invalido_usa_il_fallback(self):
        atteso = {"bg": FALLBACK_BG, "fg": readable_text_color(FALLBACK_BG)}
        for colore in ["", None, "  ", "non-un-colore"]:
            self.assertEqual(cell_colors(colore), atteso, colore)


class ColoriRisposteClienteApiTests(TestCase):
    """I colori pronti per le celle viaggiano nei payload delle viste."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "colori_user",
            "colori@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="COL01")
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="COL01-01-0001"
        )
        self.approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        self.final = StatoEsterno.objects.create(
            nome="Final - As Built", lettera="F", colore="#FFFFFF"
        )

    def test_serialize_revisione_include_i_colori_della_cella(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_colore"], "#00B050")
        self.assertEqual(data["ext_status_bg"], "#00B050")
        self.assertEqual(data["ext_status_fg"], TEXT_DARK)

    def test_serialize_revisione_su_stato_bianco_usa_testo_nero(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.final)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_bg"], "#FFFFFF")
        self.assertEqual(data["ext_status_fg"], TEXT_DARK)

    def test_serialize_revisione_senza_risposta_non_ha_colori(self):
        rev = Revisione.objects.create(documento=self.doc, rev_no=0)
        data = serialize_revisione(rev)
        self.assertEqual(data["ext_status_bg"], "")
        self.assertEqual(data["ext_status_fg"], "")

    def test_situazione_api_include_i_colori_della_cella(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, ext_status=self.approved)
        response = self.client.get(f"/api/commesse/{self.testata.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        revs = response.json()["revisioni_by_doc"][str(self.doc.pk)]
        self.assertEqual(revs[0]["ext_status_bg"], "#00B050")
        self.assertEqual(revs[0]["ext_status_fg"], TEXT_DARK)

    def test_revisioni_sbloccabili_includono_i_colori_della_cella(self):
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            rec_act_date="2025-06-01",
            ext_status=self.final,
            int_status="ricevuto",
        )
        items = list_revisioni_sbloccabili(self.testata.job)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["ext_status_bg"], "#FFFFFF")
        self.assertEqual(items[0]["ext_status_fg"], TEXT_DARK)

    def test_pdf_colora_la_cella_vendor_con_il_colore_pieno(self):
        from src.pdf import _status_cell_rgb

        self.assertEqual(_status_cell_rgb("#00B050"), ((0, 176, 80), hex_to_rgb(TEXT_DARK)))
        self.assertEqual(_status_cell_rgb("#D61D09"), ((214, 29, 9), hex_to_rgb(TEXT_LIGHT)))
        self.assertIsNone(_status_cell_rgb(""))

    def test_situazione_api_porta_i_colori_dell_inviato_al_cliente(self):
        response = self.client.get(f"/api/commesse/{self.testata.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.json()["inviato_cell"],
            {"bg": INVIATO_BG, "fg": TEXT_DARK, "label": "Inviato al Cliente"},
        )


class CellaVendorDocTests(TestCase):
    """Cella B&R Doc: il colore è l'ultimo stato del documento."""

    def setUp(self):
        self.testata = Testata.objects.create(job="ATT01")
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="ATT01-01"
        )
        self.commented = StatoEsterno.objects.create(
            nome="Commented", lettera="C", colore="#D61D09"
        )

    def _revs(self):
        return [
            serialize_revisione(r)
            for r in self.doc.revisioni.select_related("ext_status").order_by("rev_no")
        ]

    def _rev_con_risposta(self, rev_no=0):
        return Revisione.objects.create(
            documento=self.doc,
            rev_no=rev_no,
            dis_act_date=date(2026, 1, 10),
            rec_act_date=date(2026, 1, 20),
            ext_status=self.commented,
        )

    def test_colore_della_risposta_sull_ultima_revisione(self):
        self._rev_con_risposta()
        self.assertEqual(
            colori_cella_vendor(self._revs()),
            {"bg": "#D61D09", "fg": TEXT_LIGHT, "label": "Commented", "lettera": "C"},
        )

    def test_revisione_nuova_non_ancora_inviata_tiene_l_ultima_risposta(self):
        self._rev_con_risposta()
        # La Rev. 1 non è partita: il grigio è riservato a chi non ha mai avuto
        # una risposta, quindi resta il colore della risposta sulla Rev. 0.
        Revisione.objects.create(documento=self.doc, rev_no=1, dis_plan_date=date(2026, 4, 1))
        self.assertEqual(
            colori_cella_vendor(self._revs()),
            {"bg": "#D61D09", "fg": TEXT_LIGHT, "label": "Commented", "lettera": "C"},
        )

    def test_giallo_quando_la_revisione_attende_risposta(self):
        self._rev_con_risposta()
        # La revisione nuova è partita: conta lei, non la risposta sulla Rev. 0.
        Revisione.objects.create(documento=self.doc, rev_no=1, dis_act_date=date(2026, 2, 1))
        self.assertEqual(
            colori_cella_vendor(self._revs()),
            {"bg": INVIATO_BG, "fg": TEXT_DARK, "label": "Inviato al Cliente", "lettera": ""},
        )

    def test_giallo_anche_sulla_prima_revisione_appena_inviata(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, dis_act_date=date(2026, 1, 10))
        self.assertEqual(
            colori_cella_vendor(self._revs()),
            {"bg": INVIATO_BG, "fg": TEXT_DARK, "label": "Inviato al Cliente", "lettera": ""},
        )

    def test_grigio_solo_senza_risposte_e_senza_invii(self):
        Revisione.objects.create(documento=self.doc, rev_no=0, dis_plan_date=date(2026, 4, 1))
        self.assertEqual(
            colori_cella_vendor(self._revs()),
            {"bg": DA_INVIARE_BG, "fg": TEXT_DARK, "label": "Da inviare", "lettera": ""},
        )

    def test_documento_senza_revisioni_resta_senza_colore(self):
        self.assertEqual(colori_cella_vendor([]), {"bg": "", "fg": "", "label": "", "lettera": ""})

    def test_pdf_usa_gli_stessi_colori(self):
        from src.pdf import _vendor_doc_cell_color

        self._rev_con_risposta()
        self.assertEqual(_vendor_doc_cell_color(self._revs()), "#D61D09")

        Revisione.objects.create(documento=self.doc, rev_no=1, dis_act_date=date(2026, 2, 1))
        self.assertEqual(_vendor_doc_cell_color(self._revs()), INVIATO_BG)

        # Rev. 1 non più inviata: torna l'ultima risposta, non il grigio.
        Revisione.objects.filter(rev_no=1).update(dis_act_date=None)
        self.assertEqual(_vendor_doc_cell_color(self._revs()), "#D61D09")

        Revisione.objects.filter(documento=self.doc).delete()
        Revisione.objects.create(documento=self.doc, rev_no=0)
        self.assertEqual(_vendor_doc_cell_color(self._revs()), DA_INVIARE_BG)


class SituazioneOrizzontaleXlsxTests(TestCase):
    """L'Excel della vista orizzontale riproduce la tabella a schermo."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "sitxls_user",
            "sitxls@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="SITXLS1", rev_let_flag=False)
        self.commented = StatoEsterno.objects.create(
            nome="Commented", lettera="C", colore="#D61D09"
        )
        self.doc = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="SITXLS1-01", doc_title="Index"
        )
        Revisione.objects.create(
            documento=self.doc,
            rev_no=0,
            dis_act_date=date(2026, 1, 10),
            rec_act_date=date(2026, 1, 20),
            ext_status=self.commented,
        )
        Revisione.objects.create(
            documento=self.doc,
            rev_no=1,
            dis_act_date=date(2026, 2, 1),
            rec_plan_date=date(2026, 3, 1),
        )
        Documento.objects.create(testata=self.testata, item_no="002", vendor_doc="SITXLS1-02")

    def _foglio(self):

        response = self.client.get(
            f"/api/commesse/{self.testata.job}/situazione/export/?format=xlsx&vista=orizzontale"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'filename="situazione_documenti_orizzontale_SITXLS1.xlsx"',
            response["Content-Disposition"],
        )
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def _riga(self, ws, vendor_doc):
        col = [c.value for c in ws[1]].index("B&R Doc") + 1
        return next(r for r in range(3, ws.max_row + 1) if ws.cell(r, col).value == vendor_doc)

    def test_una_riga_per_documento_con_gruppi_per_revisione(self):
        ws = self._foglio()
        intestazioni = [c.value for c in ws[1]]
        self.assertEqual(
            intestazioni[:6],
            ["Client Doc N°", "Client Doc Class", "Contractor Doc N°", "B&R Doc", "Title", "Item"],
        )
        self.assertEqual(
            [v for v in intestazioni if v], intestazioni[:6] + ["Planning", "Rev. 0", "Rev. 1"]
        )
        self.assertEqual(
            [c.value for c in ws[2]][6:],
            ["Submission date", "Receipt date"] + ["Dispatch", "Received", "Status"] * 2,
        )
        # Due righe di intestazione + una riga per ciascuno dei due documenti.
        self.assertEqual(ws.max_row, 4)

    def test_valori_di_planning_e_revisioni(self):
        from ..date_fmt import format_display_date

        ws = self._foglio()
        row = self._riga(ws, "SITXLS1-01")
        col = [c.value for c in ws[1]].index("Rev. 0") + 1
        self.assertEqual(
            [ws.cell(row, c).value for c in range(col, col + 6)],
            [
                format_display_date(date(2026, 1, 10)),
                format_display_date(date(2026, 1, 20)),
                "C",
                format_display_date(date(2026, 2, 1)),
                None,
                None,
            ],
        )
        # Ultima revisione inviata: resta pianificato solo il rientro.
        self.assertEqual(ws.cell(row, col - 2).value, None)
        self.assertEqual(ws.cell(row, col - 1).value, format_display_date(date(2026, 3, 1)))

    def test_solo_b_r_doc_e_colorata_con_la_risposta(self):
        # Ultima revisione già rientrata: la cella B&R Doc segue quella risposta,
        # mentre le celle Status restano senza colore.
        doc = Documento.objects.create(testata=self.testata, item_no="003", vendor_doc="SITXLS1-03")
        Revisione.objects.create(
            documento=doc,
            rev_no=0,
            dis_act_date=date(2026, 1, 10),
            rec_act_date=date(2026, 1, 20),
            ext_status=self.commented,
        )
        ws = self._foglio()
        row = self._riga(ws, "SITXLS1-03")
        intestazioni = [c.value for c in ws[1]]
        vendor = ws.cell(row, intestazioni.index("B&R Doc") + 1)
        status_rev0 = ws.cell(row, intestazioni.index("Rev. 0") + 3)
        self.assertEqual(vendor.fill.fgColor.rgb[-6:], "D61D09")
        self.assertEqual(vendor.font.color.rgb[-6:], TEXT_LIGHT.lstrip("#"))
        self.assertEqual(status_rev0.value, "C")
        self.assertIsNone(status_rev0.fill.fill_type)

        senza_revisioni = ws.cell(self._riga(ws, "SITXLS1-02"), intestazioni.index("B&R Doc") + 1)
        self.assertIsNone(senza_revisioni.fill.fill_type)

    def test_b_r_doc_grigia_finche_la_revisione_non_e_partita(self):
        doc = Documento.objects.create(testata=self.testata, item_no="004", vendor_doc="SITXLS1-04")
        Revisione.objects.create(documento=doc, rev_no=0, dis_plan_date=date(2026, 4, 1))
        ws = self._foglio()
        intestazioni = [c.value for c in ws[1]]
        vendor = ws.cell(self._riga(ws, "SITXLS1-04"), intestazioni.index("B&R Doc") + 1)
        self.assertEqual(vendor.fill.fgColor.rgb[-6:], DA_INVIARE_BG.lstrip("#"))
        self.assertEqual(vendor.font.color.rgb[-6:], TEXT_DARK.lstrip("#"))

    def test_b_r_doc_giallo_se_l_ultima_revisione_attende_risposta(self):
        ws = self._foglio()
        row = self._riga(ws, "SITXLS1-01")
        intestazioni = [c.value for c in ws[1]]
        vendor = ws.cell(row, intestazioni.index("B&R Doc") + 1)
        status_rev1 = ws.cell(row, intestazioni.index("Rev. 1") + 3)
        # La Rev. 1 è partita e non ha risposta: vince il giallo dell'inviato al
        # cliente, non il rosso della risposta arrivata sulla Rev. 0.
        self.assertEqual(vendor.fill.fgColor.rgb[-6:], INVIATO_BG.lstrip("#"))
        self.assertEqual(vendor.font.color.rgb[-6:], TEXT_DARK.lstrip("#"))
        self.assertIsNone(status_rev1.fill.fill_type)

    def test_la_pagina_apre_il_menu_excel_pdf_in_ogni_vista(self):
        response = self.client.get(f"/commesse/{self.testata.job}/situazione/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="btn-export" type="button" onclick="toggleExportMenu(')
        self.assertNotContains(response, "exportSituazione(")

    def test_la_pagina_ha_i_controlli_di_filtro_in_barra(self):
        response = self.client.get(f"/commesse/{self.testata.job}/situazione/")
        self.assertContains(response, 'id="btn-stato-ultima-rev"')
        self.assertContains(response, "stato_ultima_rev")
        self.assertContains(response, 'id="btn-solo-ultima-rev"')
        self.assertContains(response, "toggleSoloUltimaRev(event)")


class FiltraSituazioneTests(SimpleTestCase):
    """``filtra_situazione`` restringe il payload a quello che si vede a schermo."""

    def setUp(self):
        self.payload = {
            "documenti": [{"id": 1}, {"id": 2}, {"id": 3}],
            "revisioni_by_doc": {1: [{"id": 10}, {"id": 11}], 2: [{"id": 20}]},
            "rev_let_flag": True,
            "inviato_cell": {"bg": "#F4C325"},
        }

    def test_senza_id_il_payload_resta_quello(self):
        self.assertIs(filtra_situazione(self.payload), self.payload)

    def test_documenti_nell_ordine_ricevuto_e_id_ignoti_scartati(self):
        out = filtra_situazione(self.payload, doc_ids=[3, 1, 999])
        self.assertEqual([d["id"] for d in out["documenti"]], [3, 1])
        self.assertEqual(out["revisioni_by_doc"], {3: [], 1: [{"id": 10}, {"id": 11}]})

    def test_lista_vuota_non_vale_come_nessun_filtro(self):
        out = filtra_situazione(self.payload, doc_ids=[])
        self.assertEqual(out["documenti"], [])
        self.assertEqual(out["revisioni_by_doc"], {})

    def test_rev_ids_taglia_le_revisioni_dei_documenti_tenuti(self):
        out = filtra_situazione(self.payload, doc_ids=[1, 2], rev_ids=[11, 20])
        self.assertEqual(out["revisioni_by_doc"], {1: [{"id": 11}], 2: [{"id": 20}]})

    def test_id_non_numerici_ignorati_e_resto_del_payload_intatto(self):
        out = filtra_situazione(self.payload, doc_ids=["2", None, "x"])
        self.assertEqual([d["id"] for d in out["documenti"]], [2])
        self.assertTrue(out["rev_let_flag"])
        self.assertEqual(out["inviato_cell"], {"bg": "#F4C325"})

    def test_chiavi_stringa_delle_revisioni_json(self):
        payload = {**self.payload, "revisioni_by_doc": {"1": [{"id": 10}]}}
        out = filtra_situazione(payload, doc_ids=[1])
        self.assertEqual(out["revisioni_by_doc"], {1: [{"id": 10}]})


class SituazioneExportFiltratoTests(TestCase):
    """L'export contiene solo i documenti (e le revisioni) rimasti a schermo."""

    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(
            "expfiltro_user",
            "expfiltro@brembanarolle.com",
            "pw",
            permesso=Permesso.WRITING,
        )
        self.client.force_login(self.user)
        self.testata = Testata.objects.create(job="EXPF1", rev_let_flag=False)
        self.commented = StatoEsterno.objects.create(
            nome="Commented EXPF", lettera="C", colore="#D61D09"
        )
        self.doc_a = Documento.objects.create(
            testata=self.testata, item_no="001", vendor_doc="EXPF1-01"
        )
        self.rev_a0 = Revisione.objects.create(
            documento=self.doc_a,
            rev_no=0,
            dis_act_date=date(2026, 1, 10),
            rec_act_date=date(2026, 1, 20),
            ext_status=self.commented,
        )
        self.rev_a1 = Revisione.objects.create(
            documento=self.doc_a, rev_no=1, dis_act_date=date(2026, 2, 1)
        )
        self.doc_b = Documento.objects.create(
            testata=self.testata, item_no="002", vendor_doc="EXPF1-02"
        )
        self.rev_b0 = Revisione.objects.create(documento=self.doc_b, rev_no=0)
        self.doc_c = Documento.objects.create(
            testata=self.testata, item_no="003", vendor_doc="EXPF1-03"
        )

    def _url(self, vista="orizzontale", fmt="xlsx"):
        return f"/api/commesse/{self.testata.job}/situazione/export/?format={fmt}&vista={vista}"

    def _post(self, corpo, vista="orizzontale", fmt="xlsx"):
        return self.client.post(
            self._url(vista, fmt),
            data=json.dumps(corpo),
            content_type="application/json",
        )

    def _foglio(self, response):

        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def _vendor_doc_righe(self, ws):
        col = [c.value for c in ws[1]].index("B&R Doc") + 1
        return [ws.cell(r, col).value for r in range(3, ws.max_row + 1) if ws.cell(r, col).value]

    def test_orizzontale_esporta_solo_i_documenti_indicati(self):
        ws = self._foglio(self._post({"doc_ids": [self.doc_b.pk, self.doc_a.pk]}))
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-02", "EXPF1-01"])

    def test_orizzontale_i_gruppi_revisione_seguono_i_documenti_esportati(self):
        ws = self._foglio(self._post({"doc_ids": [self.doc_b.pk]}))
        intestazioni = [c.value for c in ws[1]]
        self.assertIn("Rev. 0", intestazioni)
        self.assertNotIn("Rev. 1", intestazioni)

    def test_verticale_esporta_solo_le_revisioni_indicate(self):
        # Come con "solo ultima revisione" attivo: una riga per documento.
        ws = self._foglio(
            self._post(
                {
                    "doc_ids": [self.doc_a.pk, self.doc_b.pk],
                    "rev_ids": [self.rev_a1.pk, self.rev_b0.pk],
                },
                vista="verticale",
            )
        )
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-01", "EXPF1-02"])
        col_rev = [c.value for c in ws[1]].index("Rev.") + 1
        self.assertEqual([ws.cell(r, col_rev).value for r in (3, 4)], ["1", "0"])

    def test_verticale_documento_senza_revisioni_resta_una_riga(self):
        ws = self._foglio(
            self._post({"doc_ids": [self.doc_c.pk], "rev_ids": []}, vista="verticale")
        )
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-03"])

    def test_rev_ids_nullo_tiene_tutte_le_revisioni(self):
        ws = self._foglio(
            self._post({"doc_ids": [self.doc_a.pk], "rev_ids": None}, vista="verticale")
        )
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-01", "EXPF1-01"])

    def test_doc_ids_vuoto_non_esporta_nessuna_riga(self):
        ws = self._foglio(self._post({"doc_ids": []}))
        self.assertEqual(self._vendor_doc_righe(ws), [])

    def test_post_senza_selezione_esporta_tutto(self):
        ws = self._foglio(self._post({}))
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-01", "EXPF1-02", "EXPF1-03"])

    def test_documenti_di_altre_commesse_vengono_ignorati(self):
        altra = Testata.objects.create(job="EXPF2")
        estraneo = Documento.objects.create(testata=altra, item_no="001", vendor_doc="EXPF2-01")
        ws = self._foglio(self._post({"doc_ids": [estraneo.pk, self.doc_a.pk]}))
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-01"])

    def test_json_non_valido(self):
        res = self.client.post(self._url(), data="non-json", content_type="application/json")
        self.assertEqual(res.status_code, 400)

    def test_get_continua_a_esportare_tutto(self):
        ws = self._foglio(self.client.get(self._url()))
        self.assertEqual(self._vendor_doc_righe(ws), ["EXPF1-01", "EXPF1-02", "EXPF1-03"])

    def test_pdf_filtrato(self):
        res = self._post({"doc_ids": [self.doc_a.pk]}, fmt="pdf")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertTrue(res.content.startswith(b"%PDF"))

    def test_senza_csrf_token_il_post_e_rifiutato(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        res = client.post(
            self._url(), data=json.dumps({"doc_ids": []}), content_type="application/json"
        )
        self.assertEqual(res.status_code, 403)


class LegendaStatusPdfTests(TestCase):
    """La legenda STATUS dell'header PDF segue le risposte messe a sistema."""

    def _colori_disegnati(self, vista="orizzontale", **kwargs):
        """Colori passati a ``set_fill_color`` generando il PDF situazione."""
        from src.pdf import SituazioneDocumentiPDF, genera_situazione_documenti_pdf

        originale = SituazioneDocumentiPDF.set_fill_color
        chiamate = []

        def spia(pdf, *args):
            chiamate.append(tuple(args))
            return originale(pdf, *args)

        with patch.object(SituazioneDocumentiPDF, "set_fill_color", spia):
            genera_situazione_documenti_pdf(
                {"job": "LEG01", "client": "ACME"},
                documenti=[],
                revisioni_by_doc={},
                vista=vista,
                **kwargs,
            )
        return chiamate

    # ── Voci della legenda ───────────────────────────────────────────────────

    def test_legenda_usa_lettera_nome_e_colore_configurati(self):
        StatoEsterno.objects.create(nome="Rejected", lettera="R", colore="#8B0000")
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        self.assertEqual(
            legenda_stati_esterni(),
            [
                {"lettera": "A", "nome": "Approved", "colore": "#0070C0"},
                {"lettera": "R", "nome": "Rejected", "colore": "#8B0000"},
            ],
        )

    def test_lettera_dedotta_dal_nome_quando_non_e_configurata(self):
        StatoEsterno.objects.create(nome="Final - As Built", colore="#FFC000")
        self.assertEqual(legenda_stati_esterni()[0]["lettera"], "F")

    def test_colore_mancante_ricade_sul_default_della_lettera(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="")
        self.assertEqual(legenda_stati_esterni()[0]["colore"], DEFAULT_STATUS_COLORS["A"])

    def test_risposta_senza_lettera_riconoscibile_resta_fuori(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        StatoEsterno.objects.create(nome="Stato inventato", colore="#123456")
        self.assertEqual([voce["lettera"] for voce in legenda_stati_esterni()], ["A"])

    def test_senza_risposte_a_sistema_resta_la_legenda_canonica(self):
        self.assertEqual(legenda_stati_esterni(), legenda_default())
        self.assertEqual(
            [voce["lettera"] for voce in legenda_default()],
            ["A", "C", "F", "I", "O", "R", "S", "Z"],
        )

    # ── Righe disegnate nel PDF ──────────────────────────────────────────────

    def test_le_voci_del_pdf_usano_il_fondo_delle_celle(self):
        from src.pdf import _status_legend_entries

        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#FFFFFF")
        self.assertEqual(_status_legend_entries(), [("A", "Approved", (255, 255, 255))])

    def test_voce_con_colore_non_valido_usa_il_fallback(self):
        from src.pdf import _status_legend_entries

        voci = [{"lettera": "A", "nome": "Approved", "colore": "non-un-colore"}]
        self.assertEqual(_status_legend_entries(voci)[0][2], hex_to_rgb(FALLBACK_BG))

    def test_le_voci_si_fermano_alle_righe_disponibili(self):
        from src.pdf import _LEGEND_MAX_ROWS, _status_legend_entries

        voci = [
            {"lettera": chr(ord("A") + i), "nome": f"Stato {i}", "colore": "#00B050"}
            for i in range(_LEGEND_MAX_ROWS + 3)
        ]
        self.assertEqual(len(_status_legend_entries(voci)), _LEGEND_MAX_ROWS)

    def test_le_righe_restano_dentro_il_riquadro_status(self):
        from src.pdf import _HEADER_BOTTOM_PT, _LEGEND_MAX_ROWS, _legend_row_ys

        otto = _legend_row_ys(8)
        # Otto voci = layout storico del foglio di riferimento.
        self.assertEqual(len(otto), 8)
        self.assertAlmostEqual(otto[0], 49.5)
        self.assertAlmostEqual(otto[1] - otto[0], 9.7)
        for count in range(1, _LEGEND_MAX_ROWS + 1):
            ys = _legend_row_ys(count)
            self.assertEqual(len(ys), count)
            self.assertLess(ys[-1], _HEADER_BOTTOM_PT, count)

    # ── PDF generato ─────────────────────────────────────────────────────────

    def test_il_pdf_orizzontale_usa_i_colori_a_sistema(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        colori = self._colori_disegnati(vista="orizzontale")
        self.assertIn((0, 112, 192), colori)
        # Il vecchio verde fisso non viene più disegnato.
        self.assertNotIn((0, 176, 80), colori)

    def test_il_pdf_verticale_usa_i_colori_a_sistema(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#0070C0")
        colori = self._colori_disegnati(vista="verticale")
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)

    def test_cambiare_colore_a_sistema_cambia_la_legenda(self):
        approved = StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        self.assertIn((0, 176, 80), self._colori_disegnati())
        approved.colore = "#0070C0"
        approved.save(update_fields=["colore"])
        colori = self._colori_disegnati()
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)

    def test_legenda_passata_a_mano_ha_la_precedenza(self):
        StatoEsterno.objects.create(nome="Approved", lettera="A", colore="#00B050")
        voci = [{"lettera": "A", "nome": "Approved", "colore": "#0070C0"}]
        colori = self._colori_disegnati(status_legend=voci)
        self.assertIn((0, 112, 192), colori)
        self.assertNotIn((0, 176, 80), colori)
